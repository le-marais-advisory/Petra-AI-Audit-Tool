"""Support-tab tie-out (FA calibration, item 15).

Allocation columns are often fed from support tabs (a fee calculation, a distribution
waterfall, an expense breakout) with per-investor lookups such as
``SUMIFS('Mgmt Fee Calc'!$F:$F,'Mgmt Fee Calc'!$B:$B,$C10)``. For each such column the
dominant lookup is re-evaluated against the support tab's cached values for every
investor, so a hardcoded or mislinked cell is caught, and the support tab's investor
rows are totalled so an amount that no Allocation investor pulls is caught too.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal

from src.pipeline.workbook.cells import CENT, is_text, norm_text, to_decimal, to_money
from src.pipeline.workbook.checks._common import CheckContext, NotApplicable, Outcome, check, differs, mag, money

ZERO = Decimal("0")

# Roles with their own tie rules (ITD, Summary, Merge) or no amounts (investor register, holidays).
_NOT_SUPPORT = {"itd", "summary", "merge", "investor_data", "holiday_calendar"}

_CALL_RE = re.compile(r"\b(SUMIFS|SUMIF|INDEX|VLOOKUP)\(", re.I)
_REF_RE = re.compile(r"^(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!)?([A-Z]{1,3})(\d*)(?::([A-Z]{1,3})(\d*))?$")
_PLUG_RE = re.compile(r"\)\s*[+-]\s*\d+(?:\.\d+)?\s*$")


@dataclass(frozen=True)
class _Ref:
    sheet: str | None  # None: the Allocation sheet itself
    column: str
    first_row: int | None
    last_column: str
    last_row: int | None


@dataclass(frozen=True)
class _Lookup:
    """One per-investor lookup into a support tab."""

    sheet: str
    value_column: str
    rows: tuple[int, int] | None  # None: the whole column
    keys: tuple[tuple[str, str, str], ...]  # (support key column, "row" | "cell" | "text", source)
    first_match: bool  # INDEX/MATCH and VLOOKUP return one row; SUMIF(S) add every match


def _args(formula: str, start: int) -> tuple[list[str], int]:
    depth, current, out, i = 0, "", [], start
    in_quote = False
    while i < len(formula):
        ch = formula[i]
        if ch == "'" or ch == '"':
            in_quote = not in_quote
        if not in_quote:
            if ch == "(":
                depth += 1
            elif ch == ")":
                if depth == 0:
                    out.append(current.strip())
                    return out, i
                depth -= 1
            elif ch == "," and depth == 0:
                out.append(current.strip())
                current = ""
                i += 1
                continue
        current += ch
        i += 1
    return out, i


def _ref(text: str) -> _Ref | None:
    m = _REF_RE.match(text.replace("$", "").strip())
    if not m:
        return None
    quoted, bare, col, row, col2, row2 = m.groups()
    sheet = quoted.replace("''", "'") if quoted else bare
    return _Ref(sheet, col, int(row) if row else None, col2 or col, int(row2) if row2 else (int(row) if row else None))


def _key(arg: str, row: int, home: str) -> tuple[str, str] | None:
    """How a criterion is sourced: the investor's own row, a fixed Allocation cell, or a literal."""
    if arg.startswith('"') and arg.endswith('"'):
        return "text", arg[1:-1]
    ref = _ref(arg)
    if ref is None or ref.sheet not in (None, home) or ref.first_row is None:
        return None
    if ref.first_row == row and "$" not in arg.split("!")[-1].split(ref.column)[-1][:1]:
        return "row", ref.column
    return "cell", f"{ref.column}{ref.first_row}"


def _rows(ref: _Ref) -> tuple[int, int] | None:
    return (ref.first_row, ref.last_row) if ref.first_row is not None else None


def parse_lookup(formula: str | None, row: int, home: str) -> _Lookup | None:
    if not formula:
        return None
    m = _CALL_RE.search(formula)
    if not m:
        return None
    name = m.group(1).upper()
    args, _ = _args(formula, m.end())
    try:
        if name == "SUMIFS":
            value = _ref(args[0])
            pairs = [(_ref(args[i]), _key(args[i + 1], row, home)) for i in range(1, len(args) - 1, 2)]
            if value is None or value.sheet is None or not pairs or any(r is None or k is None for r, k in pairs):
                return None
            keys = tuple((r.column, *k) for r, k in pairs)
            return _Lookup(value.sheet, value.column, _rows(value), keys, False)
        if name == "SUMIF":
            crit_range, crit = _ref(args[0]), _key(args[1], row, home)
            value = _ref(args[2]) if len(args) > 2 else crit_range
            if crit_range is None or value is None or crit is None or value.sheet is None:
                return None
            return _Lookup(value.sheet, value.column, _rows(value), ((crit_range.column, *crit),), False)
        if name == "INDEX":
            value = _ref(args[0])
            inner = re.match(r"MATCH\(", args[1], re.I) if len(args) > 1 else None
            if value is None or value.sheet is None or not inner:
                return None
            match_args, _ = _args(args[1], inner.end())
            crit, key_range = _key(match_args[0], row, home), _ref(match_args[1])
            if key_range is None or crit is None:
                return None
            return _Lookup(value.sheet, value.column, _rows(value), ((key_range.column, *crit),), True)
        if name == "VLOOKUP":
            crit, table = _key(args[0], row, home), _ref(args[1])
            index = int(to_decimal(args[2]) or 0)
            if crit is None or table is None or table.sheet is None or index < 1:
                return None
            from openpyxl.utils.cell import column_index_from_string, get_column_letter

            value_column = get_column_letter(column_index_from_string(table.column) + index - 1)
            return _Lookup(table.sheet, value_column, _rows(table), ((table.column, *crit),), True)
    except (IndexError, ValueError):
        return None
    return None


def _norm(value) -> str:
    number = None if is_text(value) else to_decimal(value)
    return str(number.normalize()) if number is not None else norm_text(value)


def _criteria(lookup: _Lookup, alloc_sheet, row: int) -> tuple[str, ...]:
    out = []
    for _, kind, source in lookup.keys:
        if kind == "text":
            out.append(norm_text(source))
        else:
            out.append(_norm(alloc_sheet.value(f"{source}{row}" if kind == "row" else source)))
    return tuple(out)


def _is_aggregate(cell, column: str) -> bool:
    """A subtotal / total cell: its formula only reads its own column (SUM(F9:F105), F108+F112)."""
    if not cell or not cell.formula:
        return False
    refs = re.findall(r"\$?([A-Z]{1,3})\$?\d+", cell.formula.replace("SUM", ""))
    return bool(refs) and all(r == column for r in refs)


def _support_rows(ctx: CheckContext, lookup: _Lookup):
    """(row, criteria tuple, label, amount) for each investor-like row of the support tab."""
    sheet = ctx.model.sheet(lookup.sheet)
    rows = range(lookup.rows[0], lookup.rows[1] + 1) if lookup.rows else range(1, sheet.max_row + 1)
    for r in rows:
        keys = tuple(_norm(sheet.value(f"{k}{r}")) for k, _, _ in lookup.keys)
        cell = sheet.cell(f"{lookup.value_column}{r}")
        amount = to_decimal(cell.value) if cell is not None and not is_text(cell.value) else None
        if amount is None or not keys[0] or "total" in keys[0] or _is_aggregate(cell, lookup.value_column):
            continue
        yield r, keys, str(sheet.value(f"{lookup.keys[0][0]}{r}")).strip(), amount


@check("CE-TIE-SUPPORT-TABS", needs=("allocation",))
def tie_support_tabs(ctx: CheckContext, out: Outcome) -> str:
    alloc = ctx.data.allocation
    roles = {name: getattr(layout, "role", None) for name, layout in ctx.data.layouts.items()}
    investors = alloc.investors
    tied: list[str] = []
    for comp in alloc.components:
        col = comp.column
        lookups = {inv.row: parse_lookup(inv.formulas.get(col), inv.row, alloc.sheet.name) for inv in investors}
        counts = Counter(lk for lk in lookups.values() if lk is not None and lk.sheet != alloc.sheet.name
                         and ctx.model.has_sheet(lk.sheet) and roles.get(lk.sheet) not in _NOT_SUPPORT)
        if not counts:
            continue
        pattern, n = counts.most_common(1)[0]
        if n * 2 <= len(investors):
            continue  # most investors are not fed from this tab
        support = ctx.model.sheet(pattern.sheet)
        rows = list(_support_rows(ctx, pattern))
        pulled: set[tuple[str, ...]] = set()
        allocated_total = ZERO
        for inv in investors:
            criteria = _criteria(pattern, alloc.sheet, inv.row)
            pulled.add(criteria)
            matches = [amount for _, keys, _, amount in rows if keys == criteria]
            expected = (matches[0] if matches else ZERO) if pattern.first_match else sum(matches, ZERO)
            actual = inv.amounts.get(col, ZERO)
            allocated_total += actual
            if differs(mag(actual), mag(expected), CENT):
                formula = inv.formulas.get(col) or ""
                message = (f"{inv.name}: {comp.header or col} is {money(actual)} on the Allocation but "
                           f"{money(expected)} on {pattern.sheet} (column {pattern.value_column}).")
                if lookups[inv.row] == pattern and _PLUG_RE.search(formula):
                    out.review(message + " The cell adds a plug on top of the support figure.", alloc.sheet,
                               f"{col}{inv.row}")
                else:
                    out.fail(message, alloc.sheet, f"{col}{inv.row}")
        support_total = ZERO
        for r, keys, label, amount in rows:
            support_total += amount
            if amount and keys not in pulled:
                out.fail(f"{pattern.sheet} row {r} ({label}) carries {money(amount)} in column "
                         f"{pattern.value_column} that no Allocation investor pulls into {comp.header or col}.",
                         support, f"{pattern.value_column}{r}")
        if differs(mag(allocated_total), mag(support_total), CENT):
            out.fail(f"{comp.header or col}: Allocation investors total {money(allocated_total)}; {pattern.sheet} "
                     f"column {pattern.value_column} totals {money(support_total)}.", alloc.sheet, f"{col}{alloc.layout.header_row}")
        tied.append(f"{comp.header or col} <- {pattern.sheet}!{pattern.value_column} ({money(support_total)})")
    if not tied:
        raise NotApplicable("No Allocation column is fed from a support tab by per-investor lookups.")
    return "Every support-fed Allocation column ties to its support tab investor by investor and in total: " \
           + "; ".join(tied) + "."
