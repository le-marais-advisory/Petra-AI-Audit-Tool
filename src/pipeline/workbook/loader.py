"""Load an .xlsx/.xlsm workbook into a sparse, read-only model.

The workbook is read twice with openpyxl - once for stored formulas and once for the
cached values Excel saved - so both are available without recalculating. Macros are
never executed (openpyxl does not run VBA; the project is not loaded).
"""
from __future__ import annotations

import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

ERROR_VALUES = {"#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!", "#SPILL!", "#CALC!"}

# Workbooks whose formulas lack cached values (saved by a tool that does not
# recalculate) cannot be checked from values; above this share they are flagged.
_MISSING_CACHE_THRESHOLD = 0.02


@dataclass
class CellModel:
    coord: str
    row: int
    column: str
    value: Any  # cached value for formula cells, literal otherwise
    formula: str | None  # stored formula including the leading '=', or None
    number_format: str = "General"

    @property
    def is_error(self) -> bool:
        return isinstance(self.value, str) and self.value.strip() in ERROR_VALUES

    @property
    def is_formula(self) -> bool:
        return self.formula is not None

    @property
    def column_index(self) -> int:
        return column_index_from_string(self.column)


@dataclass
class SheetModel:
    name: str
    index: int  # 1-based position in the workbook
    state: str  # visible | hidden | veryHidden
    view: str  # normal | pageBreakPreview | pageLayout
    dimensions: str
    max_row: int
    max_column: int
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[str] = field(default_factory=set)
    outlined_cols: set[str] = field(default_factory=set)  # columns inside an outline (grouped) range
    merged_ranges: list[str] = field(default_factory=list)
    print_area: str | None = None
    cells: dict[str, CellModel] = field(default_factory=dict)

    def cell(self, coord: str) -> CellModel | None:
        return self.cells.get(coord.replace("$", "").upper())

    def value(self, coord: str) -> Any:
        cell = self.cell(coord)
        return None if cell is None else cell.value

    def row_cells(self, row: int) -> list[CellModel]:
        return sorted((c for c in self.cells.values() if c.row == row), key=lambda c: c.column_index)

    def column_cells(self, column: str) -> list[CellModel]:
        return sorted((c for c in self.cells.values() if c.column == column), key=lambda c: c.row)

    def rows(self) -> Iterator[tuple[int, list[CellModel]]]:
        by_row: dict[int, list[CellModel]] = {}
        for cell in self.cells.values():
            by_row.setdefault(cell.row, []).append(cell)
        for row in sorted(by_row):
            yield row, sorted(by_row[row], key=lambda c: c.column_index)

    @property
    def is_visible(self) -> bool:
        return self.state == "visible"

    @property
    def max_column_letter(self) -> str:
        return get_column_letter(max(1, self.max_column))


@dataclass
class WorkbookModel:
    path: Path
    file_name: str
    sheets: list[SheetModel]
    defined_names: list[str] = field(default_factory=list)
    formula_count: int = 0
    formulas_without_cache: int = 0

    def sheet(self, name: str) -> SheetModel:
        for sheet in self.sheets:
            if sheet.name == name:
                return sheet
        raise KeyError(f"No sheet named {name!r}")

    def has_sheet(self, name: str) -> bool:
        return any(s.name == name for s in self.sheets)

    @property
    def formulas_missing_cache(self) -> bool:
        if not self.formula_count:
            return False
        return self.formulas_without_cache / self.formula_count > _MISSING_CACHE_THRESHOLD


def _column_flags(ws) -> tuple[set[str], set[str]]:
    """(hidden columns, outlined columns): a hidden column inside an outline group is a collapsed group."""
    hidden: set[str] = set()
    outlined: set[str] = set()
    for key, dim in ws.column_dimensions.items():
        if not dim.hidden and not dim.outlineLevel:
            continue
        start = dim.min or column_index_from_string(key)
        end = dim.max or start
        for index in range(start, end + 1):
            if dim.hidden:
                hidden.add(get_column_letter(index))
            if dim.outlineLevel:
                outlined.add(get_column_letter(index))
    return hidden, outlined


def load_workbook_model(path: str | Path, file_name: str | None = None) -> WorkbookModel:
    path = Path(path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about unsupported extensions / data validation
        wb_formulas = openpyxl.load_workbook(path, data_only=False, keep_vba=False)
        wb_values = openpyxl.load_workbook(path, data_only=True, keep_vba=False)

    sheets: list[SheetModel] = []
    formula_count = 0
    missing_cache = 0
    for index, ws in enumerate(wb_formulas.worksheets, start=1):
        wv = wb_values[ws.title]
        cells: dict[str, CellModel] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                raw = cell.value
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                if formula is None and not isinstance(raw, (str, int, float, bool)) and hasattr(raw, "text"):
                    formula = str(getattr(raw, "text"))  # array / data-table formulas
                    formula = formula if formula.startswith("=") else "=" + formula
                value = wv[cell.coordinate].value if formula is not None else raw
                if formula is not None:
                    formula_count += 1
                    if value is None:
                        missing_cache += 1
                coord = cell.coordinate
                col, row_number = coordinate_from_string(coord)
                cells[coord] = CellModel(
                    coord=coord,
                    row=row_number,
                    column=col,
                    value=value,
                    formula=formula,
                    number_format=cell.number_format or "General",
                )
        print_area = ws.print_area if isinstance(ws.print_area, str) else None
        hidden_cols, outlined_cols = _column_flags(ws)
        sheets.append(
            SheetModel(
                name=ws.title,
                index=index,
                state=ws.sheet_state,
                view=ws.sheet_view.view or "normal",
                dimensions=ws.dimensions,
                max_row=ws.max_row,
                max_column=ws.max_column,
                hidden_rows={r for r, dim in ws.row_dimensions.items() if dim.hidden},
                hidden_cols=hidden_cols,
                outlined_cols=outlined_cols,
                merged_ranges=sorted(str(rng) for rng in ws.merged_cells.ranges),
                print_area=print_area,
                cells=cells,
            )
        )
    return WorkbookModel(
        path=path,
        file_name=file_name or path.name,
        sheets=sheets,
        defined_names=list(wb_formulas.defined_names.keys()),
        formula_count=formula_count,
        formulas_without_cache=missing_cache,
    )
