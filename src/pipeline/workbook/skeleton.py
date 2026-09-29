"""Compact, structure-preserving text view of a sheet for the LLM layout mapper.

Keeps every text label (headers, row labels down the sheet), the merged ranges and
hidden rows/columns, numbers only for the top of the sheet, and formulas compressed
into relative patterns: a column whose cells share one pattern becomes one line, with
the cells that deviate (e.g. a residual plug) listed as exceptions.
"""
from __future__ import annotations

import datetime as dt
import re
from collections import defaultdict

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string

from src.pipeline.workbook.loader import CellModel, SheetModel

_REF_RE = re.compile(r"(?<![A-Za-z0-9_.])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![0-9(])")
_STRING_RE = re.compile(r'"[^"]*"')
_NUMERIC_ROWS = 6  # show numeric values for this many non-empty rows at the top (covers the driver row)
_MAX_TEXT = 60
_MAX_EXCEPTIONS = 8


def estimate_tokens(text: str) -> int:
    return len(text) // 4


def _normalize(formula: str, row: int, col: int, relative_cols: bool = False) -> str:
    """Rewrite relative references as offsets from the host cell ({r}, {r-2}, {c+1})."""

    def repl(m: re.Match) -> str:
        col_abs, col_letters, row_abs, row_digits = m.groups()
        ref_row = int(row_digits)
        if row_abs:
            row_part = f"${ref_row}"
        else:
            delta = ref_row - row
            row_part = "{r}" if delta == 0 else f"{{r{delta:+d}}}"
        if relative_cols and not col_abs:
            delta_c = column_index_from_string(col_letters) - col
            col_part = "{c}" if delta_c == 0 else f"{{c{delta_c:+d}}}"
        else:
            col_part = f"{col_abs}{col_letters}"
        return f"{col_part}{row_part}"

    out, last = [], 0
    for sm in _STRING_RE.finditer(formula):
        out.append(_REF_RE.sub(repl, formula[last:sm.start()]))
        out.append(sm.group(0))
        last = sm.end()
    out.append(_REF_RE.sub(repl, formula[last:]))
    return "".join(out)


def _ranges(numbers: list[int]) -> str:
    numbers = sorted(numbers)
    parts, start, prev = [], numbers[0], numbers[0]
    for n in numbers[1:]:
        if n == prev + 1:
            prev = n
            continue
        parts.append(f"{start}" if start == prev else f"{start}-{prev}")
        start = prev = n
    parts.append(f"{start}" if start == prev else f"{start}-{prev}")
    return ",".join(parts)


def _col_ranges(cols: list[int]) -> str:
    cols = sorted(cols)
    parts, start, prev = [], cols[0], cols[0]
    for c in cols[1:]:
        if c == prev + 1:
            prev = c
            continue
        parts.append(get_column_letter(start) if start == prev else f"{get_column_letter(start)}-{get_column_letter(prev)}")
        start = prev = c
    parts.append(get_column_letter(start) if start == prev else f"{get_column_letter(start)}-{get_column_letter(prev)}")
    return ",".join(parts)


def _fmt_value(value) -> str:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.6g}" if abs(value) < 1 else f"{value:.2f}".rstrip("0").rstrip(".")
    if isinstance(value, str):
        text = value.strip().replace("\n", " ")
        return '"' + (text if len(text) <= _MAX_TEXT else text[: _MAX_TEXT - 1] + "…") + '"'
    return str(value)


def _rows_section(sheet: SheetModel) -> list[str]:
    lines = ["ROWS (text shown everywhere; numbers only near the top; [n num] = count of numeric cells):"]
    shown_numeric = 0
    for row, cells in sheet.rows():
        texts = [c for c in cells if isinstance(c.value, str) and c.value.strip()]
        others = [c for c in cells if c not in texts and c.value is not None]
        show_numbers = shown_numeric < _NUMERIC_ROWS
        shown_numeric += 1
        parts = [f"{c.column}={_fmt_value(c.value)}" for c in (cells if show_numbers else texts) if c.value is not None]
        if not show_numbers and others:
            parts.append(f"[{len(others)} num]")
        if parts:
            hidden = " (hidden)" if row in sheet.hidden_rows else ""
            lines.append(f"r{row}{hidden}: " + " ".join(parts))
    return lines


def _formulas_section(sheet: SheetModel) -> list[str]:
    formula_cells = [c for c in sheet.cells.values() if c.formula]
    if not formula_cells:
        return []
    lines = ["FORMULAS ({r} = same row, {r-2} = two rows up, {c} = same column):"]
    by_col: dict[str, list[CellModel]] = defaultdict(list)
    for cell in formula_cells:
        by_col[cell.column].append(cell)
    singletons: list[CellModel] = []
    for column in sorted(by_col, key=column_index_from_string):
        cells = sorted(by_col[column], key=lambda c: c.row)
        groups: dict[str, list[CellModel]] = defaultdict(list)
        for cell in cells:
            groups[_normalize(cell.formula, cell.row, cell.column_index)].append(cell)
        main_pattern, main_cells = max(groups.items(), key=lambda kv: len(kv[1]))
        if len(main_cells) < 3:
            singletons.extend(cells)
            continue
        exceptions: list[CellModel] = []
        for pattern, group in groups.items():
            if pattern == main_pattern:
                continue
            if len(group) >= 3:
                lines.append(f"{column} rows {_ranges([c.row for c in group])}: {pattern}")
            elif _is_near(pattern, main_pattern):
                exceptions.extend(group)
            else:
                singletons.extend(group)
        line = f"{column} rows {_ranges([c.row for c in main_cells])}: {main_pattern}"
        if exceptions:
            shown = ", ".join(f"{c.coord} {c.formula}" for c in exceptions[:_MAX_EXCEPTIONS])
            more = f" (+{len(exceptions) - _MAX_EXCEPTIONS} more)" if len(exceptions) > _MAX_EXCEPTIONS else ""
            line += f" | exceptions: {shown}{more}"
        lines.append(line)
    # Row-wise grouping of the remaining cells (subtotal rows, check rows, drivers).
    by_row: dict[tuple[int, str], list[CellModel]] = defaultdict(list)
    for cell in singletons:
        by_row[(cell.row, _normalize(cell.formula, cell.row, cell.column_index, relative_cols=True))].append(cell)
    for (row, pattern), group in sorted(by_row.items(), key=lambda kv: (kv[0][0], min(c.column_index for c in kv[1]))):
        if len(group) >= 2:
            lines.append(f"row {row} cols {_col_ranges([c.column_index for c in group])}: {pattern}")
        else:
            lines.append(f"{group[0].coord}: {group[0].formula}")
    return lines


def _is_near(pattern: str, main: str) -> bool:
    """A deviation is an 'exception' when it is the main pattern plus a trailing literal."""
    return pattern.startswith(main) and re.fullmatch(r"[+-]\d+(\.\d+)?", pattern[len(main):]) is not None


def build_skeleton(sheet: SheetModel) -> str:
    header = [
        f'SHEET "{sheet.name}" (index {sheet.index}, {sheet.state}, view={sheet.view}) used range {sheet.dimensions}',
        f"hidden columns: {', '.join(sorted(sheet.hidden_cols, key=column_index_from_string)) or '-'}"
        f" | hidden rows: {_ranges(list(sheet.hidden_rows)) if sheet.hidden_rows else '-'}",
        f"merged ranges: {', '.join(sheet.merged_ranges) or '-'}",
    ]
    return "\n".join(header + _rows_section(sheet) + _formulas_section(sheet)) + "\n"
