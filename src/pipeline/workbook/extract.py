"""Read typed capital-event data out of a workbook using validated layouts.

Everything downstream (deterministic checks, hybrid-rule facts) works on these
dataclasses, never on raw coordinates. A role whose layout is missing (not relevant
to the event, or rejected by the validator) is ``None`` here, and the checks that need
it return needs_review.
"""
from __future__ import annotations

import datetime as dt
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from openpyxl.utils.cell import range_boundaries
from pydantic import BaseModel

from src.pipeline.workbook.cells import (
    col_idx,
    col_letter,
    is_text,
    local_cell_refs,
    norm_text,
    rows_between,
    sheet_cell_refs,
    sheet_column_refs,
    to_date,
    to_decimal,
    to_money,
)
from src.pipeline.workbook.layout import (
    AllocationLayout,
    BlockComponent,
    ComponentColumn,
    EventBlock,
    HolidayCalendarLayout,
    InvestorDataLayout,
    ItdLayout,
    MergeLayout,
    MgmtFeeLayout,
    SummaryLayout,
)
from src.pipeline.workbook.loader import CellModel, SheetModel, WorkbookModel

logger = logging.getLogger("petra.pipeline")

# --- Allocation --------------------------------------------------------------------------


@dataclass
class EventInfo:
    event_type: str | None
    label: str | None
    notice_date: dt.date | None
    due_date: dt.date | None
    notice_cell: CellModel | None
    due_cell: CellModel | None
    carried_interest_rate: Decimal | None


@dataclass
class AllocationInvestor:
    name: str
    row: int
    vehicle: str
    is_gp: bool
    affiliate: bool | None
    commitment: Decimal
    commitment_pct: Decimal | None
    amounts: dict[str, Decimal]  # every component column -> amount
    formulas: dict[str, str | None]  # every component column -> stored formula
    call_total: Decimal | None
    dist_total: Decimal | None
    roll_forward: dict[str, Decimal | None]
    distribution_basis: Decimal | None = None
    received: Any = None
    received_date: Any = None
    tax_withholding: Decimal | None = None
    cells: dict[str, CellModel | None] = field(default_factory=dict)  # component column -> cell


@dataclass
class AllocationVehicle:
    name: str
    driver_row: int
    driver: dict[str, Decimal]
    totals: dict[str, Decimal]  # vehicle total row, per column
    lp_subtotals: dict[str, Decimal]
    gp_subtotals: dict[str, Decimal]
    investors: list[AllocationInvestor]
    subtotal_rows: dict[str, int | None]
    # False for a look-through block whose driver re-allocates amounts already allocated in
    # the other blocks (e.g. the GP entity's own partners), so it must not be added to them.
    additive: bool = True

    @property
    def limited_partners(self) -> list[AllocationInvestor]:
        return [i for i in self.investors if not i.is_gp]

    @property
    def rows(self) -> set[int]:
        return {i.row for i in self.investors}


@dataclass
class AllocationData:
    sheet: SheetModel
    layout: AllocationLayout
    event: EventInfo
    components: list[ComponentColumn]
    vehicles: list[AllocationVehicle]
    fund_drivers: dict[str, Decimal]  # fund-level amount per column (see ``fund_driver_row_shared``)
    grand_totals: dict[str, Decimal]
    # True when the layout's fund_driver_row is one vehicle's own driver row (a multi-vehicle
    # sheet with no separate fund-level row): ``fund_drivers`` is then the sum of the additive
    # vehicles' drivers rather than that row's values.
    fund_driver_row_shared: bool = False

    @property
    def additive_vehicles(self) -> list[AllocationVehicle]:
        return [v for v in self.vehicles if v.additive]

    @property
    def fund_investors(self) -> list[AllocationInvestor]:
        """Investors whose amounts add up to the fund: every investor of the additive vehicles."""
        return [i for v in self.additive_vehicles for i in v.investors]

    def vehicle_of(self, inv: AllocationInvestor) -> AllocationVehicle | None:
        return next((v for v in self.vehicles if inv in v.investors), None)

    @property
    def active_components(self) -> list[ComponentColumn]:
        """Components used in this event: the mapper's ``active`` flag, else a non-zero driver. A column
        flagged active that carries no driver and no amount on any investor row is not active, whatever
        the flag says (a stale header on an unused column)."""
        out = []
        for comp in self.components:
            active = comp.active if comp.active is not None else bool(self.fund_drivers.get(comp.column))
            if active and not self.fund_drivers.get(comp.column) \
                    and not any(inv.amounts.get(comp.column) for v in self.vehicles for inv in v.investors):
                active = False
            if active:
                out.append(comp)
        return out

    @property
    def inactive_components(self) -> list[ComponentColumn]:
        active = {c.column for c in self.active_components}
        return [c for c in self.components if c.column not in active]

    @property
    def event_gross(self) -> Decimal:
        return sum((self.fund_drivers.get(c.column, Decimal("0")) for c in self.active_components), Decimal("0"))

    @property
    def investors(self) -> list[AllocationInvestor]:
        return [i for v in self.vehicles for i in v.investors]

    def event_total_column(self, side: str) -> str | None:
        return next((t.column for t in self.layout.event_total_columns if t.side == side), None)


def _row_values(sheet: SheetModel, row: int | None, columns: list[str]) -> dict[str, Decimal]:
    if row is None:
        return {}
    return {col: to_money(sheet.value(f"{col}{row}")) for col in columns}


def roll_forward_columns(layout: AllocationLayout) -> dict[str, str]:
    """name -> column of every mapped roll-forward column; adjustment columns are keyed 'adjustment:<col>'."""
    out = {name: col for name, col in layout.roll_forward.model_dump().items() if isinstance(col, str) and col}
    for column in layout.roll_forward.adjustments:
        out[f"adjustment:{column}"] = column
    return out


def _numeric_columns(layout: AllocationLayout) -> list[str]:
    cols = [layout.columns.commitment, layout.columns.commitment_pct]
    cols += [c.column for c in layout.components] + [t.column for t in layout.event_total_columns]
    cols += list(roll_forward_columns(layout).values())
    for extra in (layout.columns.cash_due, layout.columns.late_interest, layout.columns.distribution_basis,
                  layout.columns.distribution_basis_pct, layout.columns.tax_withholding):
        if extra:
            cols.append(extra)
    return list(dict.fromkeys(cols))


def _affiliate(value: Any) -> bool | None:
    text = norm_text(value)
    if text in ("y", "yes", "affiliate", "true"):
        return True
    if text in ("n", "no", "false"):
        return False
    return None


def extract_allocation(model: WorkbookModel, layout: AllocationLayout) -> AllocationData:
    sheet = model.sheet(layout.sheet)
    cols = layout.columns
    numeric = _numeric_columns(layout)
    comp_cols = [c.column for c in layout.components]
    fund_drivers = _row_values(sheet, layout.fund_driver_row, comp_cols + [t.column for t in layout.event_total_columns])

    rate_cell = layout.event.carried_interest_rate_cell
    notice_cell = sheet.cell(layout.event.notice_date_cell) if layout.event.notice_date_cell else None
    due_cell = sheet.cell(layout.event.due_date_cell) if layout.event.due_date_cell else None
    label = layout.event.label
    if layout.event.label_cell and is_text(sheet.value(layout.event.label_cell)):
        label = str(sheet.value(layout.event.label_cell)).strip()
    event = EventInfo(
        event_type=layout.event.event_type,
        label=label,
        notice_date=to_date(notice_cell.value, allow_serial=True) if notice_cell else None,
        due_date=to_date(due_cell.value, allow_serial=True) if due_cell else None,
        notice_cell=notice_cell,
        due_cell=due_cell,
        carried_interest_rate=to_decimal(sheet.value(rate_cell)) if rate_cell else None,
    )

    vehicles: list[AllocationVehicle] = []
    for vrows in layout.vehicles:
        gp_rows = set(vrows.gp_rows)
        rows = [r for r in rows_between(vrows.investor_rows) if is_text(sheet.value(f"{cols.investor}{r}"))]
        rows += sorted(gp_rows)
        investors = []
        for row in rows:
            rf = {name: to_decimal(sheet.value(f"{col}{row}")) for name, col in roll_forward_columns(layout).items()}
            for name in ("commitment", "prior_contributions", "prior_recallable", "current_call",
                         "current_recallable", "remaining_commitment"):
                rf.setdefault(name, None)
            call_col = layout_total(layout, "call")
            dist_col = layout_total(layout, "distribution")
            investors.append(AllocationInvestor(
                name=str(sheet.value(f"{cols.investor}{row}")).strip(),
                row=row,
                vehicle=vrows.name,
                is_gp=row in gp_rows,
                affiliate=_affiliate(sheet.value(f"{cols.affiliate_flag}{row}")) if cols.affiliate_flag else None,
                commitment=to_money(sheet.value(f"{cols.commitment}{row}")),
                commitment_pct=to_decimal(sheet.value(f"{cols.commitment_pct}{row}")),
                amounts={c: to_money(sheet.value(f"{c}{row}")) for c in comp_cols},
                formulas={c: (sheet.cell(f"{c}{row}").formula if sheet.cell(f"{c}{row}") else None) for c in comp_cols},
                call_total=to_decimal(sheet.value(f"{call_col}{row}")) if call_col else None,
                dist_total=to_decimal(sheet.value(f"{dist_col}{row}")) if dist_col else None,
                roll_forward=rf,
                distribution_basis=to_decimal(sheet.value(f"{cols.distribution_basis}{row}"))
                if cols.distribution_basis else None,
                tax_withholding=to_decimal(sheet.value(f"{cols.tax_withholding}{row}"))
                if cols.tax_withholding else None,
                received=sheet.value(f"{cols.received}{row}") if cols.received else None,
                received_date=sheet.value(f"{cols.received_date}{row}") if cols.received_date else None,
                cells={c: sheet.cell(f"{c}{row}") for c in comp_cols},
            ))
        driver_row = vrows.driver_row or layout.fund_driver_row
        lp_subtotals = _row_values(sheet, vrows.subtotal_rows.limited_partners, numeric)
        gp_subtotals = _row_values(sheet, vrows.subtotal_rows.general_partner, numeric)
        totals = _row_values(sheet, vrows.subtotal_rows.total, numeric)
        if not totals:
            # A block without a "Total" row (e.g. LPs only): its total is the subtotal(s) it has, else
            # the sum of its rows, so vehicle-level ties still have a figure to compare.
            if lp_subtotals or gp_subtotals:
                totals = {col: lp_subtotals.get(col, Decimal("0")) + gp_subtotals.get(col, Decimal("0")) for col in numeric}
            else:
                totals = {col: sum((to_money(sheet.value(f"{col}{i.row}")) for i in investors), Decimal("0"))
                          for col in numeric}
        vehicles.append(AllocationVehicle(
            name=vrows.name,
            driver_row=driver_row,
            driver=_row_values(sheet, driver_row, comp_cols + [t.column for t in layout.event_total_columns]),
            totals=totals,
            lp_subtotals=lp_subtotals,
            gp_subtotals=gp_subtotals,
            investors=investors,
            subtotal_rows=vrows.subtotal_rows.model_dump(),
        ))
    grand_row = layout.grand_total_row or (layout.vehicles[-1].subtotal_rows.total if len(layout.vehicles) == 1 else None)
    _flag_non_additive(sheet, layout, vehicles, comp_cols, grand_row)
    shared = len(vehicles) > 1 and any(v.driver_row == layout.fund_driver_row for v in vehicles)
    if shared:
        columns = comp_cols + [t.column for t in layout.event_total_columns]
        fund_drivers = {col: sum((v.driver.get(col, Decimal("0")) for v in vehicles if v.additive), Decimal("0"))
                        for col in columns}
    return AllocationData(
        sheet=sheet,
        layout=layout,
        event=event,
        components=list(layout.components),
        vehicles=vehicles,
        fund_drivers=fund_drivers,
        grand_totals=_row_values(sheet, grand_row, numeric),
        fund_driver_row_shared=shared,
    )


def _flag_non_additive(sheet: SheetModel, layout: AllocationLayout, vehicles: list[AllocationVehicle],
                       comp_cols: list[str], grand_row: int | None) -> None:
    """Mark look-through vehicle blocks from the sheet's own formulas.

    Two signals, either of which marks a block non-additive: the grand-total row sums the
    other blocks' total rows but not this one (``=Q65+Q95+Q134``), or this block's driver row
    is built from the other blocks' investor rows (``=Q61+Q91+Q130``, the GP entity's share of
    each fund re-allocated to the GP's own partners).
    """
    if len(vehicles) < 2:
        return
    total_rows = {v.subtotal_rows.get("total"): v for v in vehicles if v.subtotal_rows.get("total")}
    if grand_row and grand_row not in total_rows:
        referenced: set[int] = set()
        ranged = False
        for col in [layout.columns.commitment] + comp_cols:
            cell = sheet.cell(f"{col}{grand_row}")
            if cell is None or not cell.formula:
                continue
            ranged = ranged or ":" in cell.formula
            referenced |= {row for _, row in local_cell_refs(cell.formula)}
        if not ranged and len(referenced & set(total_rows)) >= 2:
            for row, vehicle in total_rows.items():
                if row not in referenced:
                    vehicle.additive = False
    for vehicle in vehicles:
        if vehicle.driver_row == layout.fund_driver_row:
            continue
        others = {row for other in vehicles if other is not vehicle for row in other.rows}
        for col in comp_cols:
            cell = sheet.cell(f"{col}{vehicle.driver_row}")
            if cell is None or not cell.formula or sheet_cell_refs(cell.formula):
                continue
            if any(row in others for _, row in local_cell_refs(cell.formula)):
                vehicle.additive = False
                break


def layout_total(layout: AllocationLayout, side: str) -> str | None:
    return next((t.column for t in layout.event_total_columns if t.side == side), None)


# --- ITD ------------------------------------------------------------------------------------


@dataclass
class ItdInvestor:
    name: str
    row: int
    vehicle: str
    is_gp: bool
    cumulative: dict[str, Decimal | None]
    values: dict[str, Decimal]  # block component (and total) column -> amount
    formulas: dict[str, str | None]


@dataclass
class ItdVehicle:
    name: str
    investors: list[ItdInvestor]
    subtotal_rows: dict[str, int | None]
    totals: dict[str, Decimal]


@dataclass
class ItdData:
    sheet: SheetModel
    layout: ItdLayout
    event_blocks: list[EventBlock]
    vehicles: list[ItdVehicle]
    marks: dict[str, list[str]]  # block column -> classification categories marked 'X'
    overlay_marks: dict[str, list[str]]
    # cumulative key -> 'accumulator' (sums the marked event columns), 'derived' (computed from
    # other figures, e.g. a recycling cap) or 'value' (typed numbers)
    cumulative_kinds: dict[str, str] = field(default_factory=dict)
    # Notes from the deterministic block enumeration (e.g. a block whose side could not be read),
    # rows inside the investor span that were not treated as investors, and check columns.
    notes: list[str] = field(default_factory=list)
    excluded_rows: list[tuple[int, str]] = field(default_factory=list)
    check_columns: list[str] = field(default_factory=list)

    @property
    def derived_cumulatives(self) -> set[str]:
        return {key for key, kind in self.cumulative_kinds.items() if kind == "derived"}

    @property
    def current_block(self) -> EventBlock | None:
        current = [b for b in self.event_blocks if b.is_current]
        return current[0] if len(current) == 1 else None

    @property
    def prior_blocks(self) -> list[EventBlock]:
        return [b for b in self.event_blocks if not b.is_current]

    @property
    def investors(self) -> list[ItdInvestor]:
        return [i for v in self.vehicles for i in v.investors]

    def block_columns(self, block: EventBlock) -> list[str]:
        return [c.column for c in block.components]


def _is_mark(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() == "X"


# --- ITD event-block enumeration ------------------------------------------------------------
#
# The event header row is read in code: every text label right of the cumulative columns starts
# a block, whose span comes from the label's merged range (else it runs to the column before the
# next label). The LLM's blocks refine what the enumeration finds (event type, number, date,
# the current block, component types); they are never the only source of blocks.

_BLOCK_QUALIFIER_RE = re.compile(r"\bnon[\s-]*recallable\b|\brecallable\b", re.I)
_BLOCK_TYPE_WORDS: list[tuple[re.Pattern, str]] = [
    (re.compile(r"\btransfer"), "other"),
    (re.compile(r"\b(gains?|realized|realised)\b"), "realized_gain"),
    (re.compile(r"\b(dividends?|interest income|income)\b"), "dividend_income"),
    (re.compile(r"\b(mgmt|management)\b"), "mgmt_fee"),
    (re.compile(r"\bplacement\b"), "placement_fee"),
    (re.compile(r"\blate interest\b"), "late_interest"),
    (re.compile(r"\bcarr(y|ied)\b"), "carry"),
    (re.compile(r"\btax distribution"), "tax_distribution"),
    (re.compile(r"\b(withholding|tax wh|wh|tax)\b"), "tax_withholding"),
    (re.compile(r"\b(pref|preferred)\b"), "pref"),
    (re.compile(r"\bcatch"), "catch_up"),
    (re.compile(r"\b(roc|return of capital)\b"), "return_of_capital"),
    (re.compile(r"\b(expenses?|org|organizational|organisational|partnership|costs?)\b"), "org_expense"),
    (re.compile(r"\b(investments?|call|deemed|working capital|follow[\s-]*on|contribution)\b"), "investment"),
]
_CALL_TYPES = {"investment", "org_expense", "mgmt_fee", "placement_fee", "late_interest"}
_CALL_CATEGORIES = {"investment_contributions", "cost_contributions"}
_DIST_CATEGORIES = {"recallable_distributions", "non_recallable_distributions", "tax_withholding"}
_LABEL_DATE_RE = re.compile(r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})")


def component_type_from_header(header: Any) -> str:
    """Component type read from an ITD sub-header ('Vagaro A-2 Non-Recallable Realized Gain' ->
    realized_gain, 'Egress (Escrow) Recallable' -> return_of_capital); 'other' when nothing matches."""
    text = norm_text(header)
    stripped = _BLOCK_QUALIFIER_RE.sub(" ", text)
    for pattern, kind in _BLOCK_TYPE_WORDS:
        if pattern.search(stripped):
            return kind
    if text != stripped.strip() or "roc" in text:
        return "return_of_capital"
    return "other"


def _label_event_type(label: str) -> str:
    text = norm_text(label)
    if "transfer" in text:
        return "transfer"
    call = re.search(r"\b(call|contribution)", text) is not None
    dist = re.search(r"\bdist", text) is not None
    if call and dist:
        return "net_event"
    if call:
        return "capital_call"
    if dist:
        return "distribution"
    return "other"


def _label_date(label: str) -> str | None:
    match = _LABEL_DATE_RE.search(label)
    if not match:
        return None
    day = to_date(match.group(1).replace("-", ".").replace("/", "."))
    return day.isoformat() if day else None


def _event_number(label: str) -> int | None:
    match = re.search(r"#\s*(\d+)", label)
    return int(match.group(1)) if match else None


def _header_merges(sheet: SheetModel, header_row: int) -> dict[int, int]:
    """first column index -> last column index of every merged range on the header row."""
    spans: dict[int, int] = {}
    for rng in sheet.merged_ranges:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(rng)
        except ValueError:
            continue
        if min_row <= header_row <= max_row and max_col > min_col:
            spans[min_col] = max_col
    return spans


def itd_header_labels(sheet: SheetModel, layout: ItdLayout) -> list[tuple[int, str]]:
    """(column index, label) of every event block label on the event header row, left to right."""
    cum = [c for c in layout.cumulative_columns.model_dump().values() if c]
    start = max([col_idx(layout.investor_column)] + [col_idx(c) for c in cum])
    spans = _header_merges(sheet, layout.event_header_row)
    labels: list[tuple[int, str]] = []
    covered_until = 0
    for cell in sheet.row_cells(layout.event_header_row):
        if cell.column_index <= start or not is_text(cell.value) or cell.is_error:
            continue
        if cell.column_index <= covered_until:
            continue  # text inside another label's merged range
        labels.append((cell.column_index, str(cell.value).strip()))
        covered_until = spans.get(cell.column_index, cell.column_index)
    return labels


def _investor_span_rows(layout: ItdLayout) -> list[int]:
    rows: list[int] = []
    for vehicle in layout.vehicles:
        if len(vehicle.investor_rows) == 2:
            rows += list(rows_between(vehicle.investor_rows))
        rows += list(vehicle.gp_rows)
    return rows


def _has_values(sheet: SheetModel, column: str, rows: list[int]) -> bool:
    return any(to_decimal(sheet.value(f"{column}{r}")) for r in rows)


def _links_other_sheet(sheet: SheetModel, column: str, rows: list[int], sheets: list[str] | None) -> bool:
    for row in rows:
        cell = sheet.cell(f"{column}{row}")
        if cell is None or not cell.formula or "!" not in cell.formula:
            continue
        if not sheets:
            return True
        refs = {s for s, _, _ in sheet_cell_refs(cell.formula)} | {s for s, _ in sheet_column_refs(cell.formula)}
        if refs & set(sheets):
            return True
    return False


def enumerate_itd_blocks(sheet: SheetModel, layout: ItdLayout,
                         allocation_sheets: list[str] | None = None) -> tuple[list[EventBlock], list[str]]:
    """Event blocks read from the sheet's structure, refined by the LLM's ``layout.event_blocks``.

    Returns the blocks and notes about what could not be read (a block whose side had to be
    assumed from the sign of its values). When the header row carries no labels at all, the
    LLM's blocks are returned unchanged.
    """
    notes: list[str] = []
    labels = itd_header_labels(sheet, layout)
    if not labels:
        notes.append("no event labels were found on the event header row; the mapped blocks were used as given")
        return list(layout.event_blocks), notes
    spans = _header_merges(sheet, layout.event_header_row)
    rows = _investor_span_rows(layout)
    sub_row = layout.subheader_row
    class_rows = {cat: row for cat, row in layout.classification_rows.model_dump().items() if row}
    mapped_by_col = {col_idx(b.first_column): b for b in layout.event_blocks}
    mapped_by_label = {b.label: b for b in layout.event_blocks}
    max_col = max(sheet.max_column, max(c.column_index for c in sheet.cells.values()) if sheet.cells else 1)
    blocks: list[EventBlock] = []
    for position, (first, label) in enumerate(labels):
        next_first = labels[position + 1][0] if position + 1 < len(labels) else max_col + 1
        if first in spans:
            last = min(spans[first], next_first - 1)
        else:
            last = next_first - 1
            while last > first and not is_text(sheet.value(f"{col_letter(last)}{sub_row}")) \
                    and not _has_values(sheet, col_letter(last), rows):
                last -= 1
        mapped = mapped_by_col.get(first) or mapped_by_label.get(label) or next(
            (b for b in layout.event_blocks if first < col_idx(b.first_column) <= last), None)
        mapped_components = {c.column: c for c in mapped.components} if mapped else {}
        total_column = None
        components: list[BlockComponent] = []
        for index in range(first, last + 1):
            column = col_letter(index)
            header = sheet.value(f"{column}{sub_row}")
            has_header = is_text(header)
            if has_header and norm_text(header).startswith("total") and total_column is None:
                total_column = column
                continue
            if not has_header and not _has_values(sheet, column, rows):
                continue
            marks = {cat for cat, row in class_rows.items() if _is_mark(sheet.value(f"{column}{row}"))}
            given = mapped_components.get(column)
            kind = given.component_type if given else component_type_from_header(header if has_header else "")
            if marks & _CALL_CATEGORIES and not marks & _DIST_CATEGORIES:
                side = "call"
            elif marks & _DIST_CATEGORIES and not marks & _CALL_CATEGORIES:
                side = "distribution"
            elif given is not None:
                side = given.side
            elif kind != "other":
                side = "call" if kind in _CALL_TYPES else "distribution"
            else:
                values = [to_decimal(sheet.value(f"{column}{r}")) for r in rows]
                negatives = sum(1 for v in values if v is not None and v < 0)
                positives = sum(1 for v in values if v is not None and v > 0)
                side = "distribution" if negatives > positives else "call"
                notes.append(f"block '{label}': the side of column {column} ({header if has_header else 'no sub-header'})"
                             f" could not be read from its X marks or sub-header; assumed {side} from the sign of "
                             "its values")
            components.append(BlockComponent(column=column, component_type=kind, side=side))
        if not components and total_column is None:
            continue  # a stray label over empty columns
        last_component = components[-1].column if components else total_column
        if mapped is not None:
            event_type, number, date, is_current = (mapped.event_type, mapped.number or _event_number(label),
                                                   mapped.date or _label_date(label), mapped.is_current)
        else:
            event_type, number, date, is_current = (_label_event_type(label), _event_number(label),
                                                   _label_date(label), False)
        blocks.append(EventBlock(label=label, event_type=event_type, number=number, date=date,
                                 first_column=col_letter(first), last_column=last_component,
                                 total_column=total_column, is_current=is_current, components=components))
    if not blocks:
        notes.append("the event header labels did not yield any block; the mapped blocks were used as given")
        return list(layout.event_blocks), notes
    current = [b for b in blocks if b.is_current]
    if len(current) != 1:
        for block in blocks:
            block.is_current = False
        linked = [b for b in blocks
                  if any(_links_other_sheet(sheet, c.column, rows, allocation_sheets) for c in b.components)]
        if linked:
            linked[-1].is_current = True
            if not current:
                notes.append(f"no mapped block is marked current; '{linked[-1].label}' was taken as the current "
                             "block because its cells link to the Allocation sheet")
        else:
            notes.append("no block links to the Allocation sheet, so the current block could not be identified")
    return blocks, notes


def extract_itd(model: WorkbookModel, layout: ItdLayout, allocation_sheets: list[str] | None = None) -> ItdData:
    sheet = model.sheet(layout.sheet)
    event_blocks, notes = enumerate_itd_blocks(sheet, layout, allocation_sheets)
    block_cols = [c.column for b in event_blocks for c in b.components]
    totals_cols = [b.total_column for b in event_blocks if b.total_column]
    marks: dict[str, list[str]] = {}
    for column in block_cols:
        marks[column] = [cat for cat, row in layout.classification_rows.model_dump().items()
                         if row is not None and _is_mark(sheet.value(f"{column}{row}"))]
    overlay_marks = {column: [o.name for o in layout.overlay_rows if _is_mark(sheet.value(f"{column}{o.row}"))]
                     for column in block_cols}
    cum_cols = layout.cumulative_columns.model_dump()
    all_cols = block_cols + totals_cols
    value_cols = all_cols + [c for c in cum_cols.values() if c]
    commitment_col = cum_cols.get("commitment")
    excluded: list[tuple[int, str]] = []

    def is_investor_row(row: int) -> bool:
        """A named row with a numeric commitment, or (a transferred-out investor) with any amount."""
        name = sheet.value(f"{layout.investor_column}{row}")
        if not is_text(name):
            return False
        if commitment_col:
            commitment = sheet.value(f"{commitment_col}{row}")
            if isinstance(commitment, (int, float)) and not isinstance(commitment, bool):
                return True
            if to_decimal(commitment) is not None:
                return True
        else:
            return True
        if any(to_decimal(sheet.value(f"{c}{row}")) for c in value_cols):
            return True
        excluded.append((row, str(name).strip()))
        return False

    vehicles = []
    for vrows in layout.vehicles:
        gp_rows = set(vrows.gp_rows)
        rows = [r for r in rows_between(vrows.investor_rows) if is_investor_row(r)]
        rows += sorted(gp_rows)
        investors = []
        for row in rows:
            investors.append(ItdInvestor(
                name=str(sheet.value(f"{layout.investor_column}{row}")).strip(),
                row=row,
                vehicle=vrows.name,
                is_gp=row in gp_rows,
                cumulative={k: (to_decimal(sheet.value(f"{c}{row}")) if c else None) for k, c in cum_cols.items()},
                values={c: to_money(sheet.value(f"{c}{row}")) for c in all_cols},
                formulas={c: (sheet.cell(f"{c}{row}").formula if sheet.cell(f"{c}{row}") else None) for c in all_cols},
            ))
        vehicles.append(ItdVehicle(
            name=vrows.name,
            investors=investors,
            subtotal_rows=vrows.subtotal_rows.model_dump(),
            totals=_row_values(sheet, vrows.subtotal_rows.total, all_cols + [c for c in cum_cols.values() if c]),
        ))
    investor_rows = [inv.row for v in vehicles for inv in v.investors if not inv.is_gp]
    kinds = _cumulative_kinds(sheet, cum_cols, set(block_cols) | set(totals_cols), investor_rows)
    if excluded:
        logger.info("ITD %s: %d named row(s) inside the investor span carry no commitment or amount and were not "
                    "treated as investors: %s", sheet.name,
                    len(excluded), ", ".join(f"{r} {n[:40]}" for r, n in excluded[:8]))
    check_columns = list(layout.check_columns)
    block_set = set(all_cols)
    for cell in sheet.row_cells(layout.subheader_row):
        if cell.column in check_columns or cell.column in block_set:
            continue
        if is_text(cell.value) and re.search(r"\b(check|difference|variance)\b", norm_text(cell.value)):
            check_columns.append(cell.column)
    return ItdData(sheet=sheet, layout=layout, event_blocks=event_blocks, vehicles=vehicles,
                   marks=marks, overlay_marks=overlay_marks, cumulative_kinds=kinds, notes=notes,
                   excluded_rows=excluded, check_columns=check_columns)


_CUMULATIVE_ORDER = ("investment_contributions", "cost_contributions", "recallable_distributions",
                     "non_recallable_distributions", "total_contributions", "total_distributions", "commitment",
                     "unfunded")


def _dominant_formula(sheet: SheetModel, column: str, rows: list[int]) -> tuple[str | None, int]:
    """The most common formula shape in a column over ``rows`` (row numbers normalised), with its count."""
    shapes: dict[str, tuple[str, int]] = {}
    for row in rows:
        cell = sheet.cell(f"{column}{row}")
        formula = cell.formula if cell is not None else None
        shape = re.sub(r"(?<=[A-Z])\$?" + str(row) + r"(?!\d)", "{r}", formula) if formula else ""
        # Bare row references such as SUMIF($2:$2,"X",12:12) address the whole row.
        shape = re.sub(r"(?<![A-Za-z0-9])\$?" + str(row) + r":\$?" + str(row) + r"(?!\d)", "{r}:{r}", shape)
        count = shapes.get(shape, (formula, 0))[1] + 1
        shapes[shape] = (formula, count)
    if not shapes:
        return None, 0
    formula, count = max(shapes.values(), key=lambda fc: fc[1])
    return formula, count


def _strip_calls(text: str, names: tuple[str, ...]) -> str:
    """Remove every ``NAME(...)`` call (balanced parentheses) from an upper-cased formula."""
    out = text
    for name in names:
        while True:
            start = out.find(name)
            if start < 0:
                break
            depth, index = 0, start + len(name) - 1
            while index < len(out):
                if out[index] == "(":
                    depth += 1
                elif out[index] == ")":
                    depth -= 1
                    if depth == 0:
                        break
                index += 1
            out = out[:start] + " " + out[index + 1:]
    return out.replace("=", "").replace("-", "").replace("+", "").replace("(", "").replace(")", "").strip()


def _cumulative_kinds(sheet: SheetModel, cum_cols: dict[str, str | None], event_cols: set[str],
                      rows: list[int]) -> dict[str, str]:
    """Classify each cumulative column as accumulator, derived or value (see ``ItdData``)."""
    kinds: dict[str, str] = {}
    columns = {key: col for key, col in cum_cols.items() if col}
    accumulators = set(event_cols)
    for _ in range(2):  # a total of accumulators is an accumulator; resolve in two passes
        for key in _CUMULATIVE_ORDER:
            col = columns.get(key)
            if col is None or kinds.get(key) == "accumulator":
                continue
            formula, count = _dominant_formula(sheet, col, rows)
            if not formula or count * 2 < len(rows):
                kinds[key] = "value"
                continue
            upper = formula.upper()
            refs = {c for c, _ in local_cell_refs(formula)}
            remainder = _strip_calls(upper, ("SUMIFS(", "SUMIF("))
            if "SUMIF" in upper and not re.search(r"[A-Z0-9*/]", remainder):
                kinds[key] = "accumulator"  # only SUMIF terms, added or negated
            elif upper.startswith("=SUM(") and refs & event_cols and not re.search(r"[*/]", upper):
                kinds[key] = "accumulator"
            elif refs and refs <= accumulators and not re.search(r"[*/]|MIN\(|MAX\(|IF\(", upper.replace("SUM(", "(")):
                kinds[key] = "accumulator"
            else:
                kinds[key] = "derived"
            if kinds[key] == "accumulator":
                accumulators.add(col)
    return kinds


# --- Summary --------------------------------------------------------------------------------


@dataclass
class SummaryLineData:
    label: str | None
    amount: Decimal
    component_type: str
    side: str
    amount_cell: str


@dataclass
class SummarySectionData:
    vehicle: str | None
    title: str | None
    fund_commitment: Decimal | None
    lines: list[SummaryLineData]
    event_total: Decimal | None
    check_values: list[Decimal]
    check_cells: list[str]


@dataclass
class SummaryData:
    sheet: SheetModel
    layout: SummaryLayout
    title: str | None
    notice_date: dt.date | None
    due_date: dt.date | None
    fund_commitment: Decimal | None
    lines: list[SummaryLineData]
    event_total: Decimal
    check_values: list[Decimal]
    check_cells: list[str]
    sections: list[SummarySectionData] = field(default_factory=list)  # one per vehicle block, when repeated


def _summary_lines(sheet: SheetModel, lines) -> list[SummaryLineData]:
    return [
        SummaryLineData(
            label=str(sheet.value(line.label_cell)) if line.label_cell and sheet.value(line.label_cell) else None,
            amount=to_money(sheet.value(line.amount_cell)),
            component_type=line.component_type,
            side=line.side,
            amount_cell=line.amount_cell,
        )
        for line in lines
    ]


def extract_summary(model: WorkbookModel, layout: SummaryLayout) -> SummaryData:
    sheet = model.sheet(layout.sheet)
    lines = _summary_lines(sheet, layout.component_lines)
    sections = [
        SummarySectionData(
            vehicle=section.vehicle or (str(sheet.value(section.title_cell)).strip()
                                        if section.title_cell and is_text(sheet.value(section.title_cell)) else None),
            title=str(sheet.value(section.title_cell)) if section.title_cell and sheet.value(section.title_cell) else None,
            fund_commitment=to_decimal(sheet.value(section.fund_commitment_cell)) if section.fund_commitment_cell else None,
            lines=_summary_lines(sheet, section.component_lines),
            event_total=to_decimal(sheet.value(section.event_total_cell)) if section.event_total_cell else None,
            check_values=[to_money(sheet.value(cell)) for cell in section.check_cells],
            check_cells=list(section.check_cells),
        )
        for section in layout.sections
    ]
    notice_cell, due_cell = layout.notice_date_cell, layout.due_date_cell
    if notice_cell and due_cell and notice_cell.replace("$", "").upper() == due_cell.replace("$", "").upper():
        # One cell mapped as both dates: keep it as the due date when its text says so, else as neither.
        text = norm_text(sheet.value(due_cell))
        notice_cell = None
        if not re.search(r"\b(due|payable|payment|wire)\b", text):
            due_cell = None
    return SummaryData(
        sheet=sheet,
        layout=layout,
        title=str(sheet.value(layout.title_cell)) if layout.title_cell and sheet.value(layout.title_cell) else None,
        notice_date=to_date(sheet.value(notice_cell), allow_serial=True) if notice_cell else None,
        due_date=to_date(sheet.value(due_cell), allow_serial=True) if due_cell else None,
        fund_commitment=to_decimal(sheet.value(layout.fund_commitment_cell)) if layout.fund_commitment_cell else None,
        lines=lines,
        event_total=to_money(sheet.value(layout.event_total_cell)),
        check_values=[to_money(sheet.value(cell)) for cell in layout.check_cells],
        check_cells=list(layout.check_cells),
        sections=sections,
    )


# --- Merge ----------------------------------------------------------------------------------


@dataclass
class MergeRow:
    name: str
    row: int
    short_name: Any
    investor_id: Any
    fund_id: Any
    file_name: Any
    letter_date: Any
    due_date: Any
    commitment: Decimal | None
    amounts: dict[str, Decimal]
    event_total: Decimal | None
    values: dict[str, Decimal] = field(default_factory=dict)  # every linked numeric column -> amount
    source_row: int | None = None  # Allocation row this merge row reads


@dataclass
class MergeData:
    sheet: SheetModel
    layout: MergeLayout
    vehicle: str | None
    rows: list[MergeRow]
    source_sheet: str | None = None  # the sheet the data rows pull from (normally the Allocation)
    referenced_columns: dict[str, str] = field(default_factory=dict)  # merge column -> source column
    header_refs: dict[str, str] = field(default_factory=dict)  # merge column -> source header column
    headers: dict[str, str] = field(default_factory=dict)  # merge column -> header text
    totals: dict[str, Decimal] = field(default_factory=dict)  # total row, per merge column
    check_cells: list[tuple[str, Decimal]] = field(default_factory=list)  # (coord, value) that should be 0


def _id(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def extract_merge(model: WorkbookModel, layout: MergeLayout) -> MergeData:
    sheet = model.sheet(layout.sheet)
    cols = layout.columns

    def val(col: str | None, row: int) -> Any:
        return sheet.value(f"{col}{row}") if col else None

    data_rows = [r for r in range(layout.first_data_row, layout.last_data_row + 1) if is_text(val(cols.investor, r))]
    # Where the data rows pull from: the dominant sheet-qualified reference per column.
    per_column: dict[str, Counter] = {}
    source_rows: dict[int, Counter] = {}
    for row in data_rows:
        for cell in sheet.row_cells(row):
            refs = [(s_, c_, r_) for s_, c_, r_ in sheet_cell_refs(cell.formula) if s_ != sheet.name]
            for ref_sheet, _, ref_row in refs:
                source_rows.setdefault(row, Counter())[(ref_sheet, ref_row)] += 1
            # A column "pulls" a source column through a plain link (=Allocation!Q7) or a per-row lookup
            # keyed by the row's own cell (SUMIF(Allocation!$D:$D,$E4,Allocation!AE:AE)); a check formula
            # such as =Allocation!AI7-Q4 references the source but is not a pull of it.
            if len(refs) == 1 and _PURE_LINK_RE.match(cell.formula or ""):
                per_column.setdefault(cell.column, Counter())[refs[0][:2]] += 1
            else:
                pulled = pull_lookup(cell.formula, row, sheet.name)
                if pulled is not None:
                    per_column.setdefault(cell.column, Counter())[pulled] += 1
    sheet_votes = Counter(ref_sheet for counts in per_column.values() for (ref_sheet, _), n in counts.items()
                          for _ in range(n))
    source_sheet = sheet_votes.most_common(1)[0][0] if sheet_votes else None
    referenced_columns = {}
    for column, counts in per_column.items():
        (ref_sheet, ref_col), n = counts.most_common(1)[0]
        if ref_sheet == source_sheet and n * 2 > len(data_rows):
            referenced_columns[column] = ref_col
    linked_cols = sorted(set(referenced_columns) | {c.column for c in layout.component_columns},
                         key=lambda c: (len(c), c))
    rows = []
    for row in data_rows:
        name = val(cols.investor, row)
        votes = source_rows.get(row)
        source_row = next((r for (s, r), _ in votes.most_common() if s == source_sheet), None) if votes else None
        rows.append(MergeRow(
            name=str(name).strip(),
            row=row,
            short_name=val(cols.short_name, row),
            investor_id=_id(val(cols.investor_id, row)),
            fund_id=_id(val(cols.fund_id, row)),
            file_name=val(cols.file_name, row),
            letter_date=val(cols.letter_date, row),
            due_date=val(cols.due_date, row),
            commitment=to_decimal(val(cols.commitment, row)),
            amounts={c.column: to_money(val(c.column, row)) for c in layout.component_columns},
            event_total=to_decimal(val(cols.event_total, row)),
            values={c: to_money(val(c, row)) for c in linked_cols},
            source_row=source_row,
        ))
    header_refs, headers = {}, {}
    for cell in sheet.row_cells(layout.header_row):
        if is_text(cell.value):
            headers[cell.column] = str(cell.value).strip()
        if _PURE_LINK_RE.match(cell.formula or ""):
            for ref_sheet, ref_col, _ in sheet_cell_refs(cell.formula):
                if ref_sheet == source_sheet:
                    header_refs[cell.column] = ref_col
    totals = {}
    if layout.total_row:
        for col in linked_cols:
            cell = sheet.cell(f"{col}{layout.total_row}")
            if cell is not None and cell.formula and not is_text(cell.value):
                totals[col] = to_money(cell.value)
    check_cells: list[tuple[str, Decimal]] = []
    if cols.check:
        for row in data_rows:
            cell = sheet.cell(f"{cols.check}{row}")
            if cell is not None and not is_text(cell.value):
                check_cells.append((cell.coord, to_money(cell.value)))
    check_rows = list(layout.check_rows) or _merge_check_rows(sheet, layout, source_sheet)
    for row in check_rows:
        for cell in sheet.row_cells(row):
            if cell.formula and not is_text(cell.value) and cell.value is not None:
                check_cells.append((cell.coord, to_money(cell.value)))
    return MergeData(sheet=sheet, layout=layout, vehicle=layout.vehicle, rows=rows, source_sheet=source_sheet,
                     referenced_columns=referenced_columns, header_refs=header_refs, headers=headers, totals=totals,
                     check_cells=check_cells)


_PURE_LINK_RE = re.compile(r"^=\s*[+]?(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!\$?[A-Z]{1,3}\$?\d+\s*$")


def pull_lookup(formula: str | None, row: int, home: str) -> tuple[str, str] | None:
    """(source sheet, source column) when ``formula`` is one per-row lookup (SUMIF / SUMIFS / INDEX-MATCH /
    VLOOKUP) into another sheet keyed by a cell of the formula's own row, and nothing else."""
    if not formula or "!" not in formula:
        return None
    from src.pipeline.workbook.checks.support import parse_lookup, split_terms

    body = formula[1:] if formula.startswith("=") else formula
    if len(split_terms(formula)) != 1 or re.search(r"[*/]|\)\s*[-+]", body):
        return None
    lookup = parse_lookup(formula, row, home)
    if lookup is None or lookup.sheet == home or not any(kind == "row" for _, kind, _ in lookup.keys):
        return None
    return lookup.sheet, lookup.value_column


def _merge_check_rows(sheet: SheetModel, layout: MergeLayout, source_sheet: str | None) -> list[int]:
    """Rows below the data that compare the tab with its source: a formula subtracting a
    source-sheet cell from a cell of this sheet (``=Allocation!X95-I22``)."""
    if source_sheet is None:
        return []
    start = (layout.total_row or layout.last_data_row) + 1
    found = []
    for row in range(start, min(sheet.max_row, start + 6) + 1):
        for cell in sheet.row_cells(row):
            formula = cell.formula or ""
            if "-" in formula and any(s == source_sheet for s, _, _ in sheet_cell_refs(formula)) \
                    and local_cell_refs(formula):
                found.append(row)
                break
    return found


# --- Mgmt fee -------------------------------------------------------------------------------


@dataclass
class FeeRow:
    name: str
    row: int
    vehicle: Any
    affiliate_flag: Any
    commitment: Decimal | None
    fees: dict[str, Decimal]
    formulas: dict[str, str | None]
    is_gp: bool


@dataclass
class FeeData:
    sheet: SheetModel
    layout: MgmtFeeLayout
    rows: list[FeeRow]
    rates: list[Decimal]
    period_fractions: list[Decimal]
    period_labels: list[str]
    totals: dict[str, Decimal]

    @property
    def total(self) -> Decimal:
        return sum(self.totals.values(), Decimal("0"))


def extract_fee(model: WorkbookModel, layout: MgmtFeeLayout) -> FeeData:
    sheet = model.sheet(layout.sheet)
    cols = layout.columns
    fee_cols = [f.column for f in layout.fee_columns]
    gp_rows = set(layout.gp_rows)
    rows = []
    for row in list(rows_between(layout.investor_rows)) + sorted(gp_rows):
        name = sheet.value(f"{cols.investor}{row}")
        if not is_text(name):
            continue
        rows.append(FeeRow(
            name=str(name).strip(),
            row=row,
            vehicle=sheet.value(f"{cols.vehicle}{row}") if cols.vehicle else None,
            affiliate_flag=sheet.value(f"{cols.affiliate_flag}{row}") if cols.affiliate_flag else None,
            commitment=to_decimal(sheet.value(f"{cols.commitment}{row}")) if cols.commitment else None,
            fees={c: to_money(sheet.value(f"{c}{row}")) for c in fee_cols},
            formulas={c: (sheet.cell(f"{c}{row}").formula if sheet.cell(f"{c}{row}") else None) for c in fee_cols},
            is_gp=row in gp_rows,
        ))
    total_row = layout.subtotal_rows.total
    totals = _row_values(sheet, total_row, fee_cols) if total_row else {
        c: sum((r.fees[c] for r in rows), Decimal("0")) for c in fee_cols}
    labels = []
    for fee in layout.fee_columns:
        labels.append(fee.period_label or str(sheet.value(f"{fee.column}{layout.header_row}") or "").strip())
    return FeeData(
        sheet=sheet,
        layout=layout,
        rows=rows,
        rates=[to_money(sheet.value(c)) for c in layout.rate_cells],
        period_fractions=[to_money(sheet.value(c)) for c in layout.period_fraction_cells],
        period_labels=labels,
        totals=totals,
    )


# --- Investor data / holidays ---------------------------------------------------------------


@dataclass
class InvestorDataRow:
    name: str
    row: int
    investor_id: Any
    fund_id: Any
    fund_name: Any


@dataclass
class InvestorData:
    sheet: SheetModel
    layout: InvestorDataLayout
    rows: list[InvestorDataRow]

    @property
    def by_investor_id(self) -> dict[Any, InvestorDataRow]:
        return {r.investor_id: r for r in self.rows if r.investor_id is not None}

    @property
    def by_name(self) -> dict[str, InvestorDataRow]:
        return {r.name: r for r in self.rows}


def extract_investor_data(model: WorkbookModel, layout: InvestorDataLayout) -> InvestorData:
    sheet = model.sheet(layout.sheet)
    cols = layout.columns
    rows = []
    for row in range(layout.first_data_row, layout.last_data_row + 1):
        name = sheet.value(f"{cols.investor_name}{row}")
        if not is_text(name):
            continue
        rows.append(InvestorDataRow(
            name=str(name).strip(),
            row=row,
            investor_id=_id(sheet.value(f"{cols.investor_id}{row}")) if cols.investor_id else None,
            fund_id=_id(sheet.value(f"{cols.fund_id}{row}")) if cols.fund_id else None,
            fund_name=sheet.value(f"{cols.fund_name}{row}") if cols.fund_name else None,
        ))
    return InvestorData(sheet=sheet, layout=layout, rows=rows)


def extract_holidays(model: WorkbookModel, layout: HolidayCalendarLayout) -> list[dt.date]:
    sheet = model.sheet(layout.sheet)
    days = [to_date(sheet.value(f"{layout.date_column}{r}")) for r in range(layout.first_row, layout.last_row + 1)]
    return [d for d in days if d is not None]


# --- Workbook ---------------------------------------------------------------------------------


@dataclass
class WorkbookData:
    model: WorkbookModel
    layouts: dict[str, BaseModel]
    allocation: AllocationData | None = None
    itd: ItdData | None = None
    summary: SummaryData | None = None
    merges: list[MergeData] = field(default_factory=list)
    mgmt_fee: FeeData | None = None
    investor_data: InvestorData | None = None
    holidays: list[dt.date] = field(default_factory=list)
    has_holiday_calendar: bool = False
    extraction_errors: dict[str, str] = field(default_factory=dict)  # role -> reason
    # Sheets outside the processed set that the processed sheets' formulas reference (e.g. an
    # investment tracker, or a hidden working sheet). They are scanned by sheet-wide rules.
    reference_sheets: list[str] = field(default_factory=list)
    role_notes: list[str] = field(default_factory=list)  # e.g. a hidden empty Merge-like tab left unprocessed
    prior: "WorkbookData | None" = None  # the prior event's workbook, when supplied

    @property
    def processed_sheets(self) -> list[str]:
        return list(self.layouts)

    @property
    def scanned_sheets(self) -> list[str]:
        """Sheets that sheet-wide rules (errors, placeholders) scan."""
        return self.processed_sheets + [s for s in self.reference_sheets if s not in self.layouts]


_SHEET_REF_RE = re.compile(r"'((?:[^']|'')+)'!|(?<![A-Za-z0-9_.'])([A-Za-z_][A-Za-z0-9_.]*)!")


def _sheet_refs(model: WorkbookModel, name: str) -> set[str]:
    known = {s.name for s in model.sheets}
    refs: set[str] = set()
    for cell in model.sheet(name).cells.values():
        if not cell.formula or "!" not in cell.formula:
            continue
        for quoted, bare in _SHEET_REF_RE.findall(cell.formula):
            target = quoted.replace("''", "'") if quoted else bare
            if target in known and target != name:
                refs.add(target)
    return refs


def referenced_sheets(model: WorkbookModel, sheet_names: list[str], hubs: list[str] | None = None) -> list[str]:
    """Other sheets linked to the processed ones, in workbook order.

    That is every sheet a formula on ``sheet_names`` references (hidden or not), plus every
    visible sheet whose formulas reference a ``hubs`` sheet (the Allocation and Summary):
    support tabs such as a fee or waterfall breakout. Hidden sheets that only read the
    Allocation are allocation breakouts saved from earlier events and are skipped.
    """
    found: set[str] = set()
    for name in sheet_names:
        found |= _sheet_refs(model, name) - set(sheet_names)
    for sheet in model.sheets:
        if hubs and sheet.name not in sheet_names and sheet.is_visible and _sheet_refs(model, sheet.name) & set(hubs):
            found.add(sheet.name)
    return [s.name for s in model.sheets if s.name in found]


def extract_workbook_data(model: WorkbookModel, layouts: dict[str, BaseModel],
                          prior: "WorkbookData | None" = None) -> WorkbookData:
    data = WorkbookData(model=model, layouts=dict(layouts), prior=prior)
    hubs = [n for n, layout in layouts.items() if isinstance(layout, (AllocationLayout, SummaryLayout))]
    data.reference_sheets = referenced_sheets(model, list(layouts), hubs)
    allocation_sheets = [n for n, layout in layouts.items() if isinstance(layout, AllocationLayout)]
    for name, layout in layouts.items():
        try:
            if isinstance(layout, AllocationLayout):
                data.allocation = extract_allocation(model, layout)
            elif isinstance(layout, ItdLayout):
                data.itd = extract_itd(model, layout, allocation_sheets)
            elif isinstance(layout, SummaryLayout):
                data.summary = extract_summary(model, layout)
            elif isinstance(layout, MergeLayout):
                data.merges.append(extract_merge(model, layout))
            elif isinstance(layout, MgmtFeeLayout):
                data.mgmt_fee = extract_fee(model, layout)
            elif isinstance(layout, InvestorDataLayout):
                data.investor_data = extract_investor_data(model, layout)
            elif isinstance(layout, HolidayCalendarLayout):
                data.holidays = extract_holidays(model, layout)
                data.has_holiday_calendar = True
        except Exception as exc:  # a bad layout must not take the whole run down
            data.extraction_errors[getattr(layout, "role", name)] = f"{type(exc).__name__}: {exc}"
    return data
