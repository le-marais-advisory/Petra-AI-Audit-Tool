"""SheetLayout models and the deterministic layout validator (plan section 3, steps 3-4).

The LLM layout mapper's output is only trusted after ``validate_layout`` accepts it.
Golden layouts from the fixture generator must validate cleanly; each corruption must
be reported with a specific issue code so the mapper can be re-prompted with it.
"""
from __future__ import annotations

import copy
import importlib

import pytest

from tests.fixtures.generate_capital_event_fixtures import EVENT_TYPES, VARIANTS

CLEAN = [(e, v) for e in EVENT_TYPES for v in VARIANTS]
LAYOUT_ROLES = ("allocation", "itd", "summary", "merge", "mgmt_fee", "investor_data", "holiday_calendar")


@pytest.fixture(scope="module")
def layout_mod():
    return importlib.import_module("src.pipeline.workbook.layout")


@pytest.fixture(scope="module")
def validator():
    return importlib.import_module("src.pipeline.workbook.layout_validator")


@pytest.fixture(scope="module")
def loader():
    return importlib.import_module("src.pipeline.workbook.loader")


def _layout(manifest, role):
    return copy.deepcopy(next(lay for lay in manifest.layouts.values() if lay["role"] == role))


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_golden_layouts_parse(layout_mod, capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    for sheet, raw in manifest.layouts.items():
        parsed = layout_mod.parse_layout(raw)
        assert parsed.role == raw["role"]
        assert parsed.sheet == sheet


@pytest.mark.parametrize("role", LAYOUT_ROLES)
def test_each_role_exposes_a_json_schema_for_structured_output(layout_mod, role):
    schema = layout_mod.layout_json_schema(role)
    assert schema["type"] == "object"
    assert "role" in schema["properties"]


def _optional_params(schema):
    if isinstance(schema, dict):
        own = 0
        if schema.get("type") == "object" and "properties" in schema:
            own = len(set(schema["properties"]) - set(schema.get("required", [])))
        return own + sum(_optional_params(v) for v in schema.values())
    if isinstance(schema, list):
        return sum(_optional_params(v) for v in schema)
    return 0


@pytest.mark.parametrize("role", LAYOUT_ROLES)
def test_schemas_fit_claude_structured_output_limits(layout_mod, role):
    # Claude structured outputs reject > 16 union-typed or > 24 optional parameters.
    schema = layout_mod.layout_json_schema(role)
    assert layout_mod.count_unions(schema) == 0
    assert _optional_params(schema) == 0


def test_sentinels_map_back_to_none(layout_mod):
    raw = {"header_row": 6, "grand_total_row": 0, "event": {"label_cell": "", "label": "Capital Call #4"},
           "components": [{"active": False, "column": "H"}]}
    out = layout_mod.from_llm_output(raw)
    assert out["grand_total_row"] is None and out["event"]["label_cell"] is None
    assert out["components"][0]["active"] is False and out["header_row"] == 6


def test_unknown_role_is_rejected(layout_mod):
    with pytest.raises(ValueError):
        layout_mod.parse_layout({"role": "balance_sheet", "sheet": "BS"})


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_golden_layouts_validate(layout_mod, validator, loader, capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    model = loader.load_workbook_model(manifest.path)
    for raw in manifest.layouts.values():
        issues = validator.validate_layout(model, layout_mod.parse_layout(raw))
        assert issues == [], (raw["sheet"], [(i.code, i.message) for i in issues])


def _mutations():
    def header_off_by_one(lay):
        lay["header_row"] += 1

    def investor_range_short(lay):
        lay["vehicles"][0]["investor_rows"][1] -= 2

    def component_shifted(lay):
        lay["components"][0]["column"] = lay["columns"]["commitment"]

    def driver_below_header(lay):
        lay["fund_driver_row"] = lay["header_row"] + 1
        lay["vehicles"][0]["driver_row"] = lay["header_row"] + 1

    def column_out_of_range(lay):
        lay["roll_forward"]["remaining_commitment"] = "XFA"

    def subtotal_row_wrong(lay):
        lay["vehicles"][0]["subtotal_rows"]["total"] -= 1

    return {
        "header_off_by_one": ("allocation", header_off_by_one, "header_mismatch"),
        "investor_range_short": ("allocation", investor_range_short, "investor_range_mismatch"),
        "component_shifted": ("allocation", component_shifted, "header_mismatch"),
        "driver_below_header": ("allocation", driver_below_header, "driver_row_position"),
        "column_out_of_range": ("allocation", column_out_of_range, "column_out_of_range"),
        "subtotal_row_wrong": ("allocation", subtotal_row_wrong, "subtotal_row_mismatch"),
    }


def _itd_mutations():
    def block_overlap(lay):
        lay["event_blocks"][0]["last_column"] = lay["event_blocks"][1]["first_column"]

    def band_row_wrong(lay):
        lay["classification_rows"]["cost_contributions"] = lay["event_header_row"]

    def two_current_blocks(lay):
        lay["event_blocks"][0]["is_current"] = True

    def event_label_wrong(lay):
        lay["event_blocks"][-1]["first_column"] = lay["event_blocks"][-1]["components"][1]["column"]

    return {
        "block_overlap": ("itd", block_overlap, "block_overlap"),
        "band_row_wrong": ("itd", band_row_wrong, "classification_row_mismatch"),
        "two_current_blocks": ("itd", two_current_blocks, "current_block_count"),
        "event_label_wrong": ("itd", event_label_wrong, "header_mismatch"),
    }


def _other_mutations():
    def summary_check_empty(lay):
        lay["check_cells"] = ["Z99"]

    def merge_rows_short(lay):
        lay["last_data_row"] -= 3

    def fee_column_wrong(lay):
        lay["fee_columns"][0]["column"] = "B"

    return {
        "summary_check_empty": ("summary", summary_check_empty, "cell_empty"),
        "merge_rows_short": ("merge", merge_rows_short, "investor_range_mismatch"),
        "fee_column_wrong": ("mgmt_fee", fee_column_wrong, "header_mismatch"),
    }


ALL_MUTATIONS = {**_mutations(), **_itd_mutations(), **_other_mutations()}


@pytest.mark.parametrize("variant", ["standard", "shifted", "two_vehicles"])
@pytest.mark.parametrize("name", sorted(ALL_MUTATIONS))
def test_corrupted_layouts_are_rejected(layout_mod, validator, loader, capital_event_fixtures, name, variant):
    role, mutate, code = ALL_MUTATIONS[name]
    manifest = capital_event_fixtures.get("capital_call", variant)
    model = loader.load_workbook_model(manifest.path)
    raw = _layout(manifest, role)
    mutate(raw)
    issues = validator.validate_layout(model, layout_mod.parse_layout(raw))
    assert code in {i.code for i in issues}, [(i.code, i.message) for i in issues]


def test_issues_are_actionable(layout_mod, validator, loader, capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    raw = _layout(manifest, "allocation")
    raw["header_row"] += 1
    issue = next(i for i in validator.validate_layout(model, layout_mod.parse_layout(raw)) if i.code == "header_mismatch")
    assert issue.sheet == "Allocation"
    assert issue.cell  # points at the cell that disagreed
    assert issue.message


def test_subtotal_may_cover_blank_padding_rows(layout_mod, validator, loader, tmp_path):
    # Real Merge tabs often total over blank template rows below the last investor.
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Merge"
    for col, header in zip("ABCDE", ["Investor", "DX Investor ID", "DX Fund ID", "File Name", "Cash Due"]):
        ws[f"{col}1"] = header
    for row, (name, amount) in enumerate([("Alpha LP", 100.0), ("Beta LLC", 50.0)], start=2):
        ws[f"A{row}"], ws[f"B{row}"], ws[f"C{row}"], ws[f"D{row}"], ws[f"E{row}"] = name, 1000 + row, 9, f"9_{row}", amount
    ws["A4"] = 0  # a template row whose name formula evaluates to 0
    ws["A7"], ws["E7"] = "TOTAL:", "=SUM(E2:E6)"
    path = tmp_path / "padding.xlsx"
    wb.save(path)
    layout = {"role": "merge", "sheet": "Merge", "header_row": 1, "first_data_row": 2, "last_data_row": 3,
              "columns": {"investor": "A", "investor_id": "B", "fund_id": "C", "file_name": "D", "event_total": "E"},
              "total_row": 7}
    model = loader.load_workbook_model(path)
    assert validator.validate_layout(model, layout_mod.parse_layout(layout)) == []
    ws["A5"], ws["E5"] = "Gamma Trust", 25.0  # a real investor outside investor_rows is still caught
    wb.save(path)
    codes = {i.code for i in validator.validate_layout(loader.load_workbook_model(path), layout_mod.parse_layout(layout))}
    assert "investor_range_mismatch" in codes


def test_fee_total_column_is_not_a_billing_period(layout_mod, validator, loader, tmp_path):
    # The reference sample bills 3Q and 4Q fees with a "Total" column beside them.
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mgmt Fee Calc"
    for col, header in zip("ABCDE", ["Investor", "Commitment Amount", "Q3 2026 Mgmt Fees", "Q4 2026 Mgmt Fees",
                                     "Total Mgmt Fees"]):
        ws[f"{col}1"] = header
    for row, name in enumerate(["Alpha LP", "Beta LLC"], start=2):
        ws[f"A{row}"], ws[f"B{row}"] = name, 1000.0
        ws[f"C{row}"], ws[f"D{row}"], ws[f"E{row}"] = f"=B{row}*0.005", f"=B{row}*0.005", f"=SUM(C{row}:D{row})"
    ws["A4"] = "Total"
    for col in "BCDE":
        ws[f"{col}4"] = f"=SUM({col}2:{col}3)"
    path = tmp_path / "fees.xlsx"
    wb.save(path)
    layout = {"role": "mgmt_fee", "sheet": "Mgmt Fee Calc", "header_row": 1, "investor_rows": [2, 3],
              "columns": {"investor": "A", "commitment": "B"},
              "fee_columns": [{"column": "C", "period_label": "Q3 2026"}, {"column": "D", "period_label": "Q4 2026"}],
              "subtotal_rows": {"limited_partners": 4}}
    model = loader.load_workbook_model(path)
    assert "fee_total_column" not in {i.code for i in validator.validate_layout(model, layout_mod.parse_layout(layout))}
    layout["fee_columns"].append({"column": "E", "period_label": ""})
    codes = {i.code for i in validator.validate_layout(model, layout_mod.parse_layout(layout))}
    assert "fee_total_column" in codes


def test_date_cells_must_hold_dates(layout_mod, validator, loader, capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    raw = _layout(manifest, "summary")
    raw["due_date_cell"] = raw["fund_commitment_cell"]
    codes = {i.code for i in validator.validate_layout(model, layout_mod.parse_layout(raw))}
    assert "not_a_date" in codes


def test_dates_inside_text_are_parsed():
    import datetime as dt

    from src.pipeline.workbook.cells import to_date

    assert to_date("Capital Call - due June 10, 2026") == dt.date(2026, 6, 10)
    assert to_date("Capital Call #4") is None
