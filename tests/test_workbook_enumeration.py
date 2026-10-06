"""Deterministic ITD block enumeration, lookup-aware pulls and the layout refinements behind them.

The LLM layout mapper's event blocks are a refinement: code reads the event header row itself
(labels, merged ranges, sub-headers, X marks), so a mapper that lists one block of twenty-eight
no longer silences the cumulative, history and frozen checks.
"""
from __future__ import annotations

import copy
import importlib

import openpyxl
import pytest

from src.pipeline.workbook.cells import sheet_column_refs
from src.pipeline.workbook.extract import (
    component_type_from_header,
    enumerate_itd_blocks,
    extract_workbook_data,
    pull_lookup,
)
from src.pipeline.workbook.layout import parse_layout
from src.pipeline.workbook.loader import load_workbook_model


def _itd_layout(manifest):
    return copy.deepcopy(next(lay for lay in manifest.layouts.values() if lay["role"] == "itd"))


def _plain_itd(tmp_path, merged: bool, totals: bool, marks: bool):
    """A small ITD sheet: cumulative columns B-D, three event blocks from column F."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ITD"
    ws["A1"], ws["A2"] = "Investment Contributions", "Cost Contributions"
    ws["A5"], ws["B5"], ws["C5"], ws["D5"] = "Investor", "Commitment", "Investment Contributions", "Cost Contributions"
    blocks = [("Capital Call #1 - 01.05.2024", "F", ["Investment", "Expenses"]),
              ("Capital Call #2 - 06.25.2024", "I", ["Investment"]),
              ("Capital Call #3 - 03.09.2026", "K", ["Investment", "Expenses", "Mgmt Fees"])]
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl

    for label, first, subs in blocks:
        start = ci(first)
        ws[f"{first}4"] = label
        for offset, sub in enumerate(subs):
            col = gl(start + offset)
            ws[f"{col}5"] = sub
            if marks:
                ws[f"{col}{1 if sub == 'Investment' else 2}"] = "X"
            for row in (6, 7):
                ws[f"{col}{row}"] = 100 * row + offset
        if totals:
            ws[f"{gl(start + len(subs))}5"] = "Total"
        if merged:
            ws.merge_cells(f"{first}4:{gl(start + len(subs) - (0 if totals else 1))}4")
    ws["A6"], ws["B6"], ws["A7"], ws["B7"] = "Alpha LP", 1000, "Beta LLC", 2000
    ws["A9"], ws["A12"] = "Limited Partners", "Total Partnership"
    path = tmp_path / "itd.xlsx"
    wb.save(path)
    layout = {"role": "itd", "sheet": "ITD",
              "classification_rows": {"investment_contributions": 1, "cost_contributions": 2},
              "event_header_row": 4, "subheader_row": 5, "investor_column": "A",
              "cumulative_columns": {"commitment": "B", "investment_contributions": "C", "cost_contributions": "D"},
              "event_blocks": [{"label": "Capital Call #3 - 03.09.2026", "event_type": "capital_call", "number": 3,
                                "first_column": "K", "last_column": "M", "is_current": True, "components": []}],
              "vehicles": [{"name": "Fund", "investor_rows": [6, 7], "subtotal_rows": {"limited_partners": 9, "total": 12}}]}
    return load_workbook_model(path).sheet("ITD"), parse_layout(layout)


@pytest.mark.parametrize("merged,totals,marks", [(True, True, True), (False, False, False), (True, False, True)])
def test_enumerates_every_block_from_the_header_row(tmp_path, merged, totals, marks):
    sheet, layout = _plain_itd(tmp_path, merged, totals, marks)
    blocks, notes = enumerate_itd_blocks(sheet, layout)
    assert [b.label for b in blocks] == ["Capital Call #1 - 01.05.2024", "Capital Call #2 - 06.25.2024",
                                         "Capital Call #3 - 03.09.2026"]
    assert [b.first_column for b in blocks] == ["F", "I", "K"]
    assert [[c.column for c in b.components] for b in blocks] == [["F", "G"], ["I"], ["K", "L", "M"]]
    assert [b.total_column for b in blocks] == (["H", "J", "N"] if totals else [None, None, None])
    assert [b.event_type for b in blocks] == ["capital_call"] * 3
    assert [b.number for b in blocks] == [1, 2, 3]
    assert blocks[0].date == "2024-01-05"
    assert [b.is_current for b in blocks] == [False, False, True]  # the mapper's current block is kept
    assert [c.component_type for c in blocks[2].components] == ["investment", "org_expense", "mgmt_fee"]
    assert all(c.side == "call" for b in blocks for c in b.components)
    assert notes == []


def test_side_is_assumed_from_the_sign_with_a_note(tmp_path):
    sheet, layout = _plain_itd(tmp_path, merged=False, totals=False, marks=False)
    sheet.cells["L5"].value = "Vortx"  # neither the sub-header nor an X mark says what it is
    blocks, notes = enumerate_itd_blocks(sheet, layout)
    assert blocks[2].components[1].component_type == "other"
    assert blocks[2].components[1].side == "call"  # positive values
    assert notes and "L" in notes[0] and "assumed call" in notes[0]


def test_mapper_blocks_are_used_when_the_header_row_has_no_labels(tmp_path):
    sheet, layout = _plain_itd(tmp_path, merged=True, totals=True, marks=True)
    for coord in ("F4", "I4", "K4"):
        del sheet.cells[coord]
    blocks, notes = enumerate_itd_blocks(sheet, layout)
    assert [b.label for b in blocks] == ["Capital Call #3 - 03.09.2026"]
    assert any("mapped blocks were used" in n for n in notes)


def test_current_block_falls_back_to_the_block_linking_to_the_allocation(tmp_path):
    sheet, layout = _plain_itd(tmp_path, merged=True, totals=True, marks=True)
    raw = layout.model_dump(mode="json")
    raw["event_blocks"] = []
    layout = parse_layout(raw)
    sheet.cells["K6"].formula = "=Allocation!H7"
    blocks, notes = enumerate_itd_blocks(sheet, layout, ["Allocation"])
    assert [b.is_current for b in blocks] == [False, False, True]
    assert any("taken as the current block" in n for n in notes)


@pytest.mark.parametrize("header,expected", [
    ("Vagaro A-2 Non-Recallable Realized Gain", "realized_gain"),
    ("Egress (Escrow) Recallable", "return_of_capital"),
    ("Dist Roc", "return_of_capital"),
    ("Tango Interest Income", "dividend_income"),
    ("Vertera (Vortx) Dividend Non-Recallable", "dividend_income"),
    ("Deemed GP Contribution", "investment"),
    ("Working Capital", "investment"),
    ("Partnership Expenses / Org Costs", "org_expense"),
    ("Q3 & Q4 2025 Mgmt Fees", "mgmt_fee"),
    ("Dist Tax WH", "tax_withholding"),
    ("TangoCard Realized", "realized_gain"),
    ("Carry", "carry"),
    ("Transfer", "other"),
    ("", "other"),
])
def test_component_type_from_sub_header(header, expected):
    assert component_type_from_header(header) == expected


def test_fixture_itd_enumerates_the_golden_blocks(capital_event_fixtures):
    manifest = capital_event_fixtures.get("capital_call", "multi_vehicle")
    model = load_workbook_model(manifest.path)
    layout = parse_layout(_itd_layout(manifest))
    blocks, notes = enumerate_itd_blocks(model.sheet(layout.sheet), layout, ["Allocation"])
    golden = _itd_layout(manifest)["event_blocks"]
    assert [b.label for b in blocks] == [g["label"] for g in golden]
    assert [b.first_column for b in blocks] == [g["first_column"] for g in golden]
    assert [b.total_column for b in blocks] == [g["total_column"] for g in golden]
    assert [[c.column for c in b.components] for b in blocks] == [[c["column"] for c in g["components"]] for g in golden]
    assert notes == []


def test_blocks_the_mapper_left_out_are_still_extracted(capital_event_fixtures):
    manifest = capital_event_fixtures.get(defect="itd_blocks_in_collapsed_groups")
    model = load_workbook_model(manifest.path)
    raw = _itd_layout(manifest)
    assert len(raw["event_blocks"]) == 1  # the mapper only returned the current block
    layouts = {n: parse_layout(r) for n, r in manifest.layouts.items() if n in manifest.relevant_sheets}
    itd = extract_workbook_data(model, layouts).itd
    assert len(itd.event_blocks) == len(manifest.truth["itd"]["event_labels"])
    assert itd.current_block.label == manifest.truth["itd"]["current_block_label"]
    assert any(b.total_column is None for b in itd.prior_blocks)
    sheet = model.sheet(itd.sheet.name)
    hidden_blocks = [b for b in itd.prior_blocks if b.first_column in sheet.hidden_cols]
    assert len(hidden_blocks) >= 20 and all(b.first_column in sheet.outlined_cols for b in hidden_blocks)


def test_named_rows_without_commitment_or_amounts_are_not_investors(capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = load_workbook_model(manifest.path)
    layouts = {n: parse_layout(r) for n, r in manifest.layouts.items() if n in manifest.relevant_sheets}
    itd_layout = layouts[next(n for n, l in layouts.items() if l.role == "itd")]
    sheet = model.sheet(itd_layout.sheet)
    vehicle = itd_layout.vehicles[0]
    last = vehicle.investor_rows[1]
    # A text-only row ("Unallocated") inside the investor span, and a transferred-out LP with history.
    from src.pipeline.workbook.loader import CellModel

    sheet.cells[f"{itd_layout.investor_column}{last + 1}"] = CellModel(f"{itd_layout.investor_column}{last + 1}", last + 1,
                                                                        itd_layout.investor_column, "Unallocated", None)
    raw = itd_layout.model_dump(mode="json")
    raw["vehicles"][0]["investor_rows"][1] = last + 1
    layouts[itd_layout.sheet] = parse_layout(raw)
    itd = extract_workbook_data(model, layouts).itd
    assert "Unallocated" not in {i.name for i in itd.investors}
    assert itd.excluded_rows == [(last + 1, "Unallocated")]


# --- whole-column references and lookup pulls -----------------------------------------------


def test_sheet_column_refs():
    assert sheet_column_refs("=SUMIF(Allocation!$D:$D,$E4,Allocation!AE:AE)") == [("Allocation", "D"), ("Allocation", "AE")]
    assert sheet_column_refs("=SUMIFS('Mgmt Fee Calc'!$F:$F,'Mgmt Fee Calc'!$B:$B,$C10)") == [
        ("Mgmt Fee Calc", "F"), ("Mgmt Fee Calc", "B")]
    assert sheet_column_refs("=Allocation!Q7") == []
    assert sheet_column_refs("=SUM(Allocation!Q7:T7)") == []


def test_pull_lookup_recognises_a_lookup_keyed_by_the_rows_own_cell():
    assert pull_lookup("=SUMIF(Allocation!$D:$D,'FTV Merge'!$E4,Allocation!AE:AE)", 4, "FTV Merge") == ("Allocation", "AE")
    assert pull_lookup("=SUMIFS(Allocation!AE:AE,Allocation!$D:$D,$A12)", 12, "ITD") == ("Allocation", "AE")
    # Not pulls: a check formula, a lookup keyed by a fixed cell, a plain link, a lookup into the home sheet.
    assert pull_lookup("=SUMIF(Allocation!$D:$D,$E4,Allocation!AE:AE)-Q4", 4, "Merge") is None
    assert pull_lookup("=SUMIF(Allocation!$D:$D,$E$2,Allocation!AE:AE)", 4, "Merge") is None
    assert pull_lookup("=Allocation!AE7", 7, "Merge") is None
    assert pull_lookup("=SUMIF(Merge!$D:$D,$E4,Merge!AE:AE)", 4, "Merge") is None


def test_sumif_pulls_are_tied_like_plain_links(capital_event_fixtures):
    manifest = capital_event_fixtures.get(defect="ok_sumif_pulls")
    model = load_workbook_model(manifest.path)
    layouts = {n: parse_layout(r) for n, r in manifest.layouts.items() if n in manifest.relevant_sheets}
    data = extract_workbook_data(model, layouts)
    merge = data.merges[0]
    alloc_components = {c["column"] for c in manifest.layouts["Allocation"]["components"] if c["active"]}
    assert alloc_components <= set(merge.referenced_columns.values())
    assert merge.source_sheet == "Allocation"


# --- a single Summary date cell mapped as both dates ----------------------------------------


def test_summary_single_date_cell_counts_once_as_the_due_date(capital_event_fixtures):
    manifest = capital_event_fixtures.get(defect="ok_summary_single_text_date")
    model = load_workbook_model(manifest.path)
    raw = {n: copy.deepcopy(r) for n, r in manifest.layouts.items() if n in manifest.relevant_sheets}
    summary = next(r for r in raw.values() if r["role"] == "summary")
    summary["notice_date_cell"] = summary["due_date_cell"]  # the mapper's answer on the reference client
    validator = importlib.import_module("src.pipeline.workbook.layout_validator")
    issues = validator.validate_layout(model, parse_layout(summary))
    assert [i.code for i in issues] == ["same_date_cell"] and not validator.blocking_issues(issues)
    layouts = {n: parse_layout(r) for n, r in raw.items()}
    data = extract_workbook_data(model, layouts)
    assert data.summary.notice_date is None and data.summary.due_date is not None
    checks = importlib.import_module("src.pipeline.workbook.checks")
    rules = [{"id": "CE-DATE-ORDER", "name": "x"}, {"id": "CE-DATE-VALIDITY", "name": "y"}]
    results = checks.run_deterministic_checks(model, data, rules, {"event_type": "capital_call"})
    assert results["CE-DATE-ORDER"].verdict == "pass"
    assert any("one date" in f for f in results["CE-DATE-ORDER"].findings)
    assert results["CE-DATE-VALIDITY"].verdict == "pass"


# --- layout warnings are accepted after one re-prompt --------------------------------------


def test_block_coverage_is_a_warning_the_mapper_may_leave_unresolved(capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = load_workbook_model(manifest.path)
    validator = importlib.import_module("src.pipeline.workbook.layout_validator")
    raw = _itd_layout(manifest)
    raw["event_blocks"] = [b for b in raw["event_blocks"] if b["is_current"]]
    issues = validator.validate_layout(model, parse_layout(raw))
    assert [i.code for i in issues] == ["block_coverage"]
    assert not validator.blocking_issues(issues)
    mapper = importlib.import_module("src.pipeline.workbook.layout_mapper")

    class Partial:
        def __init__(self):
            self.calls = 0

        def complete_structured(self, system_prompt, user_content, json_schema, name="result", effort=None):
            self.calls += 1
            return copy.deepcopy(raw)

    provider = Partial()
    layout, issues = mapper.map_sheet_layout(model, raw["sheet"], "itd", provider, use_cache=False)
    assert layout is not None and provider.calls == 2  # re-prompted once, then accepted with the warning
    assert [i.code for i in issues] == ["block_coverage"]
