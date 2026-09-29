"""Workbook loader, sheet inventory, role proposal, sheet selection and skeletons.

Covers plan section 2 (runtime selection) and section 3 steps 1-2. None of this calls
an LLM.
"""
from __future__ import annotations

import importlib

import pytest

from tests.fixtures.generate_capital_event_fixtures import EVENT_TYPES, VARIANTS


@pytest.fixture(scope="module")
def loader():
    return importlib.import_module("src.pipeline.workbook.loader")


@pytest.fixture(scope="module")
def inventory_mod():
    return importlib.import_module("src.pipeline.workbook.inventory")


@pytest.fixture(scope="module")
def roles_mod():
    return importlib.import_module("src.pipeline.workbook.roles")


@pytest.fixture(scope="module")
def skeleton_mod():
    return importlib.import_module("src.pipeline.workbook.skeleton")


# --- loader -------------------------------------------------------------------------


def test_loader_reads_formulas_values_and_formats(loader, capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    assert model.file_name == manifest.path.name
    assert [s.name for s in model.sheets] == list(manifest.sheet_roles)
    alloc = model.sheet("Allocation")
    assert alloc.index == 3
    cell = alloc.cell("H7")
    assert cell.formula == "=ROUND(H$5*$F7,2)"
    assert cell.value == pytest.approx(float(manifest.truth["allocation"]["investors"]["Northgate Family Trust"]["amounts"]["H"]))
    assert "#,##0.00" in cell.number_format
    literal = alloc.cell("E7")
    assert literal.formula is None and literal.value == 12500000
    assert alloc.cell("ZZ999") is None
    assert model.formulas_missing_cache is False


def test_loader_captures_sheet_metadata(loader, capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    alloc = model.sheet("Allocation")
    assert alloc.state == "visible"
    assert alloc.view == "pageBreakPreview"
    late_interest = manifest.layouts["Allocation"]["columns"]["late_interest"]
    assert alloc.hidden_cols == {late_interest}
    assert "B2:C2" in alloc.merged_ranges
    assert model.sheet("3rd Close Rebalance").state == "hidden"


def test_loader_flags_hidden_rows_and_errors(loader, capital_event_fixtures):
    hidden = loader.load_workbook_model(capital_event_fixtures.get(defect="hidden_populated_row").path)
    assert hidden.sheet("Allocation").hidden_rows
    errors = loader.load_workbook_model(capital_event_fixtures.get(defect="formula_error").path)
    error_cells = [c for c in errors.sheet("Allocation").cells.values() if c.is_error]
    assert [c.value for c in error_cells] == ["#REF!"]


def test_loader_detects_missing_cached_values(loader, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A2"] = "=A1*2"  # openpyxl never writes a cached value
    path = tmp_path / "no_cache.xlsx"
    wb.save(path)
    assert loader.load_workbook_model(path).formulas_missing_cache is True


# --- inventory + roles ------------------------------------------------------------------


def test_inventory_covers_every_sheet_cheaply(loader, inventory_mod, capital_event_fixtures):
    model = loader.load_workbook_model(capital_event_fixtures.get().path)
    entries = inventory_mod.build_inventory(model)
    assert [e.name for e in entries] == [s.name for s in model.sheets]
    alloc = next(e for e in entries if e.name == "Allocation")
    assert alloc.state == "visible" and alloc.formula_count > 100
    assert any("Investor" in cell for row in alloc.header_preview for cell in row)
    assert next(e for e in entries if e.name == "3rd Close Rebalance").state == "hidden"


@pytest.mark.parametrize("variant", list(VARIANTS))
@pytest.mark.parametrize("event_type", EVENT_TYPES)
def test_heuristic_roles_match_ground_truth(loader, roles_mod, capital_event_fixtures, event_type, variant):
    # Includes the Merge tab named after the fund and the fee tab named "MF" (shifted variant).
    manifest = capital_event_fixtures.get(event_type, variant)
    model = loader.load_workbook_model(manifest.path)
    assert roles_mod.propose_roles(model) == manifest.sheet_roles


@pytest.mark.parametrize("event_type", EVENT_TYPES)
def test_relevant_sheets_follow_the_event_type(loader, roles_mod, capital_event_fixtures, event_type):
    manifest = capital_event_fixtures.get(event_type, "two_vehicles")
    selection = importlib.import_module("src.pipeline.workbook.selection")
    assert selection.relevant_sheets(manifest.sheet_roles, event_type) == manifest.relevant_sheets


# --- skeleton -----------------------------------------------------------------------------


@pytest.mark.parametrize("variant", list(VARIANTS))
def test_allocation_skeleton_is_compact_and_complete(loader, skeleton_mod, capital_event_fixtures, variant):
    manifest = capital_event_fixtures.get("capital_call", variant)
    model = loader.load_workbook_model(manifest.path)
    layout = next(lay for lay in manifest.layouts.values() if lay["role"] == "allocation")
    sheet = model.sheet(layout["sheet"])
    text = skeleton_mod.build_skeleton(sheet)
    naive = sum(len(c.coord) + len(str(c.value)) + len(c.formula or "") + 2 for c in sheet.cells.values())
    assert len(text) < naive * 0.35
    assert skeleton_mod.estimate_tokens(text) <= 6000
    for comp in layout["components"]:
        assert comp["header"] in text
    for label in ("Limited Partners", "General Partner", "Check", "Prior Capital Contributions"):
        assert label in text
    # Row labels down the sheet and the merged ranges are kept.
    for name in manifest.truth["allocation"]["investors"]:
        assert name in text
    for rng in sheet.merged_ranges:
        assert rng in text


def test_skeleton_lists_formula_exceptions(loader, skeleton_mod, capital_event_fixtures):
    # The single investment plug (=ROUND(...)+0.02) is an exception to the column pattern.
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    text = skeleton_mod.build_skeleton(model.sheet("Allocation"))
    assert "H9" in text and "+0.02" in text


def test_itd_skeleton_keeps_band_and_event_headers(loader, skeleton_mod, capital_event_fixtures):
    manifest = capital_event_fixtures.get()
    model = loader.load_workbook_model(manifest.path)
    text = skeleton_mod.build_skeleton(model.sheet("ITD Capital Activity"))
    for label in manifest.truth["itd"]["event_labels"]:
        assert label in text
    for label in ("Investment Contributions", "Cost Contributions", "Non-Recallable Distributions"):
        assert label in text
    assert skeleton_mod.estimate_tokens(text) <= 6000
