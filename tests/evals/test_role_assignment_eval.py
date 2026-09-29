"""Role-assignment eval: heuristic proposal + one LLM confirmation over the inventory."""
from __future__ import annotations

import importlib

import pytest

from tests.fixtures.generate_capital_event_fixtures import EVENT_TYPES, VARIANTS

pytestmark = pytest.mark.eval


@pytest.mark.parametrize("variant", list(VARIANTS))
@pytest.mark.parametrize("event_type", EVENT_TYPES)
def test_assigned_roles_match_ground_truth(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    loader = importlib.import_module("src.pipeline.workbook.loader")
    roles = importlib.import_module("src.pipeline.workbook.roles")
    model = loader.load_workbook_model(manifest.path)
    assert roles.assign_roles(model) == manifest.sheet_roles


def test_llm_corrects_a_misleading_sheet_name(capital_event_fixtures, tmp_path):
    # Rename the Merge tab to something the heuristic cannot place by name alone.
    import openpyxl

    manifest = capital_event_fixtures.get()
    wb = openpyxl.load_workbook(manifest.path)
    wb["Merge"].title = "Letters"
    path = tmp_path / manifest.path.name
    wb.save(path)
    loader = importlib.import_module("src.pipeline.workbook.loader")
    roles = importlib.import_module("src.pipeline.workbook.roles")
    assert roles.assign_roles(loader.load_workbook_model(path))["Letters"] == "merge"
