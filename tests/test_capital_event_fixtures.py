"""Self-checks for the synthetic capital-event workbook generator.

These guard the ground truth the workbook suites rely on: every formula carries a
cached value, clean fixtures tie out end to end, and each golden layout points at the
cells it claims to. They use openpyxl directly, independent of src/pipeline/workbook.
"""
from __future__ import annotations

import warnings
from decimal import Decimal

import openpyxl
import pytest

from tests.fixtures.generate_capital_event_fixtures import (
    DEFECTS,
    DETERMINISTIC_RULE_IDS,
    EVENT_TYPES,
    ROLES,
    VARIANTS,
    default_specs,
)

TOL = Decimal("0.005")
CLEAN = [(e, v) for e in EVENT_TYPES for v in VARIANTS]


def _load(path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


def _num(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _rows(span: list[int]) -> range:
    return range(span[0], span[1] + 1)


@pytest.mark.parametrize("spec", default_specs(), ids=lambda s: s.fixture_id)
def test_every_formula_has_a_cached_value(capital_event_fixtures, spec):
    manifest = capital_event_fixtures.get(spec.event_type, spec.variant, spec.defect)
    wb, wv = _load(manifest.path)
    missing = [
        f"{ws.title}!{cell.coordinate}"
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=") and wv[ws.title][cell.coordinate].value is None
    ]
    assert missing == []


@pytest.mark.parametrize("spec", default_specs(), ids=lambda s: s.fixture_id)
def test_manifest_shape(capital_event_fixtures, spec):
    manifest = capital_event_fixtures.get(spec.event_type, spec.variant, spec.defect)
    wb, _ = _load(manifest.path)
    assert set(manifest.sheet_roles) == set(wb.sheetnames)
    assert set(manifest.sheet_roles.values()) <= set(ROLES)
    assert set(manifest.layouts) <= set(wb.sheetnames)
    assert set(manifest.expected_verdicts) == set(DETERMINISTIC_RULE_IDS)
    assert {"allocation", "itd", "summary", "merge"} <= {manifest.sheet_roles[s] for s in manifest.relevant_sheets}
    # Legacy / support sheets are never part of the processed set.
    assert "3rd Close Rebalance" not in manifest.relevant_sheets
    assert "Notes" not in manifest.relevant_sheets


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_allocation_ties_out(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    _, wv = _load(manifest.path)
    (sheet_name, layout), = [(s, lay) for s, lay in manifest.layouts.items() if lay["role"] == "allocation"]
    ws = wv[sheet_name]
    fund_driver_row = layout["fund_driver_row"]
    for comp in layout["components"]:
        col = comp["column"]
        grand = Decimal("0")
        for vehicle in layout["vehicles"]:
            rows = list(_rows(vehicle["investor_rows"])) + vehicle["gp_rows"]
            per_lp = sum(_num(ws[f"{col}{r}"].value) for r in rows)
            driver = _num(ws[f"{col}{vehicle['driver_row']}"].value)
            total = _num(ws[f"{col}{vehicle['subtotal_rows']['total']}"].value)
            assert abs(per_lp - driver) <= TOL, (col, vehicle["name"])
            assert abs(total - driver) <= TOL, (col, vehicle["name"])
            grand += total
        assert abs(grand - _num(ws[f"{col}{fund_driver_row}"].value)) <= TOL, col
    for row in layout["check_rows"]:
        for cell in ws[row]:
            if isinstance(cell.value, (int, float)):
                assert abs(_num(cell.value)) <= TOL, cell.coordinate


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_roll_forward_foots(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    _, wv = _load(manifest.path)
    (sheet_name, layout), = [(s, lay) for s, lay in manifest.layouts.items() if lay["role"] == "allocation"]
    ws = wv[sheet_name]
    rf = layout["roll_forward"]
    for vehicle in layout["vehicles"]:
        for r in _rows(vehicle["investor_rows"]):
            commitment = _num(ws[f"{rf['commitment']}{r}"].value)
            prior = abs(_num(ws[f"{rf['prior_contributions']}{r}"].value))
            current = abs(_num(ws[f"{rf['current_call']}{r}"].value))
            remaining = _num(ws[f"{rf['remaining_commitment']}{r}"].value)
            assert abs(commitment - prior - current - remaining) <= Decimal("0.01"), r
            assert remaining >= 0, r


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_itd_cumulatives_and_current_block_tie(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    _, wv = _load(manifest.path)
    (itd_name, itd), = [(s, lay) for s, lay in manifest.layouts.items() if lay["role"] == "itd"]
    ws = wv[itd_name]
    band = itd["classification_rows"]
    cum = itd["cumulative_columns"]
    block_cols = [c["column"] for b in itd["event_blocks"] for c in b["components"]]
    marks = {col: [cat for cat, row in band.items() if ws[f"{col}{row}"].value == "X"] for col in block_cols}
    assert all(len(cats) == 1 for cats in marks.values()), marks
    for vehicle in itd["vehicles"]:
        for r in _rows(vehicle["investor_rows"]):
            for cat in ("investment_contributions", "cost_contributions", "non_recallable_distributions"):
                expected = sum(_num(ws[f"{col}{r}"].value) for col, cats in marks.items() if cats == [cat])
                assert abs(_num(ws[f"{cum[cat]}{r}"].value) - expected) <= TOL, (r, cat)
            unfunded = _num(ws[f"{cum['commitment']}{r}"].value) - _num(ws[f"{cum['total_contributions']}{r}"].value)
            assert abs(_num(ws[f"{cum['unfunded']}{r}"].value) - unfunded) <= TOL, r
    current = [b for b in itd["event_blocks"] if b["is_current"]]
    assert len(current) == 1
    truth = manifest.truth["allocation"]["investors"]
    alloc_layout = next(lay for lay in manifest.layouts.values() if lay["role"] == "allocation")
    alloc_cols = {c["component_type"]: c["column"] for c in alloc_layout["components"] if c["active"]}
    names = {ws[f"{itd['investor_column']}{r}"].value: r for v in itd["vehicles"] for r in _rows(v["investor_rows"])}
    for name, r in names.items():
        for comp in current[0]["components"]:
            expected = _num(truth[name]["amounts"].get(alloc_cols[comp["component_type"]], "0"))
            assert abs(_num(ws[f"{comp['column']}{r}"].value) - expected) <= TOL, (name, comp)


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_summary_and_merge_tie(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    _, wv = _load(manifest.path)
    summary = next(lay for lay in manifest.layouts.values() if lay["role"] == "summary")
    ws = wv[summary["sheet"]]
    for cell in summary["check_cells"]:
        assert abs(_num(ws[cell].value)) <= TOL
    assert abs(_num(ws[summary["event_total_cell"]].value) - _num(manifest.truth["allocation"]["event_gross"])) <= TOL

    dx = next(lay for lay in manifest.layouts.values() if lay["role"] == "investor_data")
    dxs = wv[dx["sheet"]]
    dx_ids = {dxs[f"{dx['columns']['investor_name']}{r}"].value: dxs[f"{dx['columns']['investor_id']}{r}"].value
              for r in range(dx["first_data_row"], dx["last_data_row"] + 1)}
    for merge in (lay for lay in manifest.layouts.values() if lay["role"] == "merge"):
        ms = wv[merge["sheet"]]
        for r in range(merge["first_data_row"], merge["last_data_row"] + 1):
            name = ms[f"{merge['columns']['investor']}{r}"].value
            assert dx_ids[name] == ms[f"{merge['columns']['investor_id']}{r}"].value
            file_name = ms[f"{merge['columns']['file_name']}{r}"].value
            assert file_name.startswith(f"{ms[merge['columns']['fund_id'] + str(r)].value}_{dx_ids[name]}_")


@pytest.mark.parametrize("event_type,variant", CLEAN)
def test_golden_layout_headers_point_at_labels(capital_event_fixtures, event_type, variant):
    manifest = capital_event_fixtures.get(event_type, variant)
    _, wv = _load(manifest.path)
    alloc = next(lay for lay in manifest.layouts.values() if lay["role"] == "allocation")
    ws = wv[alloc["sheet"]]
    header = alloc["header_row"]
    assert ws[f"{alloc['columns']['investor']}{header}"].value in ("Investor", "Partner Name")
    assert "ommitment" in ws[f"{alloc['columns']['commitment']}{header}"].value
    for comp in alloc["components"]:
        assert ws[f"{comp['column']}{header}"].value == comp["header"]
    for key, col in alloc["roll_forward"].items():
        assert ws[f"{col}{header}"].value, key
    itd = next(lay for lay in manifest.layouts.values() if lay["role"] == "itd")
    wi = wv[itd["sheet"]]
    for block in itd["event_blocks"]:
        assert wi[f"{block['first_column']}{itd['event_header_row']}"].value == block["label"]
        assert wi[f"{block['total_column']}{itd['subheader_row']}"].value == "Total"


@pytest.mark.parametrize("name", sorted(DEFECTS))
def test_each_defect_changes_the_workbook(capital_event_fixtures, name):
    defect = DEFECTS[name]
    event_type = defect.event_types[0]
    clean = capital_event_fixtures.get(event_type, "standard")
    broken = capital_event_fixtures.get(event_type, "standard", name)
    if name == "generic_file_name":
        assert broken.path.name == "Book1.xlsx" != clean.path.name
        return
    wb_c, wv_c = _load(clean.path)
    wb_b, wv_b = _load(broken.path)

    def snapshot(wb, wv):
        out = {}
        for ws in wb.worksheets:
            out[(ws.title, "state")] = (ws.sheet_state, ws.sheet_view.view, tuple(sorted(map(str, ws.merged_cells.ranges))),
                                        tuple(sorted(r for r, d in ws.row_dimensions.items() if d.hidden)))
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        out[(ws.title, cell.coordinate)] = (cell.value, wv[ws.title][cell.coordinate].value,
                                                            cell.number_format)
        return out

    assert snapshot(wb_c, wv_c) != snapshot(wb_b, wv_b)


@pytest.mark.parametrize("name", sorted(DEFECTS))
def test_only_calibration_probes_leave_verdicts_unasserted(name):
    defect = DEFECTS[name]
    unasserted = [rule for rule, verdict in defect.verdicts.items() if verdict is None]
    assert bool(unasserted) <= (defect.pending_calibration is not None)
