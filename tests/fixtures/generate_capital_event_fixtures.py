"""Generate synthetic capital-event workbooks for the workbook test and eval suites.

The workbooks mirror the structure of a real fund capital-event roll-forward model
(Allocation / ITD Capital Activity / Summary / Merge / DX Investor Data / Mgmt Fee
Calc / bank holidays, plus legacy sheets), with fictional investors and amounts. No
real client data is used.

Every fixture is described by a ``FixtureSpec`` (event type x layout variant x an
optional seeded defect) and ``build_fixture`` returns a ``FixtureManifest`` carrying
the ground truth the tests assert against:

  - ``sheet_roles``      sheet name -> role (allocation, itd, summary, merge, ...)
  - ``layouts``          sheet name -> golden layout dict (the contract the LLM
                         layout mapper must reproduce and the extractor consumes)
  - ``truth``            key numbers per sheet, for extraction tests
  - ``expected_verdicts`` rule id -> verdict for the deterministic rules
                         (``None`` = not asserted until the FA calibration answer)
  - ``facts``            expectations for the hybrid-rule facts pre-pass

openpyxl cannot write a formula together with its cached value, so the generator
computes every value in Python (following the same data flow Excel would) and injects
the cached ``<v>`` values into the saved sheet XML.

Run from the repo root to write the standard set used by the eval suite:
    python tests/fixtures/generate_capital_event_fixtures.py
which writes to tests/fixtures/documents/capital_event/.
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent / "documents" / "capital_event"

MONEY_FMT = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
PCT_FMT = "0.0000%"
DATE_LONG_FMT = "mmmm d, yyyy"
DATE_SHORT_FMT = "m/d/yyyy"

EXCEL_EPOCH = dt.date(1899, 12, 30)
CENT = Decimal("0.01")

FUND_NAME = "Example Growth Fund III, LP"
FUND_SHORT = "egf-iii"
MGMT_FEE_RATE = Decimal("0.02")
FEE_PERIOD_FRACTION = Decimal("0.25")
CARRY_RATE = Decimal("0.20")

EVENT_TYPES = ("capital_call", "distribution", "net_event")

# Deterministic rules (plan section 4). Hybrid rules are covered by the facts tests.
DETERMINISTIC_RULE_IDS = (
    "CE-ALLOC-PER-LP-FORMULAS",
    "CE-ALLOC-VEHICLE-TIE",
    "CE-ALLOC-GROSS-TIE",
    "CE-ALLOC-REFOOT",
    "CE-ALLOC-PRO-RATA-PARITY",
    "CE-ALLOC-PLUG-DISCIPLINE",
    "CE-ALLOC-COMMITMENTS",
    "CE-ALLOC-ROUNDING",
    "CE-ALLOC-MERGED-CELLS",
    "CE-RF-FOOTING",
    "CE-RF-CURRENT-CALL-LINK",
    "CE-ITD-PRIOR-FROZEN",
    "CE-ITD-CUMULATIVE",
    "CE-ITD-EVENT-SEQUENCE",
    "CE-SUM-CHECKS-ZERO",
    "CE-DATE-VALIDITY",
    "CE-DATE-ORDER",
    "CE-FMT-DATE-DISPLAY",
    "CE-FMT-ACCOUNTING",
    "CE-FMT-NO-FORMULA-ERRORS",
    "CE-WB-NO-HIDDEN-DATA",
    "CE-WB-PAGE-BREAK-VIEW",
    "CE-WB-FILE-NAMING",
    "CE-TIE-MGMT-FEE",
    "CE-TIE-SUPPORT-TABS",
    "CE-TIE-ITD-ALLOCATION",
    "CE-TIE-ITD-COMMITMENTS",
    "CE-TIE-MERGE",
    "CE-ID-INVESTOR-KEYS",
    "CE-DIST-ROC-LIMIT",
    # Cross-event rules: compare with the prior event's workbook (FA calibration).
    "CE-XEV-HISTORY-UNCHANGED",
    "CE-XEV-ROLL-FORWARD",
    "CE-XEV-ITD-ROLL-FORWARD",
    "CE-XEV-PLUG-CONSISTENCY",
)

CROSS_EVENT_RULE_IDS = ("CE-XEV-HISTORY-UNCHANGED", "CE-XEV-ROLL-FORWARD", "CE-XEV-ITD-ROLL-FORWARD",
                        "CE-XEV-PLUG-CONSISTENCY")

HYBRID_RULE_IDS = (
    "CE-WB-SHEETS-PRESENT",
    "CE-WB-MERGE-TABS",
    "CE-WB-NO-PLACEHOLDERS",
    "CE-ALLOC-REFERENCE-INTEGRITY",
    "CE-ALLOC-STALE-COMPONENTS",
    "CE-ALLOC-SIGNAGE",
    "CE-ITD-EVENT-BLOCK",
    "CE-TIE-SUMMARY",
    "CE-DATE-CONSISTENCY",
    "CE-NET-EVENT-STRUCTURE",
    "CE-DIST-CARRY-SPLIT",
)

ROLES = ("allocation", "itd", "summary", "merge", "mgmt_fee", "investor_data", "holiday_calendar", "other")

# Sheet roles each event type needs (mirrors config/document_types/capital_event.yaml).
RELEVANT_ROLES = {
    "capital_call": {"allocation", "itd", "summary", "merge", "mgmt_fee", "investor_data", "holiday_calendar"},
    "distribution": {"allocation", "itd", "summary", "merge", "investor_data", "holiday_calendar"},
    "net_event": {"allocation", "itd", "summary", "merge", "mgmt_fee", "investor_data", "holiday_calendar"},
}

HOLIDAYS_2026 = [
    ("New Year's Day", dt.date(2026, 1, 1)),
    ("Martin Luther King, Jr. Day", dt.date(2026, 1, 19)),
    ("Presidents' Day", dt.date(2026, 2, 16)),
    ("Good Friday", dt.date(2026, 4, 3)),
    ("Memorial Day", dt.date(2026, 5, 25)),
    ("Juneteenth", dt.date(2026, 6, 19)),
    ("Independence Day (observed)", dt.date(2026, 7, 3)),
    ("Labor Day", dt.date(2026, 9, 7)),
    ("Columbus Day", dt.date(2026, 10, 12)),
    ("Veterans Day", dt.date(2026, 11, 11)),
    ("Thanksgiving Day", dt.date(2026, 11, 26)),
    ("Christmas Day", dt.date(2026, 12, 25)),
]
NOTICE_DATE = dt.date(2026, 5, 27)


# ---------------------------------------------------------------------------
# Specs
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class LayoutVariant:
    name: str
    row_offset: int = 0
    col_offset: int = 0
    vehicles: int = 1
    alt_headers: bool = False
    sheet_names: tuple[tuple[str, str], ...] = ()
    lookthrough: str | None = None  # a GP-partners block re-allocating the vehicles' GP rows (not additive)
    shared_driver: bool = False  # no fund-level driver row: the first vehicle's driver row is the fund_driver_row
    hidden_merge: bool = False  # Merge tabs hidden once the notices are generated
    combined_prior_label: bool = False  # a prior event headed "Capital Call #2 & Distribution #1 - ..."
    fee_periods: int = 1  # quarters billed by the current event (header "Q2 2026 - Q3 2026 Mgmt Fees")
    summary_sections: bool = False  # the Summary repeats one block per vehicle

    def sheet_name(self, role: str) -> str:
        return dict(self.sheet_names).get(role, _DEFAULT_SHEET_NAMES[role])


_DEFAULT_SHEET_NAMES = {
    "allocation": "Allocation",
    "itd": "ITD Capital Activity",
    "summary": "Summary",
    "mgmt_fee": "Mgmt Fee Calc",
    "investor_data": "DX Investor Data",
    "holiday_calendar": "Bank Holidays",
}

VARIANTS: dict[str, LayoutVariant] = {
    # Mirrors the reference sample's geometry.
    "standard": LayoutVariant(name="standard"),
    # Same model with shifted anchors, renamed sheets, alternative header wording and a
    # Merge tab named after the fund (as in the reference sample).
    "shifted": LayoutVariant(
        name="shifted",
        row_offset=3,
        col_offset=2,
        alt_headers=True,
        sheet_names=(
            ("allocation", "Capital Allocation"),
            ("itd", "ITD"),
            ("summary", "Call Summary"),
            ("mgmt_fee", "MF"),
            ("holiday_calendar", "JPM Holidays"),
        ),
    ),
    # Main fund + parallel vehicle, one Merge tab per vehicle.
    "two_vehicles": LayoutVariant(name="two_vehicles", vehicles=2),
    # Three vehicles plus the GP entity's own partners as a look-through block, no fund-level driver
    # row, hidden Merge tabs (one named without "Merge"), a combined prior-event label, two fee
    # quarters billed at once from two fee-tab columns, and a Summary with one section per vehicle
    # (the pattern of the multi-vehicle reference client).
    "multi_vehicle": LayoutVariant(name="multi_vehicle", vehicles=3, lookthrough="EGF III GP Partners",
                                   shared_driver=True, hidden_merge=True, combined_prior_label=True, fee_periods=2,
                                   summary_sections=True),
}

# Clean verdicts a layout variant changes on its own (the FA-calibrated reading of its structure).
VARIANT_VERDICTS: dict[str, dict[str, str]] = {
    "multi_vehicle": {
        # Hidden Merge tabs hold the notice data: still processed and tied out, but flagged for review.
        "CE-WB-NO-HIDDEN-DATA": "needs_review",
        # With the prior workbook: the rounding residual lands in different vehicle blocks in the two
        # events, so no block carries a plug in both and the pattern cannot be compared.
        "CE-XEV-PLUG-CONSISTENCY": "not_applicable",
    },
}


@dataclass(frozen=True)
class DefectSpec:
    name: str
    description: str
    event_types: tuple[str, ...]
    verdicts: dict[str, str | None]  # overrides applied on top of the clean baseline
    facts: dict[str, dict[str, Any]] = field(default_factory=dict)
    with_prior: bool = False  # build the prior event's workbook alongside (cross-event rules)
    variant: str = "standard"  # the layout variant the defect is seeded on


def _d(name, description, event_types=("capital_call",), verdicts=None, facts=None, with_prior=False,
       variant="standard") -> DefectSpec:
    return DefectSpec(name, description, tuple(event_types), dict(verdicts or {}), dict(facts or {}), with_prior,
                      variant)


DEFECTS: dict[str, DefectSpec] = {
    d.name: d
    for d in [
        _d("hardcoded_lp_cell", "A per-LP investment cell holds a typed number instead of the cascade formula.",
           verdicts={"CE-ALLOC-PER-LP-FORMULAS": "fail"}),
        _d("concentrated_plug", "$500 moved between two LPs via formula offsets (totals still tie).",
           verdicts={"CE-ALLOC-PRO-RATA-PARITY": "fail", "CE-ALLOC-PLUG-DISCIPLINE": "fail"}),
        _d("plug_on_affiliate", "The investment rounding plug sits on the affiliate LP.",
           verdicts={"CE-ALLOC-PLUG-DISCIPLINE": "fail"}),
        _d("short_subtotal_range", "The LP subtotal of Prior Capital Contributions omits the first LP.",
           verdicts={"CE-ALLOC-REFOOT": "fail", "CE-RF-FOOTING": "fail"}),
        _d("commitment_pct_swap", "Two LPs carry typed commitment % values off by +/-0.03 pp.",
           verdicts={"CE-ALLOC-COMMITMENTS": "fail", "CE-ALLOC-PRO-RATA-PARITY": "fail"}),
        _d("live_prior_itd_link", "A prior ITD event cell is a live formula into the Allocation sheet.",
           verdicts={"CE-ITD-PRIOR-FROZEN": "fail"}),
        _d("itd_missing_x", "A current ITD block column has no classification X.",
           verdicts={"CE-TIE-ITD-COMMITMENTS": "fail"},
           facts={"CE-ITD-EVENT-BLOCK": {"unclassified_current_columns": 1}}),
        _d("itd_current_value_mismatch", "One current ITD block cell is typed and $10 off the Allocation sheet.",
           verdicts={"CE-TIE-ITD-ALLOCATION": "fail", "CE-TIE-ITD-COMMITMENTS": "fail",
                     "CE-SUM-CHECKS-ZERO": "fail"}),  # the ITD check row shows the $10
        _d("formula_error", "A #REF! error left on the Allocation sheet.",
           verdicts={"CE-FMT-NO-FORMULA-ERRORS": "fail"}),
        _d("weekend_due_date", "Due date typed as a Saturday instead of the WORKDAY formula.",
           verdicts={"CE-DATE-VALIDITY": "fail", "CE-DATE-ORDER": "needs_review"}),
        _d("due_before_notice", "Due date formula counts business days backwards from the notice date.",
           verdicts={"CE-DATE-ORDER": "fail"}),
        _d("investor_name_mismatch", "DX Investor Data spells one participating LP differently.",
           verdicts={"CE-ID-INVESTOR-KEYS": "fail"}),
        _d("fee_pulled_from_wrong_row", "One LP's Allocation fee pulls another investor's row from the fee tab.",
           verdicts={"CE-TIE-MGMT-FEE": "fail", "CE-TIE-SUPPORT-TABS": "fail", "CE-ALLOC-VEHICLE-TIE": "fail",
                     "CE-ALLOC-GROSS-TIE": "fail", "CE-SUM-CHECKS-ZERO": "fail"}),
        _d("fee_link_shifted", "One LP's fee lookup reads the fee tab's commitment column; its cached value is "
           "still right, as in the reference sample's GP row.",
           facts={"CE-ALLOC-REFERENCE-INTEGRITY": {"period_mismatches": 0, "link_pattern_exceptions": 1}}),
        _d("stale_fee_period", "The fee tab column is still labelled for the prior quarter.",
           verdicts={"CE-TIE-MGMT-FEE": "fail"},
           facts={"CE-DATE-CONSISTENCY": {"fee_period_mismatch": True},
                  "CE-ALLOC-REFERENCE-INTEGRITY": {"period_mismatches": 1}}),
        _d("hidden_populated_row", "A participating LP row on the Allocation sheet is hidden.",
           verdicts={"CE-WB-NO-HIDDEN-DATA": "fail"}),
        _d("normal_view", "The Allocation sheet is saved in Normal view.",
           verdicts={"CE-WB-PAGE-BREAK-VIEW": "fail"}),
        _d("generic_file_name", "The workbook is delivered as Book1.xlsx.",
           verdicts={"CE-WB-FILE-NAMING": "fail"}),
        _d("summary_line_hardcoded", "A Summary component line is typed 5 cents off the Allocation total.",
           verdicts={"CE-SUM-CHECKS-ZERO": "fail"}),
        _d("rf_over_contributed", "An LP's prior contributions exceed its commitment (negative unfunded).",
           verdicts={"CE-RF-FOOTING": "fail"}),
        _d("fmt_serial_date", "The notice date on the Allocation sheet renders as a serial number.",
           verdicts={"CE-FMT-DATE-DISPLAY": "fail"}),
        _d("fmt_general_money", "Per-LP investment amounts use the General number format.",
           verdicts={"CE-FMT-ACCOUNTING": "fail"}),
        _d("rf_current_call_stale", "One LP's Current Capital Call still holds the prior event's amount.",
           verdicts={"CE-RF-CURRENT-CALL-LINK": "fail", "CE-TIE-ITD-COMMITMENTS": "fail"}),
        _d("itd_event_number_gap", "The current event is numbered #5 (consistently on every tab) after #3.",
           verdicts={"CE-ITD-EVENT-SEQUENCE": "fail"},
           facts={"CE-DATE-CONSISTENCY": {"distinct_event_numbers": [5]}}),
        _d("stale_component", "An inactive component column still carries prior-event amounts.",
           verdicts={"CE-ALLOC-VEHICLE-TIE": "fail", "CE-ALLOC-GROSS-TIE": "fail"},
           facts={"CE-ALLOC-STALE-COMPONENTS": {"stale_columns": 1}}),
        _d("placeholder_left", "A Merge tab short name still holds a template token.",
           facts={"CE-WB-NO-PLACEHOLDERS": {"placeholder_count": 1}}),
        _d("merged_component_header", "Two component header cells are merged on the header row.",
           verdicts={"CE-ALLOC-MERGED-CELLS": "fail"}),
        _d("unrounded_amount", "One LP's investment formula omits ROUND(...,2).",
           verdicts={"CE-ALLOC-ROUNDING": "fail"}),
        _d("over_returned_capital", "A never-funded LP receives $5,000 of return of capital.",
           event_types=("distribution",),
           verdicts={"CE-DIST-ROC-LIMIT": "fail", "CE-ALLOC-PRO-RATA-PARITY": "fail",
                     "CE-ALLOC-PLUG-DISCIPLINE": "fail"}),
        _d("carry_split_wrong", "The GP takes 25% of the carry component although the rate cell says 20%.",
           event_types=("distribution",),
           facts={"CE-DIST-CARRY-SPLIT": {"gp_share_matches_rate": False}}),
        _d("zero_live_prior_link", "Prior ITD blocks keep live Allocation links on the $0 GP row.",
           verdicts={"CE-ITD-PRIOR-FROZEN": "needs_review"}),  # FA: a warning when every value is $0
        _d("referenced_hidden_sheet_error", "The Allocation references a hidden legacy sheet that carries #REF!.",
           verdicts={"CE-FMT-NO-FORMULA-ERRORS": "fail"}),
        _d("stale_label_on_active", "An active component column is still headed with the prior event's label.",
           facts={"CE-ALLOC-STALE-COMPONENTS": {"active_with_prior_label": 1}}),
        # --- cross-event defects: built together with the prior event's workbook.
        _d("prior_block_edited", "A frozen prior-event ITD value was edited after the prior event was issued.",
           verdicts={"CE-XEV-HISTORY-UNCHANGED": "fail", "CE-XEV-ROLL-FORWARD": "fail",
                     "CE-XEV-ITD-ROLL-FORWARD": "fail"}, with_prior=True),
        _d("prior_not_most_recent", "The 'prior' workbook uploaded is from two events back (Capital Call #2), so "
           "the ITD balances do not roll forward from it by the current event alone.",
           verdicts={"CE-XEV-ITD-ROLL-FORWARD": "fail", "CE-XEV-ROLL-FORWARD": "fail",
                     "CE-XEV-PLUG-CONSISTENCY": "not_applicable"},  # Capital Call #2 carried no plug
           with_prior=True),
        _d("plug_pattern_changed", "Plugs are spread over the top LPs although the prior event used a single plug.",
           verdicts={"CE-XEV-PLUG-CONSISTENCY": "fail"}, with_prior=True),
        # --- accepted patterns (FA calibration): seen in the reference sample and confirmed
        # as legitimate, so every rule must keep its clean verdict.
        _d("ok_spread_plug", "Two LPs tie for the largest commitment; residual plugs spread over the top three LPs."),
        _d("ok_whole_dollar_fees", "Fees rounded to whole dollars while other components are in cents."),
        _d("ok_itd_overlay_rows", "ITD band carries overlay rows (Mgmt Fees, Late Interest) marking fee columns twice.",
           facts={"CE-ITD-EVENT-BLOCK": {"overlay_rows": 2}}),
        _d("ok_transfer_block", "A non-event 'Transfers' block sits between two ITD events."),
        # FA item 15: a call does not always carry a management fee.
        _d("ok_call_without_fee", "The current call has no management-fee component; the fee tab stays from "
           "prior events.", verdicts={"CE-TIE-MGMT-FEE": "not_applicable", "CE-TIE-SUPPORT-TABS": "not_applicable"}),
        # FA item 15: a distribution may break the carry out on a waterfall support tab.
        _d("ok_waterfall_support_tab", "The Allocation carry column pulls each investor's carry from a "
           "'Distribution Waterfall' support tab with SUMIFS.", event_types=("distribution",),
           verdicts={"CE-TIE-SUPPORT-TABS": "pass"}),
        _d("waterfall_row_not_pulled", "The waterfall spells one LP differently, so the Allocation's SUMIFS "
           "pulls $0 carry for that LP.", event_types=("distribution",),
           verdicts={"CE-TIE-SUPPORT-TABS": "fail", "CE-ALLOC-VEHICLE-TIE": "fail", "CE-ALLOC-GROSS-TIE": "fail",
                     "CE-SUM-CHECKS-ZERO": "fail"}),
        _d("ok_inactive_investor_na", "A transferred-out LP row on the Merge tab shows #N/A.",
           facts={"CE-WB-MERGE-TABS": {"inactive_ignored": 1}}),
        _d("ok_hidden_legacy_errors", "A hidden legacy sheet that nothing references carries #REF! errors."),
        _d("ok_tbd_pending", "A 'TBD' wire date on the portfolio tracker while cash has not moved yet.",
           facts={"CE-WB-NO-PLACEHOLDERS": {"tbd_cells": 1, "tbd_referenced": 0}}),
        _d("ok_stale_label_inactive", "An unused distribution column still carries the prior distribution's label.",
           facts={"CE-ALLOC-STALE-COMPONENTS": {"stale_columns": 0, "active_with_prior_label": 0}}),
        # --- Merge tabs (the notice data) and multi-vehicle / recycling patterns.
        _d("merge_stale_columns", "The Merge tab still pulls the prior event's component columns and omits this "
           "event's expenses column, so every notice is short.",
           verdicts={"CE-TIE-MERGE": "fail", "CE-SUM-CHECKS-ZERO": "fail"}),
        _d("merge_wrong_row_link", "One Merge row's investment cell links to the next investor's Allocation row.",
           verdicts={"CE-TIE-MERGE": "fail", "CE-SUM-CHECKS-ZERO": "fail"}),
        _d("hidden_column_short_subtotal", "A hidden admin column ('Outstanding') holds per-LP amounts and its LP "
           "subtotal SUM range skips the first rows of the block.",
           verdicts={"CE-ALLOC-REFOOT": "fail", "CE-WB-NO-HIDDEN-DATA": "fail"}),
        _d("combined_label_dist_gap", "A prior block is headed 'Capital Call #2 & Distribution #1' and the current "
           "distribution is numbered #3, skipping #2.", event_types=("distribution",),
           verdicts={"CE-ITD-EVENT-SEQUENCE": "fail"},
           facts={"CE-DATE-CONSISTENCY": {"distinct_event_numbers": [3]}}),
        _d("ok_derived_recallable", "The ITD 'Recallable Distributions' cumulative is a recycling-cap formula "
           "(MIN of the cap and the recalled amounts), not a sum of the marked columns.",
           verdicts={"CE-ITD-CUMULATIVE": "needs_review"}),
        _d("multi_vehicle_prior_edited", "On the multi-vehicle workbook, a frozen prior-event value of an Executive "
           "Fund investor was edited after the prior event was issued.",
           verdicts={"CE-XEV-HISTORY-UNCHANGED": "fail", "CE-XEV-ROLL-FORWARD": "fail",
                     "CE-XEV-ITD-ROLL-FORWARD": "fail"}, with_prior=True, variant="multi_vehicle"),
        # --- patterns from the second distribution client (whole-dollar rounding, waiver column, typed carry
        # driver, a long ITD history in collapsed groups, SUMIF pulls, a leftover Merge template, no DX tab,
        # a Summary with one text date, an ITD check column, tax withholding). Each accepted pattern has a
        # counterpart that seeds the real defect in the same shape.
        _d("ok_whole_dollar_spread_plugs", "Distribution components rounded to whole dollars with +/-1 plugs spread "
           "over the largest LPs.", event_types=("distribution",)),
        _d("whole_dollar_plug_on_affiliate", "Whole-dollar rounding with a +/-1 plug on the affiliate LP.",
           event_types=("distribution",), verdicts={"CE-ALLOC-PLUG-DISCIPLINE": "fail"}),
        _d("rf_waiver_adjustment_column", "A 'Waiver Adjustment' column between the roll-forward columns reduces one "
           "LP's remaining commitment on the Allocation and the ITD alike."),
        _d("rf_sum_omits_adjustment", "The remaining-commitment SUM range stops before the waiver column, so the "
           "Allocation remaining ignores the waiver the ITD applies.",
           verdicts={"CE-RF-FOOTING": "fail", "CE-TIE-ITD-COMMITMENTS": "fail"}),
        _d("typed_driver_in_gp_row", "The carry driver cell is blank; the GP row holds the typed carry total and the "
           "per-LP formulas divide that cell.", event_types=("distribution",),
           verdicts={"CE-ALLOC-PER-LP-FORMULAS": "fail"}),
        _d("itd_blocks_in_collapsed_groups", "24 extra prior capital calls sit in collapsed outline groups, one block "
           "has no Total column, and the mapper was given only the current block.",
           verdicts={"CE-WB-NO-HIDDEN-DATA": "fail"}, with_prior=True),
        _d("prior_block_header_live_link", "A prior ITD block's sub-header is a live link to the Allocation header row.",
           verdicts={"CE-ITD-PRIOR-FROZEN": "fail"}),
        _d("itd_block_relabelled", "A prior ITD block was relabelled since the prior workbook (same columns).",
           verdicts={"CE-XEV-HISTORY-UNCHANGED": "needs_review"}, with_prior=True),
        _d("ok_sumif_pulls", "The Merge tab and the current ITD block pull with $-anchored SUMIF lookups keyed by "
           "the row's own investor cell."),
        _d("ok_empty_hidden_merge_tab", "A hidden Merge-like template tab full of #REF! that nothing references."),
        _d("referenced_merge_tab_ref_errors", "A hidden Merge-like tab full of #REF! that the Summary references.",
           verdicts={"CE-FMT-NO-FORMULA-ERRORS": "fail"}),
        _d("ok_no_investor_data_sheet", "The workbook has no DX Investor Data tab.",
           verdicts={"CE-ID-INVESTOR-KEYS": "not_applicable"}),
        _d("ok_summary_single_text_date", "The Summary carries one date, inside its title text."),
        _d("ok_itd_check_column", "The ITD sheet has a 'Contributions Check' column that is zero on every row."),
        _d("itd_check_column_nonzero", "The ITD 'Contributions Check' column omits the current call, so it is non-zero "
           "on every investor row.", verdicts={"CE-SUM-CHECKS-ZERO": "fail"}),
        _d("ok_withholding_in_itd", "A per-investor Tax Withholding column on the Allocation, carried into the "
           "current ITD block.", event_types=("distribution",)),
        _d("withholding_missing_from_itd", "A per-investor Tax Withholding column on the Allocation that the current "
           "ITD block does not carry.", event_types=("distribution",),
           verdicts={"CE-TIE-ITD-ALLOCATION": "fail"}),
    ]
}

WHOLE_DOLLAR_DEFECTS = ("ok_whole_dollar_spread_plugs", "whole_dollar_plug_on_affiliate")
WAIVER_DEFECTS = ("rf_waiver_adjustment_column", "rf_sum_omits_adjustment")
CHECK_COLUMN_DEFECTS = ("ok_itd_check_column", "itd_check_column_nonzero")
WITHHOLDING_DEFECTS = ("ok_withholding_in_itd", "withholding_missing_from_itd")
HIDDEN_MERGE_DEFECTS = ("ok_empty_hidden_merge_tab", "referenced_merge_tab_ref_errors")
EXTRA_CALLS = 24  # itd_blocks_in_collapsed_groups: prior calls #4..#27 in collapsed groups


@dataclass(frozen=True)
class FixtureSpec:
    event_type: str = "capital_call"
    variant: str = "standard"
    defect: str | None = None
    with_prior: bool = False  # also build the prior event's workbook (always on for cross-event defects)

    @property
    def has_prior(self) -> bool:
        return self.with_prior or bool(self.defect and DEFECTS[self.defect].with_prior)

    @property
    def fixture_id(self) -> str:
        parts = [self.event_type, self.variant, self.defect or "clean"]
        if self.with_prior and not (self.defect and DEFECTS[self.defect].with_prior):
            parts.append("with_prior")
        return "__".join(parts)

    def __post_init__(self) -> None:
        if self.event_type not in EVENT_TYPES:
            raise ValueError(f"unknown event type {self.event_type!r}")
        if self.variant not in VARIANTS:
            raise ValueError(f"unknown layout variant {self.variant!r}")
        if self.defect is not None:
            defect = DEFECTS[self.defect]
            if self.event_type not in defect.event_types:
                raise ValueError(f"defect {self.defect!r} does not apply to {self.event_type!r}")
        if self.has_prior and (self.event_type != "capital_call" or self.variant not in ("standard", "two_vehicles",
                                                                                           "multi_vehicle")):
            raise ValueError("prior-event workbooks are generated for capital-call fixtures only")


@dataclass
class FixtureManifest:
    spec: FixtureSpec
    path: Path
    sheet_roles: dict[str, str]
    relevant_sheets: list[str]
    layouts: dict[str, dict[str, Any]]
    truth: dict[str, Any]
    expected_verdicts: dict[str, str | None]
    facts: dict[str, dict[str, Any]]
    reference_sheets: list[str] = field(default_factory=list)  # other sheets the relevant ones reference
    prior: "FixtureManifest | None" = None  # the prior event's workbook, when generated

    @property
    def options(self) -> dict[str, Any]:
        return {"event_type": self.spec.event_type}

    def to_json(self) -> dict[str, Any]:
        return {
            "fixture_id": self.spec.fixture_id,
            "event_type": self.spec.event_type,
            "variant": self.spec.variant,
            "defect": self.spec.defect,
            "file_name": self.path.name,
            "sheet_roles": self.sheet_roles,
            "relevant_sheets": self.relevant_sheets,
            "layouts": self.layouts,
            "truth": self.truth,
            "expected_verdicts": self.expected_verdicts,
            "facts": self.facts,
            "reference_sheets": self.reference_sheets,
            "prior": self.prior.to_json() if self.prior else None,
        }


def clean_verdicts(event_type: str, with_prior: bool = False, variant: str = "standard") -> dict[str, str | None]:
    verdicts: dict[str, str | None] = {rule_id: "pass" for rule_id in DETERMINISTIC_RULE_IDS}
    verdicts.update(VARIANT_VERDICTS.get(variant, {}))
    if event_type == "capital_call":
        verdicts["CE-DIST-ROC-LIMIT"] = "not_applicable"
    if event_type == "distribution":
        verdicts["CE-TIE-MGMT-FEE"] = "not_applicable"
        verdicts["CE-TIE-SUPPORT-TABS"] = "not_applicable"  # no fee tab; see ok_waterfall_support_tab
    if not with_prior:
        # The prior event's workbook was not supplied (and the run is not marked as the first event).
        for rule_id in CROSS_EVENT_RULE_IDS:
            verdicts[rule_id] = "needs_review"
    return verdicts


def default_specs() -> list[FixtureSpec]:
    """Clean fixtures for every event type x layout variant, plus every defect."""
    specs = [FixtureSpec(event_type=e, variant=v) for e in EVENT_TYPES for v in VARIANTS]
    specs.append(FixtureSpec(with_prior=True))
    specs.append(FixtureSpec(variant="multi_vehicle", with_prior=True))
    for defect in DEFECTS.values():
        specs.append(FixtureSpec(event_type=defect.event_types[0], variant=defect.variant, defect=defect.name))
    return specs


# ---------------------------------------------------------------------------
# Domain model
# ---------------------------------------------------------------------------


@dataclass
class Investor:
    name: str
    vehicle: str
    commitment: Decimal
    affiliate: bool = False
    is_gp: bool = False
    investor_id: int = 0
    fund_id: int = 0
    late_closer: bool = False  # admitted after the last prior call; nothing contributed yet


@dataclass(frozen=True)
class ComponentType:
    key: str
    header: str
    component_type: str  # investment | org_expense | mgmt_fee | placement_fee | return_of_capital | realized_gain | carry | late_interest
    side: str  # call | distribution
    classification: str  # ITD band category


C_INVESTMENT = ComponentType("investment", "Investment", "investment", "call", "investment_contributions")
C_EXPENSES = ComponentType("expenses", "Partnership Expenses / Org Costs", "org_expense", "call", "cost_contributions")
C_FEE = ComponentType("mgmt_fee", "Mgmt Fees", "mgmt_fee", "call", "cost_contributions")
C_PLACEMENT = ComponentType("placement", "Placement Fees", "placement_fee", "call", "cost_contributions")
C_ROC = ComponentType("roc", "Return of Capital", "return_of_capital", "distribution", "non_recallable_distributions")
C_GAIN = ComponentType("gain", "Realized Gain", "realized_gain", "distribution", "non_recallable_distributions")
C_CARRY = ComponentType("carry", "Realized Carry", "carry", "distribution", "non_recallable_distributions")

C_WITHHOLDING = ComponentType("withholding", "Tax Withholding", "tax_withholding", "distribution", "tax_withholding")

WATERFALL_SHEET = "Distribution Waterfall"

CALL_COMPONENTS = (C_INVESTMENT, C_EXPENSES, C_FEE, C_PLACEMENT)
DIST_COMPONENTS = (C_ROC, C_GAIN, C_CARRY)


@dataclass
class EventDef:
    kind: str  # capital_call | distribution | net_event
    number: int
    date: dt.date  # due / payment date shown in the ITD header
    drivers: dict[str, Decimal]  # component key -> fund-level amount (distributions negative)
    fee_period: str | None = None
    label: str | None = None  # ITD header text when it is not "<word> #<n> - <date>" (combined events)

    @property
    def word(self) -> str:
        return {"capital_call": "Capital Call", "distribution": "Distribution", "net_event": "Net Capital Call"}[self.kind]

    @property
    def short_label(self) -> str:
        return f"{self.word} #{self.number}"

    @property
    def itd_label(self) -> str:
        return self.label or f"{self.short_label} - {self.date:%m.%d.%Y}"

    def header_for(self, comp: ComponentType) -> str:
        if comp.key == "mgmt_fee" and self.fee_period:
            return f"{self.fee_period} Mgmt Fees"
        return comp.header


def _investors(variant: LayoutVariant, defect: str | None) -> list[Investor]:
    main = [
        ("Northgate Family Trust", "12500000"),
        ("Crescent Ridge Investments LP", "10000000"),
        ("Harborview Capital Partners, LLC", "25000000"),
        ("Alder & Finch Holdings, LLC", "7750000"),
        ("Meridian Endowment Fund", "6300000"),
        ("Juniper Hollow Partners", "5000000"),
        ("Pemberton Street FLP", "3333333"),
        ("Oakmere Ventures, Ltd.", "2750000"),
        ("Silverline Retirement Plan", "2000000"),
        ("Tamsin Rourke", "1250000"),
        ("Wexford Lane Capital, LLC", "900000"),
    ]
    if defect == "ok_spread_plug":
        main[0] = ("Northgate Family Trust", "25000000")
    investors = [Investor(n, "Main Fund", Decimal(c)) for n, c in main]
    investors.append(Investor("Brightwater Affiliates Fund, LP", "Main Fund", Decimal("4000000"), affiliate=True))
    investors.append(Investor("Vireo Late Close Partners", "Main Fund", Decimal("1500000"), late_closer=True))
    investors.append(Investor("Example Growth Fund III GP, LLC", "Main Fund", Decimal("0"), is_gp=True))
    if variant.vehicles >= 2:
        for n, c in [
            ("Kestrel Offshore Feeder, Ltd.", "9000000"),
            ("Lindqvist Pension Stiftung", "6500000"),
            ("Quarry Road Partners", "2250000"),
            ("Ottoline Varga", "750000"),
        ]:
            investors.append(Investor(n, "Parallel Fund", Decimal(c)))
        investors.append(Investor("EGF III Parallel GP, LLC", "Parallel Fund", Decimal("0"), is_gp=True))
    if variant.vehicles >= 3:
        for n, c in [
            ("Halloran Executive Partners", "1500000"),
            ("Ines Marchetti", "750000"),
            ("Rowan & Sable Capital", "2250000"),
        ]:
            investors.append(Investor(n, "Executive Fund", Decimal(c)))
        investors.append(Investor("EGF III Executive GP, LLC", "Executive Fund", Decimal("0"), is_gp=True))
    if variant.lookthrough:
        # The GP entity's own partners: their block re-allocates the GP rows of the vehicles above.
        # Their commitments add up to the GP entity's commitments across the vehicles (3 x 2,000,000).
        for n, c in [("Brian Castleberry", "1890000"), ("Marta Quill", "1500000"), ("Devin Oyelaran", "2610000")]:
            investors.append(Investor(n, variant.lookthrough, Decimal(c)))
    if variant.lookthrough:
        for inv in investors:
            if inv.is_gp:
                inv.commitment = Decimal("2000000")  # the GP entity invests alongside the LPs
    fund_ids = {"Main Fund": 901, "Parallel Fund": 902, "Executive Fund": 903}
    next_id = 20001
    for inv in investors:
        inv.fund_id = fund_ids.get(inv.vehicle, 904)
        if not inv.is_gp:
            inv.investor_id = next_id
            next_id += 1
    return investors


def _events(event_type: str, defect: str | None, variant: LayoutVariant | None = None) -> tuple[list[EventDef], EventDef]:
    prior = [
        EventDef("capital_call", 1, dt.date(2024, 1, 22),
                 {"investment": Decimal("20000000"), "expenses": Decimal("400000"), "mgmt_fee": Decimal("0")}, "Q1 2024"),
        EventDef("capital_call", 2, dt.date(2024, 6, 25),
                 {"investment": Decimal("10000000"), "mgmt_fee": Decimal("0")}, "Q2 2024"),
        EventDef("distribution", 1, dt.date(2025, 1, 17), {"roc": Decimal("-1500000")}),
        EventDef("capital_call", 3, dt.date(2026, 3, 9),
                 {"investment": Decimal("5000000"), "expenses": Decimal("250000"), "mgmt_fee": Decimal("0")}, "Q1 2026"),
    ]
    if (variant and variant.combined_prior_label) or defect == "combined_label_dist_gap":
        # Capital Call #2 and Distribution #1 issued together under one ITD header.
        prior[1:3] = [EventDef("net_event", 2, dt.date(2024, 6, 25),
                               {"investment": Decimal("10000000"), "mgmt_fee": Decimal("0"), "roc": Decimal("-1500000")},
                               "Q2 2024", label="Capital Call #2 & Distribution #1 - 06.25.2024")]
    due = dt.date(2026, 6, 10)
    if event_type == "capital_call":
        current = EventDef("capital_call", 4, due,
                           {"investment": Decimal("8500000"), "expenses": Decimal("275000"), "mgmt_fee": Decimal("0")},
                           "Q3 2026")
    elif event_type == "distribution":
        current = EventDef("distribution", 2, due,
                           {"roc": Decimal("-3000000"), "gain": Decimal("-1200000"), "carry": Decimal("-400000")})
    else:
        current = EventDef("net_event", 4, due,
                           {"expenses": Decimal("150000"), "mgmt_fee": Decimal("0"), "roc": Decimal("-2000000")},
                           "Q3 2026")
    if defect == "itd_blocks_in_collapsed_groups":
        # A long history: calls #4..#27 after Capital Call #3, dated a day apart, each a small investment.
        for index in range(EXTRA_CALLS):
            prior.append(EventDef("capital_call", 4 + index, dt.date(2026, 3, 10) + dt.timedelta(days=index),
                                  {"investment": Decimal("100000")}))
        current.number = 4 + EXTRA_CALLS
    if defect == "typed_driver_in_gp_row":
        current.drivers["carry"] = Decimal("0")  # the driver cell stays blank; the GP row carries the total
    if defect == "itd_event_number_gap":
        current.number = 5
    if defect == "combined_label_dist_gap":
        current.number = 3  # the combined header carried Distribution #1, so #2 is skipped
    if defect == "ok_call_without_fee":
        del current.drivers["mgmt_fee"]
        current.fee_period = None
    if variant and variant.fee_periods == 2 and "mgmt_fee" in current.drivers:
        current.fee_period = "Q2 2026 - Q3 2026"
    return prior, current


def r2(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def r0(value: Decimal) -> Decimal:
    return value.quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def _serial(value: dt.date) -> int:
    return (value - EXCEL_EPOCH).days


def _workday(start: dt.date, days: int, holidays: set[dt.date]) -> dt.date:
    step = 1 if days >= 0 else -1
    remaining = abs(days)
    current = start
    while remaining:
        current += dt.timedelta(days=step)
        if current.weekday() < 5 and current not in holidays:
            remaining -= 1
    return current


# ---------------------------------------------------------------------------
# Allocation math
# ---------------------------------------------------------------------------


@dataclass
class Allocation:
    """Per-investor amounts for one event's components, with the plug bookkeeping."""

    amounts: dict[str, dict[str, Decimal]]  # component key -> investor name -> amount
    offsets: dict[str, dict[str, Decimal]]  # component key -> investor name -> literal offset in formula
    unrounded: set[tuple[str, str]] = field(default_factory=set)  # (component, investor)


def _fee(inv: Investor, affiliate_override: bool | None = None, whole_dollars: bool = False) -> Decimal:
    affiliate = inv.affiliate if affiliate_override is None else affiliate_override
    if inv.is_gp or affiliate:
        return Decimal("0")
    raw = inv.commitment * MGMT_FEE_RATE * FEE_PERIOD_FRACTION
    return r0(raw) if whole_dollars else r2(raw)


def _vehicle_split(total: Decimal, investors: list[Investor], vehicles: list[str]) -> dict[str, Decimal]:
    """Split a fund-level driver across vehicles by commitment; the first vehicle takes the residual."""
    if len(vehicles) == 1:
        return {vehicles[0]: total}
    fund_commitment = sum(i.commitment for i in investors)
    split = {v: r2(total * sum(i.commitment for i in investors if i.vehicle == v) / fund_commitment) for v in vehicles}
    split[vehicles[0]] += total - sum(split.values())
    return split


def _plug_target(investors: list[Investor]) -> Investor:
    eligible = [i for i in investors if not i.is_gp and not i.affiliate]
    return max(eligible, key=lambda i: i.commitment)  # max() keeps the topmost row on ties


def _allocate_pro_rata(
    driver: Decimal,
    investors: list[Investor],
    weights: dict[str, Decimal],
    plug: Investor | None,
    unrounded: set[str] = frozenset(),
    whole_dollars: bool = False,
) -> tuple[dict[str, Decimal], dict[str, Decimal]]:
    total_weight = sum(weights.values())
    amounts: dict[str, Decimal] = {}
    offsets: dict[str, Decimal] = {}
    for inv in investors:
        share = weights.get(inv.name, Decimal("0")) / total_weight if total_weight else Decimal("0")
        raw = driver * share
        amounts[inv.name] = raw if inv.name in unrounded else (r0(raw) if whole_dollars else r2(raw))
    residual = driver - sum(amounts.values())
    if plug is not None and residual:
        amounts[plug.name] += residual
        offsets[plug.name] = residual
    return amounts, offsets


# ---------------------------------------------------------------------------
# Sheet builder
# ---------------------------------------------------------------------------


@dataclass
class CellSpec:
    value: Any = None  # literal (number, str, date) or None
    formula: str | None = None  # without leading '='
    cached: Any = None  # cached value for formula cells
    fmt: str | None = None


class SheetBuilder:
    def __init__(self, name: str, role: str, row_offset: int = 0, col_offset: int = 0) -> None:
        self.name = name
        self.role = role
        self.row_offset = row_offset
        self.col_offset = col_offset
        self.cells: dict[str, CellSpec] = {}
        self.merges: list[str] = []
        self.hidden_rows: set[int] = set()
        self.hidden_cols: set[str] = set()
        self.outlined_cols: set[str] = set()  # columns in an outline group (collapsed when also hidden)
        self.view: str | None = None
        self.state = "visible"

    def row(self, r: int) -> int:
        return r + self.row_offset

    def col(self, c: int) -> str:
        return get_column_letter(c + self.col_offset)

    def coord(self, r: int, c: int) -> str:
        return f"{self.col(c)}{self.row(r)}"

    def put(self, r: int, c: int, value: Any = None, *, formula: str | None = None, cached: Any = None,
            fmt: str | None = None) -> str:
        coord = self.coord(r, c)
        self.put_at(coord, value, formula=formula, cached=cached, fmt=fmt)
        return coord

    def put_at(self, coord: str, value: Any = None, *, formula: str | None = None, cached: Any = None,
               fmt: str | None = None) -> str:
        self.cells[coord] = CellSpec(value=value, formula=formula, cached=cached, fmt=fmt)
        return coord

    def get(self, coord: str) -> CellSpec:
        return self.cells[coord]


def q(sheet: str) -> str:
    return sheet if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", sheet) else "'" + sheet.replace("'", "''") + "'"


def _num(value: Decimal | int | float) -> float | int:
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


# ---------------------------------------------------------------------------
# Fixture assembly
# ---------------------------------------------------------------------------


class _Builder:
    def __init__(self, spec: FixtureSpec, stage: str = "current") -> None:
        """``stage="prior"`` builds the prior event's workbook (Capital Call #3) for ``spec``."""
        self.spec = spec
        self.stage = stage
        self.variant = VARIANTS[spec.variant]
        self.defect = spec.defect if stage == "current" else None
        self.investors = _investors(self.variant, self.defect)
        # A defect that shapes the fund's history applies to the prior workbook too.
        history_defect = spec.defect if spec.defect == "itd_blocks_in_collapsed_groups" else self.defect
        self.prior_events, self.current = _events(spec.event_type, history_defect, self.variant)
        self.holidays = {d for _, d in HOLIDAYS_2026}
        self.notice_date = NOTICE_DATE
        if stage == "prior":
            # The prior event is the last prior capital call; late closers had not closed yet.
            self.investors = [i for i in self.investors if not i.late_closer]
            last = -1
            if spec.defect == "prior_not_most_recent":
                last = max(i for i, e in enumerate(self.prior_events[:-1]) if e.kind == "capital_call") - len(
                    self.prior_events)
            self.current = self.prior_events[last]
            self.prior_events = self.prior_events[:last]
            self.notice_date = _workday(self.current.date, -10, self.holidays)
            assert _workday(self.notice_date, 10, self.holidays) == self.current.date
        self.vehicles = list(dict.fromkeys(i.vehicle for i in self.investors))
        self.lookthrough = self.variant.lookthrough
        self.lp_vehicles = [v for v in self.vehicles if v != self.lookthrough]
        # Billing two quarters at once is a feature of the current event; the prior event billed one.
        self.fee_periods = self.variant.fee_periods if stage == "current" else 1
        self.fee_quarters = ["Q2 2026", "Q3 2026"] if self.fee_periods == 2 else None
        self.sheets: list[SheetBuilder] = []
        self.layouts: dict[str, dict[str, Any]] = {}
        self.truth: dict[str, Any] = {}
        self.facts: dict[str, dict[str, Any]] = {}
        self.names = {role: self.variant.sheet_name(role) for role in _DEFAULT_SHEET_NAMES}
        if self.variant.vehicles == 1:
            self.merge_names = {"Main Fund": "Example Growth Fund III" if self.variant.name == "shifted" else "Merge"}
        else:
            self.merge_names = {v: f"Merge - {v}" for v in self.vehicles}
            if self.variant.name == "multi_vehicle":
                # As in the reference client: one tab named after the vehicle without "Merge".
                self.merge_names["Executive Fund"] = "EGF III Executive"
        self.whole_dollar_fees = self.defect == "ok_whole_dollar_fees"
        # FA item 15: the carry is broken out on a waterfall support tab the Allocation pulls from.
        self.waterfall = self.defect in ("ok_waterfall_support_tab", "waterfall_row_not_pulled")
        self.waterfall_rows: list[dict[str, Any]] = []

    # -- helpers -----------------------------------------------------------------

    def lps(self, vehicle: str | None = None) -> list[Investor]:
        return [i for i in self.investors if not i.is_gp and (vehicle is None or i.vehicle == vehicle)]

    def in_vehicle(self, vehicle: str) -> list[Investor]:
        return [i for i in self.investors if i.vehicle == vehicle]

    def by_name(self, name: str) -> Investor:
        return next(i for i in self.investors if i.name == name)

    def fee_for(self, inv: Investor) -> Decimal:
        if inv.vehicle == self.lookthrough:
            return Decimal("0")  # the GP's partners pay no fee; their block re-allocates the GP rows ($0 fee)
        return _fee(inv, whole_dollars=self.whole_dollar_fees) * self.fee_periods

    def is_lookthrough(self, name: str) -> bool:
        return self.lookthrough is not None and self.by_name(name).vehicle == self.lookthrough

    # -- event allocations ----------------------------------------------------------

    def allocate_event(self, event: EventDef, contributed: dict[str, Decimal], *, current: bool) -> Allocation:
        """Allocate every component of ``event``.

        Late closers were admitted after the prior events, so they are excluded from
        (and carry $0 on) every prior event.
        """
        amounts: dict[str, dict[str, Decimal]] = {}
        offsets: dict[str, dict[str, Decimal]] = {}
        unrounded: set[tuple[str, str]] = set()

        def eligible(inv: Investor) -> bool:
            return current or not inv.late_closer

        for key, fund_driver in event.drivers.items():
            amounts[key] = {i.name: Decimal("0") for i in self.investors}
            offsets[key] = {}
            if key == "mgmt_fee":
                for inv in self.investors:
                    if eligible(inv) and inv.vehicle != self.lookthrough:
                        amounts[key][inv.name] = self.fee_for(inv) if current else _fee(inv)
                if current and self.defect == "fee_pulled_from_wrong_row":
                    amounts[key]["Juniper Hollow Partners"] = self.fee_for(self.by_name("Meridian Endowment Fund"))
                continue
            lp_investors = [i for i in self.investors if i.vehicle != self.lookthrough and eligible(i)]
            split = _vehicle_split(fund_driver, lp_investors, self.lp_vehicles)
            for vehicle in self.lp_vehicles:
                members = [i for i in self.in_vehicle(vehicle) if eligible(i)]
                driver = split[vehicle]
                plug = _plug_target(members)
                if key == "carry" and current and self.defect == "typed_driver_in_gp_row":
                    # The GP's carry (typed on its row) is taken from the LPs pro rata: the column nets to zero.
                    gp = next(i for i in members if i.is_gp)
                    lps = [i for i in members if not i.is_gp]
                    weights = {i.name: contributed.get(i.name, Decimal("0")) for i in lps}
                    lp_amounts, lp_offsets = _allocate_pro_rata(TYPED_CARRY, lps, weights, plug)
                    amounts[key].update(lp_amounts)
                    amounts[key][gp.name] = -TYPED_CARRY
                    offsets[key].update(lp_offsets)
                    continue
                if key == "carry":
                    gp = next(i for i in members if i.is_gp)
                    rate = self.gp_carry_rate if current else CARRY_RATE
                    gp_share = r2(driver * rate)
                    lps = [i for i in members if not i.is_gp]
                    weights = {i.name: contributed.get(i.name, Decimal("0")) for i in lps}
                    lp_amounts, lp_offsets = _allocate_pro_rata(driver - gp_share, lps, weights, plug)
                    amounts[key].update(lp_amounts)
                    amounts[key][gp.name] = gp_share
                    offsets[key].update(lp_offsets)
                    continue
                if C_BY_KEY[key].side == "distribution":
                    weights = {i.name: contributed.get(i.name, Decimal("0")) for i in members}
                else:
                    weights = {i.name: i.commitment for i in members}
                main = vehicle == "Main Fund"
                if current and main and key == "investment" and self.defect == "commitment_pct_swap":
                    weights = self._swapped_pct_weights(members)
                if current and main and key == "investment" and self.defect == "plug_on_affiliate":
                    plug = next(i for i in members if i.affiliate)
                skip_round: set[str] = set()
                if current and main and key == "investment" and self.defect == "unrounded_amount":
                    skip_round = {"Oakmere Ventures, Ltd."}
                    unrounded.add((key, "Oakmere Ventures, Ltd."))
                spread = current and main and key in ("investment", "expenses") \
                    and self.defect in ("ok_spread_plug", "plug_pattern_changed")
                whole = current and key in ("roc", "gain") and self.defect in WHOLE_DOLLAR_DEFECTS
                vehicle_amounts, vehicle_offsets = _allocate_pro_rata(
                    driver, members, weights, None if (spread or whole) else plug, skip_round, whole_dollars=whole)
                if spread:
                    vehicle_offsets = _spread_residual(driver, vehicle_amounts, members)
                if whole:
                    vehicle_offsets = _spread_units(driver, vehicle_amounts, members,
                                                    on_affiliate=self.defect == "whole_dollar_plug_on_affiliate")
                amounts[key].update(vehicle_amounts)
                offsets[key].update(vehicle_offsets)
            if self.lookthrough:
                # The GP entity's share of every vehicle, re-allocated to its own partners by commitment.
                partners = [i for i in self.in_vehicle(self.lookthrough) if eligible(i)]
                gp_total = sum((amounts[key][i.name] for i in self.investors if i.is_gp), Decimal("0"))
                if C_BY_KEY[key].side == "distribution":
                    lt_weights = {i.name: contributed.get(i.name, Decimal("0")) for i in partners}
                else:
                    lt_weights = {i.name: i.commitment for i in partners}
                lt_amounts, lt_offsets = _allocate_pro_rata(gp_total, partners, lt_weights,
                                                            _plug_target(partners) if partners else None)
                amounts[key].update(lt_amounts)
                offsets[key].update(lt_offsets)
            if current and key == "investment" and self.defect == "concentrated_plug":
                self._move(amounts, offsets, key, "Juniper Hollow Partners", "Silverline Retirement Plan", Decimal("500"))
            if current and key == "carry" and self.defect == "waterfall_row_not_pulled":
                # The waterfall spells the investor differently, so the Allocation's SUMIFS finds nothing.
                self.unpulled_carry = amounts[key]["Juniper Hollow Partners"]
                amounts[key]["Juniper Hollow Partners"] = Decimal("0")
            if current and key == "roc" and self.defect == "over_returned_capital":
                self._move(amounts, offsets, key, "Harborview Capital Partners, LLC", "Vireo Late Close Partners",
                           Decimal("-5000"))
        return Allocation(amounts, offsets, unrounded)

    @property
    def gp_carry_rate(self) -> Decimal:
        return Decimal("0.25") if self.defect == "carry_split_wrong" else CARRY_RATE

    def _swapped_pct_weights(self, members: list[Investor]) -> dict[str, Decimal]:
        total = sum(i.commitment for i in members)
        weights = {i.name: i.commitment / total for i in members}
        weights["Tamsin Rourke"] += Decimal("0.0003")
        weights["Wexford Lane Capital, LLC"] -= Decimal("0.0003")
        self.pct_overrides = {n: weights[n] for n in ("Tamsin Rourke", "Wexford Lane Capital, LLC")}
        return weights

    @staticmethod
    def _move(amounts, offsets, key, from_name, to_name, amount: Decimal) -> None:
        amounts[key][from_name] -= amount
        amounts[key][to_name] += amount
        offsets[key][from_name] = offsets[key].get(from_name, Decimal("0")) - amount
        offsets[key][to_name] = offsets[key].get(to_name, Decimal("0")) + amount

    # -- build ------------------------------------------------------------------------

    def build(self, out_dir: Path) -> FixtureManifest:
        self.pct_overrides: dict[str, Decimal] = {}
        # Prior events: contributions accumulate event by event.
        contributed: dict[str, Decimal] = {i.name: Decimal("0") for i in self.investors}
        distributed: dict[str, Decimal] = {i.name: Decimal("0") for i in self.investors}
        self.prior_allocations: list[Allocation] = []
        for event in self.prior_events:
            members_before = {n: v for n, v in contributed.items()}
            alloc = self.allocate_event(event, members_before, current=False)
            if self.defect == "rf_over_contributed" and event.number == 1 and event.kind == "capital_call":
                alloc.amounts["investment"]["Tamsin Rourke"] += Decimal("800000")
            self.prior_allocations.append(alloc)
            for key, per_inv in alloc.amounts.items():
                for name, amount in per_inv.items():
                    if C_BY_KEY[key].side == "call":
                        contributed[name] += amount
                    else:
                        distributed[name] += amount
        if self.defect in ("prior_block_edited", "multi_vehicle_prior_edited"):
            # Edited after the prior events were issued: one frozen cell, nothing re-derived from it.
            edited = next(a for e, a in zip(self.prior_events, self.prior_allocations)
                          if e.kind in ("capital_call", "net_event") and e.number == 2)
            victim = "Northgate Family Trust" if self.defect == "prior_block_edited" else "Ines Marchetti"
            edited.amounts["investment"][victim] += Decimal("1000")
            contributed[victim] += Decimal("1000")
        self.prior_contributed = contributed
        self.prior_distributed = distributed
        self.current_alloc = self.allocate_event(self.current, contributed, current=True)
        self.withholding = self.defect in WITHHOLDING_DEFECTS
        if self.withholding:
            # Two foreign LPs: 10% of their realized gain is withheld (positive: it reduces the cash paid).
            foreign = ("Crescent Ridge Investments LP", "Juniper Hollow Partners")
            gains = self.current_alloc.amounts.get("gain", {})
            self.current_alloc.amounts["withholding"] = {
                i.name: (r2(abs(gains.get(i.name, Decimal("0"))) * Decimal("0.1")) if i.name in foreign else Decimal("0"))
                for i in self.investors}

        self._build_holidays()
        self._build_allocation_plan()
        self._build_itd()
        self._build_allocation()
        self._build_summary()
        self._build_merge_tabs()
        if self.defect != "ok_no_investor_data_sheet":
            self._build_investor_data()
        self._build_fee_tab()
        self._build_waterfall()
        self._build_other_sheets()
        self._apply_cell_defects()

        order = [self.names["holiday_calendar"], self.names["summary"], self.names["allocation"], self.names["itd"],
                 *self.merge_names.values(), self.names["investor_data"], "Notes", "Portfolio Investment Tracker",
                 self.names["mgmt_fee"], *([WATERFALL_SHEET] if self.waterfall else []), "3rd Close Rebalance",
                 *(["Old Merge"] if self.defect in HIDDEN_MERGE_DEFECTS else [])]
        by_name = {s.name: s for s in self.sheets}
        ordered = [by_name[n] for n in order if n in by_name]

        file_name = self._file_name()
        target_dir = out_dir / self.spec.fixture_id if self.stage == "current" else out_dir
        target_dir.mkdir(parents=True, exist_ok=True)
        path = target_dir / file_name
        _write_workbook(ordered, path)

        sheet_roles = {s.name: s.role for s in ordered}
        relevant = [s.name for s in ordered if s.role in RELEVANT_ROLES[self.spec.event_type]]
        references = _referenced_sheets(ordered, relevant)
        verdicts = clean_verdicts(self.spec.event_type, with_prior=self.spec.has_prior, variant=self.spec.variant)
        facts = _clean_facts(self.spec.event_type)
        if self.defect:
            verdicts.update(DEFECTS[self.defect].verdicts)
            for rule_id, extra in DEFECTS[self.defect].facts.items():
                facts.setdefault(rule_id, {}).update(extra)
        return FixtureManifest(
            spec=self.spec,
            path=path,
            sheet_roles=sheet_roles,
            relevant_sheets=relevant,
            layouts={name: layout for name, layout in self.layouts.items()},
            truth=_jsonable(self.truth),
            expected_verdicts=verdicts,
            facts=facts,
            reference_sheets=references,
            prior=(_Builder(self.spec, stage="prior").build(target_dir / "prior")
                   if self.stage == "current" and self.spec.has_prior else None),
        )

    def _file_name(self) -> str:
        if self.defect == "generic_file_name":
            return "Book1.xlsx"
        slug = re.sub(r"[^a-z0-9]+", "-", self.current.short_label.lower()).strip("-")
        return f"{FUND_SHORT}_{self.current.number:03d}_{slug}_allocation_summary.xlsx"

    # -- Bank holidays --------------------------------------------------------------------

    def _build_holidays(self) -> None:
        sheet = SheetBuilder(self.names["holiday_calendar"], "holiday_calendar")
        sheet.put(1, 1, "Holiday")
        sheet.put(1, 2, "Date")
        for idx, (label, day) in enumerate(HOLIDAYS_2026, start=2):
            sheet.put(idx, 1, label)
            sheet.put(idx, 2, day, fmt=DATE_SHORT_FMT)
        self.holiday_range = f"{q(sheet.name)}!$B$2:$B${len(HOLIDAYS_2026) + 1}"
        self.sheets.append(sheet)
        self.layouts[sheet.name] = {
            "role": "holiday_calendar",
            "sheet": sheet.name,
            "header_row": 1,
            "date_column": "B",
            "first_row": 2,
            "last_row": len(HOLIDAYS_2026) + 1,
        }

    # -- Allocation geometry (planned first: the ITD sheet references it) ---------------

    def _build_allocation_plan(self) -> None:
        v = self.variant
        sheet = SheetBuilder(self.names["allocation"], "allocation", v.row_offset, v.col_offset)
        self.alloc_sheet = sheet
        event = self.current
        has_dist_basis = event.kind in ("distribution", "net_event")
        cols: dict[str, int] = {}
        c = 2
        for key in ("close", "investor", "affiliate", "commitment", "commitment_pct"):
            cols[key] = c
            c += 1
        if has_dist_basis:
            cols["contributed"] = c
            cols["contributed_pct"] = c + 1
            c += 2
        c += 1
        self.call_comp_cols: dict[str, int] = {}
        for comp in CALL_COMPONENTS:
            self.call_comp_cols[comp.key] = c
            c += 1
        cols["call_total"] = c
        c += 1
        cols["stale"] = c  # spare column right of the call total (used by the stale_component defect)
        c += 2
        self.dist_comp_cols: dict[str, int] = {}
        for comp in DIST_COMPONENTS:
            self.dist_comp_cols[comp.key] = c
            c += 1
        cols["dist_total"] = c
        c += 2
        if self.defect in WITHHOLDING_DEFECTS:
            cols["withholding"] = c - 1  # right after the distribution total, as on the reference client
        cols["late_interest"] = c
        c += 1
        cols["cash_due"] = c
        c += 2
        rf_keys = ["rf_commitment", "rf_prior", "rf_prior_recallable", "rf_current_call", "rf_current_recallable"]
        if self.defect in WAIVER_DEFECTS:
            rf_keys.append("rf_waiver")
        rf_keys.append("rf_remaining")
        for key in rf_keys:
            cols[key] = c
            c += 1
        c += 1
        for key in ("prior_gross_dist", "current_dist", "total_gross_dist"):
            cols[key] = c
            c += 1
        c += 1
        cols["received"] = c
        cols["received_date"] = c + 1
        cols["error_cell"] = c + 3
        self.acols = cols
        self.comp_col = {**self.call_comp_cols, **self.dist_comp_cols}
        if "withholding" in cols:
            self.comp_col["withholding"] = cols["withholding"]

        # Row plan
        rows: dict[str, Any] = {"fund_name": 2, "notice": 3, "due": 4, "driver": 5, "header": 6}
        r = 7
        blocks = []
        for index, vehicle in enumerate(self.vehicles):
            block: dict[str, Any] = {"name": vehicle}
            if len(self.vehicles) > 1 and not (v.shared_driver and index == 0):
                block["title_row"] = r
                block["driver_row"] = r + 1
                r += 2
            else:
                block["driver_row"] = rows["driver"]  # the row above the header
            block["investor_rows"] = {}
            for inv in self.in_vehicle(vehicle):
                if inv.is_gp:
                    continue
                block["investor_rows"][inv.name] = r
                r += 1
            r += 1
            block["lp_subtotal"] = r
            r += 2
            gp = next((i for i in self.in_vehicle(vehicle) if i.is_gp), None)
            if gp is not None:
                block["gp_row"] = r
                block["gp_name"] = gp.name
                r += 1
                block["gp_subtotal"] = r
                r += 2
            else:  # the look-through block has no general partner row
                block["gp_row"] = block["gp_name"] = block["gp_subtotal"] = None
            block["total_row"] = r
            r += 2 if len(self.vehicles) > 1 else 1
            blocks.append(block)
        if len(self.vehicles) > 1:
            rows["grand_total"] = r
            r += 1
        else:
            rows["grand_total"] = blocks[0]["total_row"]
        rows["check"] = r
        if v.shared_driver:
            rows["driver"] = blocks[0]["driver_row"]  # no fund-level row: the first block's driver row
        if self.defect == "hidden_column_short_subtotal":
            cols["outstanding"] = cols["received_date"] + 1
        self.arows = rows
        self.ablocks = blocks
        self.comp_col_extra = {"withholding": cols["withholding"]} if "withholding" in cols else {}
        self.investor_row = {n: sheet.row(rr) for b in blocks for n, rr in b["investor_rows"].items()}
        for b in blocks:
            if b["gp_name"]:
                self.investor_row[b["gp_name"]] = sheet.row(b["gp_row"])

    # -- ITD ----------------------------------------------------------------------------

    def _itd_blocks(self) -> list[tuple[EventDef, Allocation, bool]]:
        return [(e, a, False) for e, a in zip(self.prior_events, self.prior_allocations)] + \
            [(self.current, self.current_alloc, True)]

    def _build_itd(self) -> None:
        v = self.variant
        sheet = SheetBuilder(self.names["itd"], "itd", v.row_offset, v.col_offset)
        self.itd_sheet = sheet
        alloc = self.alloc_sheet
        band = ["investment_contributions", "cost_contributions", "recallable_distributions",
                "non_recallable_distributions", "tax_withholding"]
        band_labels = ["Investment Contributions", "Cost Contributions", "Recallable Distributions",
                       "Non-Recallable Distributions", "Tax Withholding"]
        band_rows = {cat: i + 1 for i, cat in enumerate(band)}
        overlay_rows: dict[str, int] = {}
        for cat, label in zip(band, band_labels):
            sheet.put(band_rows[cat], 1, label)
        if self.defect == "ok_itd_overlay_rows":
            overlay_rows = {"late_interest": 6, "mgmt_fees": 7}
            sheet.put(6, 1, "Late Interest")
            sheet.put(7, 1, "Mgmt Fees")
        fund_row, event_row, sub_row = 9, 10, 11
        sheet.put(fund_row, 1, formula=f"{q(alloc.name)}!{alloc.coord(self.arows['fund_name'], 2)}", cached=FUND_NAME)

        cum_keys = ["commitment", "investment_contributions", "cost_contributions", "total_contributions",
                    "unfunded", "recallable_distributions", "non_recallable_distributions", "total_distributions"]
        cum_headers = ["Commitment", "Investment Contributions", "Cost Contributions",
                       "Total Paid-In Capital" if v.alt_headers else "Total Contributions", "Unfunded",
                       "Recallable Distributions", "Non-Recallable Distributions", "Total Distributions"]
        if self.defect in WAIVER_DEFECTS:
            # As on the reference client: a waiver column between the contributions and the unfunded.
            cum_keys.insert(4, "waiver")
            cum_headers.insert(4, "Waiver Adjustment")
        if self.defect in CHECK_COLUMN_DEFECTS:
            cum_keys.append("check")
            cum_headers.append("Contributions Check")
        cum_cols = {k: i + 2 for i, k in enumerate(cum_keys)}
        sheet.put(sub_row, 1, "Partner" if v.alt_headers else "Investor")
        for k, h in zip(cum_keys, cum_headers):
            sheet.put(sub_row, cum_cols[k], h)
        waiver_itd = {"Northgate Family Trust": WAIVER_AMOUNT} if self.defect in WAIVER_DEFECTS else {}

        # Column plan for event blocks.
        c = max(cum_cols.values()) + 2
        first_block_col = c
        blocks_plan = []
        itd_blocks = self._itd_blocks()
        if self.defect == "ok_transfer_block":
            itd_blocks.insert(2, (EventDef("transfer", 0, dt.date(2024, 9, 30), {"investment": Decimal("0")}),
                                  Allocation({"investment": {i.name: Decimal("0") for i in self.investors}}, {}),
                                  False))
        for event, allocation, is_current in itd_blocks:
            comps = [C_BY_KEY[k] for k in event.drivers]
            if is_current and self.defect == "ok_withholding_in_itd":
                comps.append(C_WITHHOLDING)
            plan = {"event": event, "alloc": allocation, "current": is_current, "cols": {}}
            for comp in comps:
                plan["cols"][comp.key] = c
                c += 1
            # itd_blocks_in_collapsed_groups: the extra calls are folded into collapsed outline groups
            # and one of them has no Total column.
            extra = (self.defect == "itd_blocks_in_collapsed_groups" and not is_current
                     and event.kind == "capital_call" and event.number >= 4)
            if extra and event.number == 10:
                plan["total_col"] = None
                c += 1
            else:
                plan["total_col"] = c
                c += 2
            plan["collapsed"] = extra
            blocks_plan.append(plan)
        last_col = c
        band_range_first = sheet.col(first_block_col)
        band_range_last = sheet.col(last_col)

        # Row plan: investors per vehicle.
        r = sub_row + 1
        vblocks = []
        for vehicle in self.vehicles:
            vb: dict[str, Any] = {"name": vehicle, "investor_rows": {}}
            if len(self.vehicles) > 1:
                vb["title_row"] = r
                sheet.put(r, 1, vehicle)
                r += 1
            for inv in self.in_vehicle(vehicle):
                if inv.is_gp:
                    continue
                vb["investor_rows"][inv.name] = r
                r += 1
            r += 1
            vb["lp_subtotal"] = r
            r += 2
            gp = next((i for i in self.in_vehicle(vehicle) if i.is_gp), None)
            if gp is not None:
                vb["gp_row"] = r
                vb["gp_name"] = gp.name
                r += 1
                vb["gp_subtotal"] = r
                r += 2
            else:
                vb["gp_row"] = vb["gp_name"] = vb["gp_subtotal"] = None
            vb["total_row"] = r
            r += 2
            vblocks.append(vb)
        check_row = r

        # Values per investor.
        values: dict[str, dict[str, Decimal]] = {}  # investor -> "block_index:comp" -> amount
        for inv in self.investors:
            values[inv.name] = {}
            for bi, plan in enumerate(blocks_plan):
                for key in plan["cols"]:
                    amount = plan["alloc"].amounts.get(key, {}).get(inv.name, Decimal("0"))
                    if plan["current"] and self.defect == "itd_current_value_mismatch" \
                            and key == self._primary_current_key() and inv.name == "Crescent Ridge Investments LP":
                        amount += Decimal("10")
                    values[inv.name][f"{bi}:{key}"] = amount

        # Classification marks.
        marks: dict[int, str] = {}  # column index -> category
        for bi, plan in enumerate(blocks_plan):
            if plan["event"].kind == "transfer":
                continue
            for key, col in plan["cols"].items():
                marks[col] = C_BY_KEY[key].classification
        current_plan = blocks_plan[-1]
        if self.defect == "itd_missing_x":
            dropped = current_plan["cols"][self._primary_current_key()]
            marks.pop(dropped)
        for col, cat in marks.items():
            sheet.put(band_rows[cat], col, "X")
        overlay_marks: dict[int, str] = {}
        if overlay_rows:
            for plan in blocks_plan:
                for key, col in plan["cols"].items():
                    if key == "mgmt_fee":
                        sheet.put(overlay_rows["mgmt_fees"], col, "X")
                        overlay_marks[col] = "mgmt_fees"

        # Event headers + sub-headers (+ merged label ranges).
        for bi, plan in enumerate(blocks_plan):
            event = plan["event"]
            first = min(plan["cols"].values())
            label = self._block_label(bi, event)
            sheet.put(event_row, first, label)
            last_col = plan["total_col"] or max(plan["cols"].values())
            sheet.merges.append(f"{sheet.coord(event_row, first)}:{sheet.coord(event_row, last_col)}")
            for key, col in plan["cols"].items():
                sheet.put(sub_row, col, event.header_for(C_BY_KEY[key]) if event.kind != "transfer" else "Transfer")
            if plan["total_col"]:
                sheet.put(sub_row, plan["total_col"], "Total")
            if plan.get("collapsed"):
                for col in [*plan["cols"].values(), *([plan["total_col"]] if plan["total_col"] else [])]:
                    sheet.hidden_cols.add(sheet.col(col))
                    sheet.outlined_cols.add(sheet.col(col))
        if self.defect == "prior_block_header_live_link":
            # The first prior block's investment sub-header reads the Allocation header row.
            col = blocks_plan[0]["cols"]["investment"]
            sheet.put(sub_row, col, formula=f"{q(alloc.name)}!{alloc.coord(self.arows['header'], self.comp_col['investment'])}",
                      cached=self.current.header_for(C_INVESTMENT))

        # Investor rows.
        cum_values: dict[str, dict[str, Decimal]] = {}
        for vb in vblocks:
            rows_in_block = list(vb["investor_rows"].items()) + ([(vb["gp_name"], vb["gp_row"])] if vb["gp_row"] else [])
            for name, rr in rows_in_block:
                inv = self.by_name(name)
                arow = self.investor_row[name]
                a_inv_col = alloc.col(self.acols["investor"])
                sheet.put(rr, 1, formula=f"{q(alloc.name)}!{a_inv_col}{arow}", cached=name)
                inv_ref = f"${sheet.col(1)}{sheet.row(rr)}"
                a_commit = alloc.col(self.acols["commitment"])
                sheet.put(rr, cum_cols["commitment"], formula=(
                    f"SUMIFS({q(alloc.name)}!${a_commit}:${a_commit},{q(alloc.name)}!${a_inv_col}:${a_inv_col},{inv_ref})"),
                    cached=_num(inv.commitment), fmt=MONEY_FMT)
                for bi, plan in enumerate(blocks_plan):
                    for key, col in plan["cols"].items():
                        amount = values[name][f"{bi}:{key}"]
                        if plan["current"]:
                            a_col = alloc.col(self.comp_col[key])
                            if self.defect == "ok_sumif_pulls":
                                formula = (f"SUMIF({q(alloc.name)}!${a_inv_col}:${a_inv_col},{inv_ref},"
                                           f"{q(alloc.name)}!{a_col}:{a_col})")
                            else:
                                formula = (f"SUMIFS({q(alloc.name)}!{a_col}:{a_col},{q(alloc.name)}!${a_inv_col}:${a_inv_col},"
                                           f"{inv_ref})")
                            if self.defect == "itd_current_value_mismatch" and key == self._primary_current_key() \
                                    and name == "Crescent Ridge Investments LP":
                                sheet.put(rr, col, _num(amount), fmt=MONEY_FMT)
                            else:
                                sheet.put(rr, col, formula=formula, cached=_num(amount), fmt=MONEY_FMT)
                        else:
                            sheet.put(rr, col, _num(amount), fmt=MONEY_FMT)
                    block_cols = sorted(plan["cols"].values())
                    total = sum(values[name][f"{bi}:{k}"] for k in plan["cols"])
                    if plan["total_col"]:
                        sheet.put(rr, plan["total_col"], formula=(
                            f"SUM({sheet.coord(rr, block_cols[0])}:{sheet.coord(rr, block_cols[-1])})"),
                            cached=_num(total), fmt=MONEY_FMT)
                # Cumulative columns.
                cum = {cat: Decimal("0") for cat in band}
                for bi, plan in enumerate(blocks_plan):
                    for key, col in plan["cols"].items():
                        cat = marks.get(col)
                        if cat:
                            cum[cat] += values[name][f"{bi}:{key}"]
                row_ref = sheet.row(rr)

                def sumif(cat: str) -> str:
                    band_row = sheet.row(band_rows[cat])
                    return (f"SUMIF(${band_range_first}${band_row}:${band_range_last}${band_row},\"X\","
                            f"${band_range_first}{row_ref}:${band_range_last}{row_ref})")

                total_contrib = cum["investment_contributions"] + cum["cost_contributions"]
                waiver = waiver_itd.get(name, Decimal("0"))
                unfunded = inv.commitment - total_contrib - cum["recallable_distributions"] - waiver
                total_dist = cum["recallable_distributions"] + cum["non_recallable_distributions"]
                recallable_formula = sumif("recallable_distributions")
                if self.defect == "ok_derived_recallable":
                    # A recycling cap: the recallable room is the lesser of 15% of the commitment and
                    # the recalled amounts, not a plain accumulation of the marked columns.
                    recallable_formula = (f"-MIN(ROUND({sheet.coord(rr, cum_cols['commitment'])}*0.15,2),"
                                          f"-({recallable_formula}))")
                unfunded_formula = (f"{sheet.coord(rr, cum_cols['commitment'])}-{sheet.coord(rr, cum_cols['total_contributions'])}"
                                    f"-{sheet.coord(rr, cum_cols['recallable_distributions'])}")
                if "waiver" in cum_cols:
                    sheet.put(rr, cum_cols["waiver"], _num(waiver), fmt=MONEY_FMT)
                    unfunded_formula += f"-{sheet.coord(rr, cum_cols['waiver'])}"
                if "check" in cum_cols:
                    # Contributions to date less the Allocation's prior contributions and current call: zero.
                    a_prior = alloc.col(self.acols["rf_prior"])
                    a_call = alloc.col(self.acols["call_total"])
                    check_formula = (f"{sheet.coord(rr, cum_cols['total_contributions'])}+SUMIFS({q(alloc.name)}!${a_prior}:"
                                     f"${a_prior},{q(alloc.name)}!${a_inv_col}:${a_inv_col},{inv_ref})")
                    call_total = sum((self.current_alloc.amounts[k].get(name, Decimal("0")) for k in self.current.drivers
                                      if C_BY_KEY[k].side == "call"), Decimal("0"))
                    if self.defect == "itd_check_column_nonzero":
                        check_value = call_total  # the current call is left out, so every called LP shows it
                    else:
                        check_formula += (f"-SUMIFS({q(alloc.name)}!${a_call}:${a_call},{q(alloc.name)}!${a_inv_col}:"
                                          f"${a_inv_col},{inv_ref})")
                        check_value = Decimal("0")
                    sheet.put(rr, cum_cols["check"], formula=check_formula, cached=_num(check_value), fmt=MONEY_FMT)
                cells = {
                    "investment_contributions": (sumif("investment_contributions"), cum["investment_contributions"]),
                    "cost_contributions": (sumif("cost_contributions"), cum["cost_contributions"]),
                    "total_contributions": (
                        f"SUM({sheet.coord(rr, cum_cols['investment_contributions'])}:"
                        f"{sheet.coord(rr, cum_cols['cost_contributions'])})", total_contrib),
                    "unfunded": (unfunded_formula, unfunded),
                    "recallable_distributions": (recallable_formula, cum["recallable_distributions"]),
                    "non_recallable_distributions": (sumif("non_recallable_distributions"),
                                                     cum["non_recallable_distributions"]),
                    "total_distributions": (
                        f"SUM({sheet.coord(rr, cum_cols['recallable_distributions'])}:"
                        f"{sheet.coord(rr, cum_cols['non_recallable_distributions'])})", total_dist),
                }
                for key, (formula, cached) in cells.items():
                    sheet.put(rr, cum_cols[key], formula=formula, cached=_num(cached), fmt=MONEY_FMT)
                cum_values[name] = {
                    "commitment": inv.commitment,
                    "total_contributions": total_contrib,
                    "unfunded": unfunded,
                    "total_distributions": total_dist,
                    "recallable_distributions": cum["recallable_distributions"],
                    "non_recallable_distributions": cum["non_recallable_distributions"],
                }
            # Subtotals.
            lp_rows = list(vb["investor_rows"].values())
            numeric_cols = sorted(set(cum_cols.values()) - {1}) + sorted(
                col for plan in blocks_plan for col in [*plan["cols"].values(), plan["total_col"]] if col)
            for col in numeric_cols:
                lp_sum = sum(_cached_decimal(sheet, sheet.coord(rr, col)) for rr in lp_rows)
                sheet.put(vb["lp_subtotal"], col, formula=(
                    f"SUM({sheet.coord(lp_rows[0], col)}:{sheet.coord(lp_rows[-1], col)})"),
                    cached=_num(lp_sum), fmt=MONEY_FMT)
                if vb["gp_row"]:
                    gp_val = _cached_decimal(sheet, sheet.coord(vb["gp_row"], col))
                    sheet.put(vb["gp_subtotal"], col, formula=f"SUM({sheet.coord(vb['gp_row'], col)})",
                              cached=_num(gp_val), fmt=MONEY_FMT)
                    total_formula = f"{sheet.coord(vb['lp_subtotal'], col)}+{sheet.coord(vb['gp_subtotal'], col)}"
                else:
                    gp_val = Decimal("0")
                    total_formula = sheet.coord(vb["lp_subtotal"], col)
                sheet.put(vb["total_row"], col, formula=total_formula, cached=_num(lp_sum + gp_val), fmt=MONEY_FMT)
            sheet.put(vb["lp_subtotal"], 1, "Limited Partners")
            if vb["gp_subtotal"]:
                sheet.put(vb["gp_subtotal"], 1, "General Partner")
            sheet.put(vb["total_row"], 1, f"Total {vehicle}" if len(self.vehicles) > 1 else "Total Partnership")

        # Prior-block live links on the $0 GP row (calibration probe) / on an LP (defect).
        if self.defect in ("live_prior_itd_link", "zero_live_prior_link"):
            plan = blocks_plan[0]
            key = "investment"
            col = plan["cols"][key]
            target = "Harborview Capital Partners, LLC" if self.defect == "live_prior_itd_link" else vblocks[0]["gp_name"]
            rr = vblocks[0]["gp_row"] if target == vblocks[0]["gp_name"] else vblocks[0]["investor_rows"][target]
            coord = sheet.coord(rr, col)
            a_col = alloc.col(self.comp_col[key])
            a_inv_col = alloc.col(self.acols["investor"])
            cell = sheet.get(coord)
            sheet.put_at(coord, formula=(f"SUMIFS({q(alloc.name)}!{a_col}:{a_col},{q(alloc.name)}!${a_inv_col}:"
                                         f"${a_inv_col},${sheet.col(1)}{sheet.row(rr)})"),
                         cached=cell.value, fmt=MONEY_FMT)

        # Check row: current block totals vs Allocation grand totals (the look-through block re-allocates
        # the others' GP rows and is outside both).
        sheet.put(check_row, 1, "Check vs Allocation")
        grand_itd_rows = [vb["total_row"] for vb in vblocks if vb["name"] != self.lookthrough]
        for key, col in current_plan["cols"].items():
            a_col = alloc.col(self.comp_col[key])
            itd_total = sum(_cached_decimal(sheet, sheet.coord(rr, col)) for rr in grand_itd_rows)
            alloc_total = sum(a for n, a in self.current_alloc.amounts[key].items() if not self.is_lookthrough(n))
            sheet.put(check_row, col, formula=(
                "+".join(sheet.coord(rr, col) for rr in grand_itd_rows) +
                f"-{q(alloc.name)}!{a_col}{alloc.row(self.arows['grand_total'])}"),
                cached=_num(itd_total - alloc_total), fmt=MONEY_FMT)

        self.itd_cum_cols = cum_cols
        self.itd_cum = cum_values
        self.itd_current_total_col = current_plan["total_col"]
        self.itd_values = values
        self.itd_blocks_plan = blocks_plan
        self.sheets.append(sheet)
        sheet.view = "pageBreakPreview"

        # Golden layout.
        layout_blocks = []
        for bi, plan in enumerate(blocks_plan):
            event = plan["event"]
            block = {
                "label": self._block_label(bi, event),
                "event_type": "transfer" if event.kind == "transfer" else event.kind,
                "number": None if event.kind == "transfer" else event.number,
                "date": None if event.kind == "transfer" else event.date.isoformat(),
                "first_column": sheet.col(min(plan["cols"].values())),
                "last_column": sheet.col(max(plan["cols"].values())),
                "total_column": sheet.col(plan["total_col"]) if plan["total_col"] else None,
                "is_current": plan["current"],
                "components": [
                    {"column": sheet.col(col), "component_type": C_BY_KEY[key].component_type,
                     "side": C_BY_KEY[key].side}
                    for key, col in plan["cols"].items()
                ],
            }
            layout_blocks.append(block)
        self.layouts[sheet.name] = {
            "role": "itd",
            "sheet": sheet.name,
            "classification_rows": {cat: sheet.row(r) for cat, r in band_rows.items()},
            "overlay_rows": {cat: sheet.row(r) for cat, r in overlay_rows.items()},
            "event_header_row": sheet.row(event_row),
            "subheader_row": sheet.row(sub_row),
            "investor_column": sheet.col(1),
            "cumulative_columns": {k: sheet.col(c) for k, c in cum_cols.items() if k not in ("waiver", "check")},
            # The mapper on the reference client returned one block of twenty-eight: code enumerates the rest.
            "event_blocks": [b for b in layout_blocks if b["is_current"]]
            if self.defect == "itd_blocks_in_collapsed_groups" else layout_blocks,
            "check_columns": [sheet.col(cum_cols["check"])] if "check" in cum_cols else [],
            "vehicles": [
                {
                    "name": vb["name"],
                    "title_row": sheet.row(vb["title_row"]) if "title_row" in vb else None,
                    "investor_rows": [sheet.row(min(vb["investor_rows"].values())),
                                      sheet.row(max(vb["investor_rows"].values()))],
                    "gp_rows": [sheet.row(vb["gp_row"])] if vb["gp_row"] else [],
                    "subtotal_rows": {"limited_partners": sheet.row(vb["lp_subtotal"]),
                                      "general_partner": sheet.row(vb["gp_subtotal"]) if vb["gp_subtotal"] else None,
                                      "total": sheet.row(vb["total_row"])},
                }
                for vb in vblocks
            ],
            "check_rows": [sheet.row(check_row)],
        }
        self.truth["itd"] = {
            "current_block_label": current_plan["event"].itd_label,
            "event_labels": [b["label"] for b in layout_blocks],
            "cumulative": {n: {k: str(v) for k, v in d.items()} for n, d in cum_values.items()},
            "current_block": {
                n: {sheet.col(col): str(values[n][f"{len(blocks_plan) - 1}:{key}"])
                    for key, col in current_plan["cols"].items()}
                for n in values
            },
        }
        self.facts_itd_unclassified = 1 if self.defect == "itd_missing_x" else 0

    def _primary_current_key(self) -> str:
        return next(iter(self.current.drivers))

    def _block_label(self, index: int, event: EventDef) -> str:
        if event.kind == "transfer":
            return "3Q24 Transfers"
        if self.stage == "current" and self.defect == "itd_block_relabelled" and index == 1:
            return f"{event.itd_label} (final)"  # relabelled since the prior workbook, same columns
        return event.itd_label

    # -- Allocation -----------------------------------------------------------------------

    def _build_allocation(self) -> None:
        sheet = self.alloc_sheet
        v = self.variant
        cols, rows = self.acols, self.arows
        event = self.current
        alloc = self.current_alloc
        active = {k for k in event.drivers}
        itd = self.itd_sheet
        holidays = self.holidays

        # Header area.
        sheet.put(rows["fund_name"], 2, FUND_NAME)
        sheet.merges.append(f"{sheet.coord(rows['fund_name'], 2)}:{sheet.coord(rows['fund_name'], 3)}")
        word = "Distribution" if event.kind == "distribution" else "Capital Call"
        notice_fmt = "General" if self.defect == "fmt_serial_date" else DATE_LONG_FMT
        sheet.put(rows["notice"], 2, f"{word} - Notice Date")
        notice_coord = sheet.put(rows["notice"], 3, self.notice_date, fmt=notice_fmt)
        due_label = "Payment Date" if event.kind == "distribution" else "Due Date"
        sheet.put(rows["due"], 2, f"{word} - {due_label}")
        if self.defect == "weekend_due_date":
            due_date = dt.date(2026, 6, 13)
            due_coord = sheet.put(rows["due"], 3, due_date, fmt=DATE_LONG_FMT)
        else:
            days = -3 if self.defect == "due_before_notice" else 10
            due_date = _workday(self.notice_date, days, holidays)
            due_coord = sheet.put(rows["due"], 3, formula=f"WORKDAY({notice_coord},{days},{self.holiday_range})",
                                  cached=_serial(due_date), fmt=DATE_LONG_FMT)
        self.due_date = due_date
        self.notice_coord, self.due_coord = notice_coord, due_coord
        if event.kind in ("distribution", "net_event"):
            sheet.put(rows["due"], cols["commitment"], "Carried Interest Rate")
            self.carry_rate_coord = sheet.put(rows["due"], cols["commitment_pct"], _num(CARRY_RATE), fmt="0.00%")
        # Decorative banners above the header row.
        call_cols = sorted(self.call_comp_cols.values())
        dist_cols = sorted(self.dist_comp_cols.values())
        sheet.put(rows["notice"], call_cols[0], "Capital Call")
        sheet.merges.append(f"{sheet.coord(rows['notice'], call_cols[0])}:{sheet.coord(rows['notice'], cols['call_total'])}")
        sheet.put(rows["notice"], dist_cols[0], "Distribution")
        sheet.merges.append(f"{sheet.coord(rows['notice'], dist_cols[0])}:{sheet.coord(rows['notice'], cols['dist_total'])}")

        # Headers.
        header = rows["header"]
        alt = v.alt_headers
        labels = {
            "close": "Close #",
            "investor": "Partner Name" if alt else "Investor",
            "affiliate": "Affiliate",
            "commitment": "Commitment" if alt else "Total Commitment",
            "commitment_pct": "% of Fund" if alt else "Commitment %",
            "contributed": "Contributed Capital",
            "contributed_pct": "Contributed Capital %",
            "late_interest": "Late Interest",
            "cash_due": "Cash Due",
            "rf_commitment": "Commitment Amount",
            "rf_prior": "Prior Capital Contributions",
            "rf_prior_recallable": "Prior Recallable Distributions",
            "rf_current_call": "Current Capital Call",
            "rf_current_recallable": "Current Recallable Distributions",
            "rf_waiver": "Waiver Adjustment",
            "rf_remaining": "Partners' Remaining Commitment",
            "withholding": "Tax Withholding",
            "prior_gross_dist": "Prior Gross Distributions",
            "current_dist": "Current Distributions",
            "total_gross_dist": "Total Gross Distributions",
            "received": "Received",
            "received_date": "Date",
        }
        for key, label in labels.items():
            if key in cols:
                sheet.put(header, cols[key], label)
        current_label = event.short_label
        comp_headers: dict[str, str] = {}
        for comp in (*CALL_COMPONENTS, *DIST_COMPONENTS):
            text = event.header_for(comp)
            if comp.key not in active:
                text = f"{comp.header} (not used {current_label})"
                if self.defect == "ok_stale_label_inactive" and comp.key == "roc":
                    text = f"{comp.header} (Distribution #1)"  # left over from the prior distribution
            elif self.defect == "stale_label_on_active" and comp.key == "investment":
                text = f"{comp.header} (Capital Call #3)"
            comp_headers[comp.key] = text
            sheet.put(header, self.comp_col[comp.key], text)
        self.comp_headers = comp_headers
        call_side_active = any(C_BY_KEY[k].side == "call" for k in active)
        dist_side_active = any(C_BY_KEY[k].side == "distribution" for k in active)
        call_total_label = current_label if call_side_active else f"Capital Call (not used {current_label})"
        dist_total_label = current_label if (dist_side_active and not call_side_active) else (
            "Distribution Total" if dist_side_active else f"Distribution (not used {current_label})")
        if self.defect == "ok_stale_label_inactive":
            dist_total_label = "Distribution #1"
        self.call_total_header = sheet.put(header, cols["call_total"], call_total_label)
        self.dist_total_header = sheet.put(header, cols["dist_total"], dist_total_label)
        self.event_label_coord = self.call_total_header if call_side_active else self.dist_total_header
        if self.defect == "stale_component":
            sheet.put(header, cols["stale"], "Placement Fees (Capital Call #3)")
        if self.defect == "merged_component_header":
            sheet.merges.append(f"{sheet.coord(header, self.call_comp_cols['investment'])}:"
                                f"{sheet.coord(header, self.call_comp_cols['expenses'])}")
            del sheet.cells[sheet.coord(header, self.call_comp_cols["expenses"])]
        sheet.hidden_cols.add(sheet.col(cols["late_interest"]))

        fee_sheet = self.names["mgmt_fee"]
        fee_amount_col = "F"
        fee_name_col = "B"
        fee_vehicle_col = "C"

        # Fund-level driver row.
        fund_driver_row = rows["driver"]
        drivers: dict[str, Decimal] = {}
        for comp in (*CALL_COMPONENTS, *DIST_COMPONENTS):
            if comp.key in active:
                if comp.key == "mgmt_fee":
                    drivers[comp.key] = self.fee_tab_total()
                else:
                    drivers[comp.key] = event.drivers[comp.key]
            else:
                drivers[comp.key] = Decimal("0")
        call_driver_total = sum(drivers[c.key] for c in CALL_COMPONENTS)
        dist_driver_total = sum(drivers[c.key] for c in DIST_COMPONENTS)

        withholding = self.current_alloc.amounts.get("withholding", {})
        withholding_total = sum(withholding.values(), Decimal("0"))

        def put_event_totals(row: int, call_amount: Decimal, dist_amount: Decimal) -> None:
            sheet.put(row, cols["call_total"], formula=(
                f"SUM({sheet.coord(row, call_cols[0])}:{sheet.coord(row, call_cols[-1])})"),
                cached=_num(call_amount), fmt=MONEY_FMT)
            sheet.put(row, cols["dist_total"], formula=(
                f"SUM({sheet.coord(row, dist_cols[0])}:{sheet.coord(row, dist_cols[-1])})"),
                cached=_num(dist_amount), fmt=MONEY_FMT)
            sheet.put(row, cols["late_interest"], 0, fmt=MONEY_FMT)
            terms = [sheet.coord(row, cols["call_total"]), sheet.coord(row, cols["dist_total"]),
                     sheet.coord(row, cols["late_interest"])]
            extra = Decimal("0")
            if self.withholding:
                sheet.put(row, cols["withholding"], _num(withholding_total), fmt=MONEY_FMT)  # typed, as on the client
                terms.append(sheet.coord(row, cols["withholding"]))
                extra = withholding_total
            sheet.put(row, cols["cash_due"], formula=f"SUM({','.join(terms)})",
                      cached=_num(call_amount + dist_amount + extra), fmt=MONEY_FMT)

        if not v.shared_driver:
            for comp_key, amount in drivers.items():
                col = self.comp_col[comp_key]
                if comp_key == "carry" and self.defect == "typed_driver_in_gp_row":
                    continue  # the driver cell is left blank; the GP row holds the typed total
                if comp_key == "mgmt_fee" and comp_key in active:
                    sheet.put(fund_driver_row, col, formula=f"{q(fee_sheet)}!${fee_amount_col}${self._fee_total_row()}",
                              cached=_num(amount), fmt=MONEY_FMT)
                elif comp_key == "investment" and comp_key in active and self.tracker_current_row():
                    # As in the reference sample, the investment driver pulls the current deal from the tracker.
                    sheet.put(fund_driver_row, col, formula=f"{q('Portfolio Investment Tracker')}!$E${self.tracker_current_row()}",
                              cached=_num(amount), fmt=MONEY_FMT)
                else:
                    sheet.put(fund_driver_row, col, _num(amount), fmt=MONEY_FMT)
            put_event_totals(fund_driver_row, call_driver_total, dist_driver_total)
            if self.defect == "stale_component":
                sheet.put(fund_driver_row, cols["stale"], 0, fmt=MONEY_FMT)

        # Vehicle blocks.
        itd_inv = f"{q(itd.name)}!${itd.col(1)}:${itd.col(1)}"
        itd_total_contrib = itd.col(self.itd_cum_cols["total_contributions"])
        itd_total_dist = itd.col(self.itd_cum_cols["total_distributions"])
        truth_investors: dict[str, dict[str, Any]] = {}
        vehicle_layouts = []
        grand: dict[int, Decimal] = {}
        for block in self.ablocks:
            vehicle = block["name"]
            members = self.in_vehicle(vehicle)
            lps = [i for i in members if not i.is_gp]
            vdriver_row = block["driver_row"]
            vehicle_commitment = sum(i.commitment for i in members)
            if len(self.vehicles) > 1:
                if "title_row" in block:
                    sheet.put(block["title_row"], cols["investor"], vehicle)
                sheet.put(vdriver_row, cols["investor"], f"{vehicle} Allocation")
                vehicle_call = vehicle_dist = Decimal("0")
                for comp in (*CALL_COMPONENTS, *DIST_COMPONENTS):
                    col = self.comp_col[comp.key]
                    amount = sum(alloc.amounts.get(comp.key, {}).get(i.name, Decimal("0")) for i in members) \
                        if comp.key in active else Decimal("0")
                    if comp.side == "call":
                        vehicle_call += amount
                    else:
                        vehicle_dist += amount
                    if vehicle == self.lookthrough:
                        # The look-through driver is the GP rows of the other blocks added up.
                        gp_cells = [sheet.coord(b["gp_row"], col) for b in self.ablocks if b["gp_row"]]
                        sheet.put(vdriver_row, col, formula="+".join(gp_cells), cached=_num(amount), fmt=MONEY_FMT)
                    elif comp.key == "mgmt_fee" and comp.key in active:
                        terms = [f"SUMIFS({q(fee_sheet)}!${c}:${c},{q(fee_sheet)}!${fee_vehicle_col}:${fee_vehicle_col},"
                                 f"\"{vehicle}\")" for c in self.fee_cols()]
                        sheet.put(vdriver_row, col, formula="+".join(terms), cached=_num(amount), fmt=MONEY_FMT)
                    else:
                        sheet.put(vdriver_row, col, _num(amount), fmt=MONEY_FMT)
                if v.shared_driver:
                    put_event_totals(vdriver_row, vehicle_call, vehicle_dist)
            total_row_commit = sheet.coord(block["total_row"], cols["commitment"])
            contributed_total = sum(self.prior_contributed[i.name] for i in members)
            gp_member = next((i for i in members if i.is_gp), None)
            rows_for_block = [(i, block["investor_rows"][i.name]) for i in lps] + \
                             ([(gp_member, block["gp_row"])] if gp_member is not None else [])
            col_sums_lp: dict[int, Decimal] = {}
            col_sums_gp: dict[int, Decimal] = {}
            for inv, rr in rows_for_block:
                row_vals: dict[int, Decimal] = {}

                def put(col_key_or_idx, value=None, formula=None, cached=None, fmt=MONEY_FMT):
                    col = cols[col_key_or_idx] if isinstance(col_key_or_idx, str) else col_key_or_idx
                    sheet.put(rr, col, value, formula=formula, cached=cached, fmt=fmt)
                    num = value if formula is None else cached
                    if isinstance(num, (int, float, Decimal)) and not isinstance(num, bool):
                        row_vals[col] = Decimal(str(num))

                put("close", 1 if not inv.late_closer else 5, fmt="General")
                sheet.put(rr, cols["investor"], inv.name)
                sheet.put(rr, cols["affiliate"], "GP" if inv.is_gp else ("Y" if inv.affiliate else "N"))
                put("commitment", _num(inv.commitment))
                pct = (inv.commitment / vehicle_commitment) if vehicle_commitment else Decimal("0")
                if inv.name in self.pct_overrides:
                    put("commitment_pct", float(self.pct_overrides[inv.name]), fmt=PCT_FMT)
                else:
                    put("commitment_pct", formula=f"{sheet.coord(rr, cols['commitment'])}/{_abs_row(total_row_commit)}",
                        cached=float(pct), fmt=PCT_FMT)
                pct_ref = f"${sheet.col(cols['commitment_pct'])}{sheet.row(rr)}"
                if "contributed" in cols:
                    contributed = self.prior_contributed[inv.name]
                    put("contributed", formula=f"-{sheet.coord(rr, cols['rf_prior'])}", cached=_num(contributed))
                    cpct = contributed / contributed_total if contributed_total else Decimal("0")
                    put("contributed_pct", formula=(
                        f"{sheet.coord(rr, cols['contributed'])}/{_abs_row(sheet.coord(block['total_row'], cols['contributed']))}"),
                        cached=float(cpct), fmt=PCT_FMT)
                    cpct_ref = f"${sheet.col(cols['contributed_pct'])}{sheet.row(rr)}"
                row_amounts: dict[str, Decimal] = {}
                for comp in (*CALL_COMPONENTS, *DIST_COMPONENTS):
                    col = self.comp_col[comp.key]
                    driver_ref = f"{sheet.col(col)}${sheet.row(vdriver_row)}"
                    amount = alloc.amounts.get(comp.key, {}).get(inv.name, Decimal("0")) if comp.key in active \
                        else Decimal("0")
                    row_amounts[comp.key] = amount
                    offset = alloc.offsets.get(comp.key, {}).get(inv.name)
                    basis_ref = cpct_ref if comp.side == "distribution" and "contributed_pct" in cols else pct_ref
                    if comp.key == "mgmt_fee" and comp.key in active and self.defect == "fee_pulled_from_wrong_row" \
                            and inv.name == "Juniper Hollow Partners":
                        wrong_row = self._fee_rows()["lp_rows"]["Meridian Endowment Fund"]
                        formula = f"{q(fee_sheet)}!${fee_amount_col}${wrong_row}"
                    elif comp.key == "mgmt_fee" and comp.key in active and self.defect == "fee_link_shifted" \
                            and inv.name == "Juniper Hollow Partners":
                        formula = (f"SUMIFS({q(fee_sheet)}!$E:$E,{q(fee_sheet)}!"
                                   f"${fee_name_col}:${fee_name_col},${sheet.col(cols['investor'])}{sheet.row(rr)})")
                    elif comp.key == "mgmt_fee" and comp.key in active and vehicle == self.lookthrough:
                        formula = f"ROUND({driver_ref}*{basis_ref},2)"  # the GP's partners share the GP's $0 fee
                    elif comp.key == "mgmt_fee" and comp.key in active:
                        formula = "+".join(
                            f"SUMIFS({q(fee_sheet)}!${c}:${c},{q(fee_sheet)}!${fee_name_col}:${fee_name_col},"
                            f"${sheet.col(cols['investor'])}{sheet.row(rr)})" for c in self.fee_cols())
                    elif comp.key == "carry" and comp.key in active and self.waterfall:
                        rate_ref = _abs(self.carry_rate_coord)
                        unpulled = self.defect == "waterfall_row_not_pulled" and inv.name == "Juniper Hollow Partners"
                        self.waterfall_rows.append({
                            "name": inv.name + (" LP" if unpulled else ""), "vehicle": vehicle, "is_gp": inv.is_gp,
                            "amount": self.unpulled_carry if unpulled else amount,
                            "offset": offset, "driver": f"{sheet.col(col)}${sheet.row(vdriver_row)}",
                            "basis": basis_ref, "rate": rate_ref})
                        formula = (f"SUMIFS({q(WATERFALL_SHEET)}!$E:$E,{q(WATERFALL_SHEET)}!$B:$B,"
                                   f"${sheet.col(cols['investor'])}{sheet.row(rr)})")
                        offset = None  # the waterfall carries the rounding plug
                    elif comp.key == "carry" and comp.key in active and self.defect == "typed_driver_in_gp_row":
                        if inv.is_gp:
                            put(col, _num(amount), fmt=MONEY_FMT)  # typed: this cell drives the LP formulas
                            continue
                        formula = f"ROUND(-{sheet.col(col)}${sheet.row(block['gp_row'])}*{basis_ref},2)"
                    elif comp.key == "carry" and comp.key in active and inv.is_gp:
                        rate_ref = "0.25" if self.defect == "carry_split_wrong" else _abs(self.carry_rate_coord)
                        formula = f"ROUND({driver_ref}*{rate_ref},2)"
                    elif comp.key == "carry" and comp.key in active and vehicle == self.lookthrough:
                        formula = f"ROUND({driver_ref}*{basis_ref},2)"  # the GP's carry, shared by its partners
                    elif comp.key == "carry" and comp.key in active:
                        formula = (f"ROUND(({driver_ref}-{sheet.col(col)}${sheet.row(block['gp_row'])})*{basis_ref},2)")
                    elif (comp.key, inv.name) in alloc.unrounded:
                        formula = f"{driver_ref}*{basis_ref}"
                    elif comp.key in ("roc", "gain") and self.defect in WHOLE_DOLLAR_DEFECTS:
                        formula = f"ROUND({driver_ref}*{basis_ref},0)"
                    else:
                        formula = f"ROUND({driver_ref}*{basis_ref},2)"
                    if offset:
                        formula += f"{'+' if offset > 0 else '-'}{_fmt_offset(abs(offset))}"
                    fmt = "General" if (self.defect == "fmt_general_money" and comp.key == "investment") else MONEY_FMT
                    if self.defect == "hardcoded_lp_cell" and comp.key == "investment" \
                            and inv.name == "Alder & Finch Holdings, LLC":
                        put(col, _num(amount), fmt=fmt)
                    else:
                        put(col, formula=formula, cached=_num(amount), fmt=fmt)
                call_total = sum(row_amounts[c.key] for c in CALL_COMPONENTS)
                dist_total = sum(row_amounts[c.key] for c in DIST_COMPONENTS)
                put("call_total", formula=(f"SUM({sheet.coord(rr, call_cols[0])}:{sheet.coord(rr, call_cols[-1])})"),
                    cached=_num(call_total))
                put("dist_total", formula=(f"SUM({sheet.coord(rr, dist_cols[0])}:{sheet.coord(rr, dist_cols[-1])})"),
                    cached=_num(dist_total))
                if self.defect == "stale_component" and not inv.is_gp:
                    stale = r2(Decimal("60000") * inv.commitment / vehicle_commitment)
                    put("stale", _num(stale))
                late_row = vdriver_row if v.shared_driver else fund_driver_row
                put("late_interest", formula=f"ROUND({sheet.col(cols['late_interest'])}${sheet.row(late_row)}*{pct_ref},2)",
                    cached=0)
                cash_formula = (f"{sheet.coord(rr, cols['call_total'])}+{sheet.coord(rr, cols['dist_total'])}+"
                                f"{sheet.coord(rr, cols['late_interest'])}")
                withheld = Decimal("0")
                if self.withholding:
                    withheld = withholding.get(inv.name, Decimal("0"))
                    put("withholding", _num(withheld))
                    cash_formula += f"+{sheet.coord(rr, cols['withholding'])}"
                put("cash_due", formula=cash_formula, cached=_num(call_total + dist_total + withheld))
                # Roll-forward.
                itd_e = self.itd_cum[inv.name]["total_contributions"]
                itd_dist = self.itd_cum[inv.name]["total_distributions"]
                prior = itd_e - call_total  # AJ = -(ITD E) + call total  -> -prior
                put("rf_commitment", formula=sheet.coord(rr, cols["commitment"]), cached=_num(inv.commitment))
                inv_ref = f"${sheet.col(cols['investor'])}{sheet.row(rr)}"
                put("rf_prior", formula=(f"-SUMIFS({q(itd.name)}!${itd_total_contrib}:${itd_total_contrib},{itd_inv},"
                                         f"{inv_ref})+{sheet.coord(rr, cols['call_total'])}"), cached=_num(-prior))
                put("rf_prior_recallable", 0)
                current_call = -call_total
                if self.defect == "rf_current_call_stale" and inv.name == "Pemberton Street FLP":
                    stale_amount = -sum(self.prior_allocations[-1].amounts[k].get(inv.name, Decimal("0"))
                                        for k in self.prior_allocations[-1].amounts if C_BY_KEY[k].side == "call")
                    current_call = stale_amount
                    put("rf_current_call", _num(current_call))
                else:
                    put("rf_current_call", formula=f"-{sheet.coord(rr, cols['call_total'])}", cached=_num(current_call))
                put("rf_current_recallable", 0)
                remaining = inv.commitment - prior + current_call
                if "rf_waiver" in cols:
                    waiver = -WAIVER_AMOUNT if inv.name == "Northgate Family Trust" else Decimal("0")
                    put("rf_waiver", _num(waiver))
                    if self.defect == "rf_sum_omits_adjustment":
                        last_rf = cols["rf_current_recallable"]  # the SUM stops before the waiver column
                    else:
                        last_rf = cols["rf_waiver"]
                        remaining += waiver
                else:
                    last_rf = cols["rf_current_recallable"]
                put("rf_remaining", formula=(
                    f"SUM({sheet.coord(rr, cols['rf_commitment'])}:{sheet.coord(rr, last_rf)})"),
                    cached=_num(remaining))
                prior_dist = itd_dist - dist_total
                put("prior_gross_dist", formula=(
                    f"SUMIFS({q(itd.name)}!${itd_total_dist}:${itd_total_dist},{itd_inv},{inv_ref})-"
                    f"{sheet.coord(rr, cols['dist_total'])}"), cached=_num(prior_dist))
                put("current_dist", formula=sheet.coord(rr, cols["dist_total"]), cached=_num(dist_total))
                put("total_gross_dist", formula=(
                    f"{sheet.coord(rr, cols['prior_gross_dist'])}+{sheet.coord(rr, cols['current_dist'])}"),
                    cached=_num(prior_dist + dist_total))
                if "outstanding" in cols and not inv.is_gp:
                    # A hidden admin column ("Outstanding" = cash due less received), as in the reference client.
                    put("outstanding", formula=f"{sheet.coord(rr, cols['cash_due'])}-{sheet.coord(rr, cols['received'])}",
                        cached=_num(call_total + dist_total))
                target = col_sums_gp if inv.is_gp else col_sums_lp
                for col, val in row_vals.items():
                    target[col] = target.get(col, Decimal("0")) + val
                truth_investors[inv.name] = {
                    "vehicle": vehicle,
                    "row": sheet.row(rr),
                    "commitment": str(inv.commitment),
                    "is_gp": inv.is_gp,
                    "affiliate": inv.affiliate,
                    "amounts": {sheet.col(self.comp_col[k]): str(a) for k, a in row_amounts.items() if k in active},
                    "call_total": str(call_total),
                    "dist_total": str(dist_total),
                    "prior_contributions": str(prior),
                    "remaining_commitment": str(remaining),
                }
                if self.defect == "hidden_populated_row" and inv.name == "Juniper Hollow Partners":
                    sheet.hidden_rows.add(sheet.row(rr))

            # Subtotals: LP, GP, vehicle total.
            numeric_cols = sorted(set(col_sums_lp) | set(col_sums_gp))
            first_lp = min(block["investor_rows"].values())
            last_lp = max(block["investor_rows"].values())
            sheet.put(block["lp_subtotal"], cols["investor"], "Limited Partners")
            if block["gp_subtotal"]:
                sheet.put(block["gp_subtotal"], cols["investor"], "General Partner")
            sheet.put(block["total_row"], cols["investor"],
                      f"Total {vehicle}" if len(self.vehicles) > 1 else "Total Partnership")
            for col in numeric_cols:
                if col == cols["close"]:
                    continue
                lp_first = first_lp
                lp_value = col_sums_lp.get(col, Decimal("0"))
                if self.defect == "short_subtotal_range" and col == cols["rf_prior"] and block is self.ablocks[0]:
                    lp_first = first_lp + 1  # the subtotal skips the first LP
                    lp_value -= Decimal(str(sheet.get(sheet.coord(first_lp, col)).cached))
                if col == cols.get("outstanding") and block is self.ablocks[0]:
                    lp_first = first_lp + 2  # the hidden column's subtotal skips the first two LPs
                    lp_value -= sum(_cached_decimal(sheet, sheet.coord(r_, col)) for r_ in (first_lp, first_lp + 1))
                fmt = PCT_FMT if col in (cols["commitment_pct"], cols.get("contributed_pct")) else MONEY_FMT
                sheet.put(block["lp_subtotal"], col, formula=(
                    f"SUM({sheet.coord(lp_first, col)}:{sheet.coord(last_lp, col)})"), cached=_num(lp_value), fmt=fmt)
                gp_value = col_sums_gp.get(col, Decimal("0"))
                if block["gp_row"]:
                    sheet.put(block["gp_subtotal"], col, formula=f"SUM({sheet.coord(block['gp_row'], col)})",
                              cached=_num(gp_value), fmt=fmt)
                    total_formula = f"{sheet.coord(block['lp_subtotal'], col)}+{sheet.coord(block['gp_subtotal'], col)}"
                else:
                    total_formula = sheet.coord(block["lp_subtotal"], col)
                sheet.put(block["total_row"], col, formula=total_formula, cached=_num(lp_value + gp_value), fmt=fmt)
                if vehicle != self.lookthrough:
                    grand[col] = grand.get(col, Decimal("0")) + lp_value + gp_value
            vehicle_layouts.append({
                "name": vehicle,
                "title_row": sheet.row(block["title_row"]) if "title_row" in block else None,
                "driver_row": sheet.row(vdriver_row),
                "investor_rows": [sheet.row(first_lp), sheet.row(last_lp)],
                "gp_rows": [sheet.row(block["gp_row"])] if block["gp_row"] else [],
                "subtotal_rows": {"limited_partners": sheet.row(block["lp_subtotal"]),
                                  "general_partner": sheet.row(block["gp_subtotal"]) if block["gp_subtotal"] else None,
                                  "total": sheet.row(block["total_row"])},
            })

        additive_blocks = [b for b in self.ablocks if b["name"] != self.lookthrough]
        if len(self.vehicles) > 1:
            gt = rows["grand_total"]
            sheet.put(gt, cols["investor"], "Grand Total")
            for col, value in grand.items():
                if col == cols["close"]:
                    continue
                fmt = PCT_FMT if col in (cols["commitment_pct"], cols.get("contributed_pct")) else MONEY_FMT
                sheet.put(gt, col, formula="+".join(sheet.coord(b["total_row"], col) for b in additive_blocks),
                          cached=_num(value), fmt=fmt)

        # Check row: grand totals minus fund drivers (the vehicle drivers added up when no fund row exists).
        check = rows["check"]
        sheet.put(check, cols["investor"], "Check")
        for key in [*self.comp_col, "call_total", "dist_total", "cash_due"]:
            col = self.comp_col.get(key) or cols[key]
            total = grand.get(col, Decimal("0"))
            if v.shared_driver:
                driver_cells = [sheet.coord(b["driver_row"], col) for b in additive_blocks]
                driver_value = sum(_cached_decimal(sheet, c) for c in driver_cells)
                formula = f"ROUND({sheet.coord(rows['grand_total'], col)}-({'+'.join(driver_cells)}),2)"
            else:
                driver_cached = sheet.cells.get(sheet.coord(fund_driver_row, col))
                driver_value = Decimal("0") if driver_cached is None else Decimal(str(
                    driver_cached.cached if driver_cached.formula else driver_cached.value))
                formula = f"ROUND({sheet.coord(rows['grand_total'], col)}-{sheet.coord(fund_driver_row, col)},2)"
            sheet.put(check, col, formula=formula, cached=_num(r2(total - driver_value)), fmt=MONEY_FMT)
        if "outstanding" in cols:
            sheet.put(header, cols["outstanding"], "Outstanding")
            sheet.hidden_cols.add(sheet.col(cols["outstanding"]))

        if self.defect == "referenced_hidden_sheet_error":
            memo = r2(self.lps()[0].commitment * Decimal("0.0021"))
            sheet.put(rows["notice"], cols["error_cell"], formula=f"{q('3rd Close Rebalance')}!$B$4", cached=_num(memo),
                      fmt=MONEY_FMT)
        if self.defect == "formula_error":
            sheet.put(self.ablocks[0]["lp_subtotal"], cols["error_cell"], formula="#REF!", cached="#REF!")

        sheet.view = "normal" if self.defect == "normal_view" else "pageBreakPreview"
        self.sheets.append(sheet)
        self.alloc_grand = grand

        components_layout = []
        for comp in (*CALL_COMPONENTS, *DIST_COMPONENTS):
            components_layout.append({
                "column": sheet.col(self.comp_col[comp.key]),
                "header": comp_headers[comp.key],
                "component_type": comp.component_type,
                "side": comp.side,
                "active": comp.key in active,
            })
        if self.defect == "stale_component":
            components_layout.append({"column": sheet.col(cols["stale"]), "header": "Placement Fees (Capital Call #3)",
                                      "component_type": "placement_fee", "side": "call", "active": False})
        column_map = {
            "investor": sheet.col(cols["investor"]),
            "affiliate_flag": sheet.col(cols["affiliate"]),
            "commitment": sheet.col(cols["commitment"]),
            "commitment_pct": sheet.col(cols["commitment_pct"]),
            "late_interest": sheet.col(cols["late_interest"]),
            "cash_due": sheet.col(cols["cash_due"]),
            "received": sheet.col(cols["received"]),
            "received_date": sheet.col(cols["received_date"]),
        }
        if "contributed_pct" in cols:
            column_map["distribution_basis"] = sheet.col(cols["contributed"])
            column_map["distribution_basis_pct"] = sheet.col(cols["contributed_pct"])
        if "withholding" in cols:
            column_map["tax_withholding"] = sheet.col(cols["withholding"])
        self.layouts[sheet.name] = {
            "role": "allocation",
            "sheet": sheet.name,
            "header_row": sheet.row(header),
            "fund_driver_row": sheet.row(fund_driver_row),
            "event": {
                "event_type": event.kind,
                "label": event.short_label,
                "label_cell": self.event_label_coord,
                "notice_date_cell": notice_coord,
                "due_date_cell": due_coord,
                "carried_interest_rate_cell": getattr(self, "carry_rate_coord", None),
            },
            "columns": column_map,
            "components": components_layout,
            "event_total_columns": [
                {"column": sheet.col(cols["call_total"]), "side": "call"},
                {"column": sheet.col(cols["dist_total"]), "side": "distribution"},
            ],
            "roll_forward": {
                "commitment": sheet.col(cols["rf_commitment"]),
                "prior_contributions": sheet.col(cols["rf_prior"]),
                "prior_recallable": sheet.col(cols["rf_prior_recallable"]),
                "current_call": sheet.col(cols["rf_current_call"]),
                "current_recallable": sheet.col(cols["rf_current_recallable"]),
                "remaining_commitment": sheet.col(cols["rf_remaining"]),
                "adjustments": [sheet.col(cols["rf_waiver"])] if "rf_waiver" in cols else [],
            },
            "vehicles": vehicle_layouts,
            "grand_total_row": sheet.row(rows["grand_total"]),
            "check_rows": [sheet.row(check)],
        }
        self.truth["allocation"] = {
            "investors": truth_investors,
            "drivers": {sheet.col(self.comp_col[k]): str(a) for k, a in drivers.items() if k in active},
            "event_gross": str(sum(drivers[k] for k in active)),
            "notice_date": self.notice_date.isoformat(),
            "due_date": due_date.isoformat(),
            "additive_vehicles": [vname for vname in self.vehicles if vname != self.lookthrough],
        }

    # -- Summary --------------------------------------------------------------------------

    def _build_summary(self) -> None:
        v = self.variant
        sheet = SheetBuilder(self.names["summary"], "summary", v.row_offset, v.col_offset)
        alloc = self.alloc_sheet
        a = q(alloc.name)
        event = self.current
        cols = self.acols
        sheet.put(2, 2, formula=f"{a}!{alloc.coord(self.arows['fund_name'], 2)}", cached=FUND_NAME)
        word = "Distribution - payable " if event.kind == "distribution" else "Capital Call - due "
        if self.defect == "ok_summary_single_text_date":
            word = "Capital Call Summary due "  # the only date on the sheet is inside the title
        title = f"{word}{self.due_date:%B} {self.due_date.day}, {self.due_date.year}"
        title_coord = sheet.put(3, 2, formula=f"\"{word}\"&TEXT({a}!{self.due_coord},\"mmmm d, yyyy\")", cached=title)
        if self.defect == "ok_summary_single_text_date":
            notice_cell, due_cell = None, title_coord
        else:
            sheet.put(5, 3, "Notice Date")
            notice_cell = sheet.put(5, 4, formula=f"{a}!{self.notice_coord}", cached=_serial(self.notice_date), fmt=DATE_LONG_FMT)
            sheet.put(6, 3, "Payment Date" if event.kind == "distribution" else "Due Date")
            due_cell = sheet.put(6, 4, formula=f"{a}!{self.due_coord}", cached=_serial(self.due_date), fmt=DATE_LONG_FMT)
        if self.defect == "referenced_merge_tab_ref_errors":
            sheet.put(40, 8, formula="'Old Merge'!F2", cached="#REF!")  # a stray link into the leftover template
        # One block for the fund, or one block per vehicle (its total row) when the Summary repeats them.
        if v.summary_sections:
            blocks = [(b["name"], alloc.row(b["total_row"]), None) for b in self.ablocks]
        else:
            blocks = [(None, alloc.row(self.arows["grand_total"]), self.alloc_grand)]
        r = 8
        sections_layout = []
        section_totals: list[Decimal] = []
        for index, (vehicle, total_row, grand) in enumerate(blocks):

            def amount_at(col_index: int) -> Decimal:
                if grand is not None:
                    return grand[col_index]
                return _cached_decimal(alloc, alloc.coord(total_row - alloc.row_offset, col_index))

            if vehicle is not None:
                block = self.ablocks[index]
                if "title_row" in block:
                    sheet.put(r, 2, formula=f"{a}!{alloc.coord(block['title_row'], cols['investor'])}", cached=vehicle)
                else:
                    sheet.put(r, 2, vehicle)
                heading_cell = sheet.coord(r, 2)
                r += 1
            else:
                heading_cell = None
            sheet.put(r, 2, "Total Commitments" if v.alt_headers else "Total Fund Commitments")
            total_commit = amount_at(cols["commitment"])
            commit_cell = sheet.put(r, 4, formula=f"{a}!{alloc.col(cols['commitment'])}{total_row}",
                                    cached=_num(total_commit), fmt=MONEY_FMT)
            sheet.put(r, 5, "% of commitment")
            r += 2
            lines = []
            sides = []
            for side, title_text in (("call", "Current Capital Call:"), ("distribution", "Current Distribution:")):
                keys = [k for k in event.drivers if C_BY_KEY[k].side == side]
                if not keys:
                    continue
                sheet.put(r, 3, title_text)
                r += 1
                first = r
                for key in keys:
                    col = alloc.col(self.comp_col[key])
                    amount = amount_at(self.comp_col[key])
                    label = self.comp_headers[key]
                    sheet.put(r, 3, formula=f"{a}!{col}{alloc.row(self.arows['header'])}", cached=label)
                    if self.defect == "summary_line_hardcoded" and key == "expenses" and index == 0:
                        amount = amount - Decimal("0.05")
                        amount_cell = sheet.put(r, 4, _num(amount), fmt=MONEY_FMT)
                    else:
                        amount_cell = sheet.put(r, 4, formula=f"{a}!{col}{total_row}", cached=_num(amount), fmt=MONEY_FMT)
                    sheet.put(r, 5, formula=f"{amount_cell}/{_abs(commit_cell)}",
                              cached=float(amount / total_commit) if total_commit else 0.0, fmt=PCT_FMT)
                    lines.append({"label_cell": sheet.coord(r, 3), "amount_cell": amount_cell,
                                  "component_type": C_BY_KEY[key].component_type, "side": side})
                    r += 1
                if side == "distribution" and self.withholding:
                    sheet.put(r, 3, "Tax Withholding")
                    amount = amount_at(cols["withholding"])
                    amount_cell = sheet.put(r, 4, formula=f"{a}!{alloc.col(cols['withholding'])}{total_row}",
                                            cached=_num(amount), fmt=MONEY_FMT)
                    lines.append({"label_cell": sheet.coord(r, 3), "amount_cell": amount_cell,
                                  "component_type": "tax_withholding", "side": side})
                    r += 1
                total = sum(Decimal(str(sheet.get(line["amount_cell"]).cached if sheet.get(line["amount_cell"]).formula
                                        else sheet.get(line["amount_cell"]).value))
                            for line in lines if line["side"] == side)
                label = "Total Current Capital Call" if side == "call" else "Total Current Distribution"
                sheet.put(r, 3, label)
                total_cell = sheet.put(r, 4, formula=f"SUM({sheet.coord(first, 4)}:{sheet.coord(r - 1, 4)})",
                                       cached=_num(total), fmt=MONEY_FMT)
                sides.append({"side": side, "total_cell": total_cell, "total": total})
                r += 2
            net = sum(x["total"] for x in sides)
            sheet.put(r, 3, "Total Net Cash Due" if len(sides) > 1 or sides[0]["side"] == "call"
                      else "Total Cash to LPs")
            net_cell = sheet.put(r, 4, formula="+".join(x["total_cell"] for x in sides), cached=_num(net), fmt=MONEY_FMT)
            sheet.put(r, 6, "check")
            alloc_cash = amount_at(cols["cash_due"])
            check_cell = sheet.put(r, 7, formula=f"{a}!{alloc.col(cols['cash_due'])}{total_row}-{net_cell}",
                                   cached=_num(alloc_cash - net), fmt=MONEY_FMT)
            r += 3
            section_totals.append(net)
            sections_layout.append({
                "vehicle": vehicle,
                "title_cell": heading_cell,
                "fund_commitment_cell": commit_cell,
                "component_lines": lines,
                "section_totals": [{"side": x["side"], "cell": x["total_cell"]} for x in sides],
                "event_total_cell": net_cell,
                "check_cells": [check_cell],
            })
        first = sections_layout[0]
        self.sheets.append(sheet)
        self.layouts[sheet.name] = {
            "role": "summary",
            "sheet": sheet.name,
            "title_cell": title_coord,
            "notice_date_cell": notice_cell,
            "due_date_cell": due_cell,
            "fund_commitment_cell": first["fund_commitment_cell"],
            "component_lines": first["component_lines"],
            "section_totals": first["section_totals"],
            "event_total_cell": first["event_total_cell"],
            "check_cells": first["check_cells"],
            "sections": sections_layout if v.summary_sections else [],
        }
        self.truth["summary"] = {"event_total": str(section_totals[0]),
                                 "check": str(_cached_decimal(sheet, first["check_cells"][0])),
                                 "section_totals": [str(t) for t in section_totals]}

    # -- Merge tabs -----------------------------------------------------------------------

    def _build_merge_tabs(self) -> None:
        alloc = self.alloc_sheet
        a = q(alloc.name)
        event = self.current
        cols = self.acols
        alt_ids = self.variant.name == "shifted"
        active = [k for k in event.drivers]
        label_cached = alloc.get(self.event_label_coord).value
        for vehicle, tab in self.merge_names.items():
            sheet = SheetBuilder(tab, "merge")
            if self.variant.hidden_merge:
                sheet.state = "hidden"  # hidden once the notices were generated (reference client)
            tab_keys = [k for k in active if not (self.defect == "merge_stale_columns" and k == "expenses")]
            headers = ["Investor", "Short Name", "Letter Date", "Due (Wire) Date",
                       "DX Investor ID" if alt_ids else "Investor ID", "DX Fund ID" if alt_ids else "Fund ID",
                       "File Name", "Commitment Amount", "Commitment %"]
            comp_cols: dict[str, int] = {}
            for i, h in enumerate(headers, start=1):
                sheet.put(1, i, h)
            c = len(headers) + 1
            for key in tab_keys:
                comp_cols[key] = c
                sheet.put(1, c, formula=f"{a}!{alloc.col(self.comp_col[key])}{alloc.row(self.arows['header'])}",
                          cached=self.comp_headers[key])
                c += 1
            if self.withholding:
                comp_cols["withholding"] = c
                sheet.put(1, c, "Tax Withholding")
                c += 1
            total_col = c
            sheet.put(1, total_col, "Cash Due")
            check_col = c + 1
            sheet.put(1, check_col, "Check")
            sheet.put(2, 1, formula=f"{a}!{alloc.coord(self.arows['fund_name'], 2)}", cached=FUND_NAME)
            r = 4
            first_row = r
            lps = [i for i in self.in_vehicle(vehicle) if not i.is_gp]
            totals: dict[int, Decimal] = {}
            for inv in lps:
                arow = self.investor_row[inv.name]
                a_inv = alloc.col(cols["investor"])
                sheet.put(r, 1, formula=f"{a}!{a_inv}{arow}", cached=inv.name)
                short = "{Investor Short Name}" if (self.defect == "placeholder_left" and inv.name == "Tamsin Rourke") \
                    else inv.name
                if short != inv.name:
                    sheet.put(r, 2, short)
                else:
                    sheet.put(r, 2, formula=f"A{r}", cached=inv.name)
                sheet.put(r, 3, formula=f"{a}!{_abs(self.notice_coord)}", cached=_serial(self.notice_date), fmt=DATE_LONG_FMT)
                sheet.put(r, 4, formula=f"{a}!{_abs(self.due_coord)}", cached=_serial(self.due_date), fmt=DATE_LONG_FMT)
                sheet.put(r, 5, inv.investor_id)
                sheet.put(r, 6, inv.fund_id)
                file_name = f"{inv.fund_id}_{inv.investor_id}_{FUND_NAME}_{label_cached}"
                sheet.put(r, 7, formula=f"$F{r}&\"_\"&$E{r}&\"_\"&$A$2&\"_\"&{a}!{_abs(self.event_label_coord)}",
                          cached=file_name)
                sheet.put(r, 8, formula=f"{a}!{alloc.col(cols['commitment'])}{arow}", cached=_num(inv.commitment),
                          fmt=MONEY_FMT)
                pct_cell = alloc.get(f"{alloc.col(cols['commitment_pct'])}{arow}")
                sheet.put(r, 9, formula=f"{a}!{alloc.col(cols['commitment_pct'])}{arow}",
                          cached=pct_cell.cached if pct_cell.formula else pct_cell.value, fmt=PCT_FMT)
                row_total = Decimal("0")
                full_total = sum((self.current_alloc.amounts[k].get(inv.name, Decimal("0")) for k in active), Decimal("0"))
                if self.withholding:
                    full_total += self.current_alloc.amounts["withholding"].get(inv.name, Decimal("0"))
                for key, col in comp_cols.items():
                    source_row = arow
                    amount = self.current_alloc.amounts[key].get(inv.name, Decimal("0"))
                    if self.defect == "merge_wrong_row_link" and key == "investment" \
                            and inv.name == "Juniper Hollow Partners":
                        source_row = self.investor_row["Meridian Endowment Fund"]  # links to the next investor's row
                        amount = self.current_alloc.amounts[key]["Meridian Endowment Fund"]
                    source_col = alloc.col(self.comp_col[key])
                    if self.defect == "ok_sumif_pulls":
                        formula = f"SUMIF({a}!${a_inv}:${a_inv},$A{r},{a}!{source_col}:{source_col})"
                    else:
                        formula = f"{a}!{source_col}{source_row}"
                    sheet.put(r, col, formula=formula, cached=_num(amount), fmt=MONEY_FMT)
                    totals[col] = totals.get(col, Decimal("0")) + amount
                    row_total += amount
                last_key = "withholding" if self.withholding else tab_keys[-1]
                sheet.put(r, total_col, formula=f"SUM({sheet.coord(r, comp_cols[tab_keys[0]])}:"
                                                f"{sheet.coord(r, comp_cols[last_key])})",
                          cached=_num(row_total), fmt=MONEY_FMT)
                totals[total_col] = totals.get(total_col, Decimal("0")) + row_total
                a_cash = alloc.col(cols["cash_due"])
                sheet.put(r, check_col, formula=(f"-{sheet.coord(r, total_col)}+SUMIFS({a}!${a_cash}:${a_cash},"
                                                 f"{a}!${a_inv}:${a_inv},$A{r})"), cached=_num(full_total - row_total),
                          fmt=MONEY_FMT)
                r += 1
            last_row = r - 1
            if self.defect == "ok_inactive_investor_na" and vehicle == "Main Fund":
                sheet.put(r, 1, "Harlan Transfer Trust (transferred)")
                sheet.put(r, 5, formula="INDEX('DX Investor Data'!E:E,MATCH(A%d,'DX Investor Data'!D:D,0))" % r,
                          cached="#N/A")
                sheet.put(r, 7, formula=f"$F{r}&\"_\"&$E{r}", cached="#N/A")
                last_row = r
                r += 1
            r += 1
            total_row = r
            sheet.put(total_row, 1, "TOTAL:")
            for col, value in totals.items():
                sheet.put(total_row, col, formula=f"SUM({sheet.coord(first_row, col)}:{sheet.coord(last_row, col)})",
                          cached=_num(value), fmt=MONEY_FMT)
            self.sheets.append(sheet)
            self.layouts[sheet.name] = {
                "role": "merge",
                "sheet": sheet.name,
                "vehicle": vehicle,
                "header_row": 1,
                "first_data_row": first_row,
                "last_data_row": last_row,
                "columns": {
                    "investor": "A", "short_name": "B", "letter_date": "C", "due_date": "D",
                    "investor_id": "E", "fund_id": "F", "file_name": "G", "commitment": "H", "commitment_pct": "I",
                    "event_total": get_column_letter(total_col), "check": get_column_letter(check_col),
                },
                "component_columns": [
                    {"column": get_column_letter(col), "component_type": C_BY_KEY[key].component_type,
                     "side": C_BY_KEY[key].side}
                    for key, col in comp_cols.items()
                ],
                "total_row": total_row,
            }

    # -- DX Investor Data -------------------------------------------------------------------

    def _build_investor_data(self) -> None:
        sheet = SheetBuilder(self.names["investor_data"], "investor_data")
        for i, h in enumerate(["Fund Name", "Fund ID", "Fund Tax ID", "Investor Name", "Investor ID",
                               "Investor Tax ID"], start=1):
            sheet.put(1, i, h)
        r = 2
        for inv in self.investors:
            if inv.is_gp:
                continue
            name = inv.name
            if self.defect == "investor_name_mismatch" and name == "Crescent Ridge Investments LP":
                name = "Crescent Ridge Investments, L.P."
            sheet.put(r, 1, FUND_NAME)
            sheet.put(r, 2, inv.fund_id)
            sheet.put(r, 4, name)
            sheet.put(r, 5, inv.investor_id)
            r += 1
        # A transferred-out investor that no longer participates (out of scope for identity checks).
        sheet.put(r, 1, FUND_NAME)
        sheet.put(r, 2, 901)
        sheet.put(r, 4, "Harlan Transfer Trust")
        sheet.put(r, 5, 29999)
        self.sheets.append(sheet)
        self.layouts[sheet.name] = {
            "role": "investor_data",
            "sheet": sheet.name,
            "header_row": 1,
            "first_data_row": 2,
            "last_data_row": r,
            "columns": {"fund_name": "A", "fund_id": "B", "investor_name": "D", "investor_id": "E"},
        }

    # -- Distribution waterfall support tab ---------------------------------------------------

    def _build_waterfall(self) -> None:
        if not self.waterfall:
            return
        sheet = SheetBuilder(WATERFALL_SHEET, "other")
        a = q(self.alloc_sheet.name)
        sheet.put(2, 2, "Carried Interest Waterfall")
        for i, h in enumerate(["Investor", "Vehicle", "Distribution Basis", "Carried Interest"], start=2):
            sheet.put(4, i, h)
        first = 5
        gp_rows = {e["vehicle"]: first + idx for idx, e in enumerate(self.waterfall_rows) if e["is_gp"]}
        for idx, entry in enumerate(self.waterfall_rows):
            r = first + idx
            sheet.put(r, 2, entry["name"])
            sheet.put(r, 3, entry["vehicle"])
            basis = self.alloc_sheet.get(entry["basis"].replace("$", ""))
            sheet.put(r, 4, formula=f"{a}!{entry['basis']}", cached=basis.cached, fmt=PCT_FMT)
            if entry["is_gp"]:
                formula = f"ROUND({a}!{entry['driver']}*{a}!{entry['rate']},2)"
            else:
                formula = f"ROUND(({a}!{entry['driver']}-$E${gp_rows[entry['vehicle']]})*{a}!{entry['basis']},2)"
            if entry["offset"]:
                formula += f"{'+' if entry['offset'] > 0 else '-'}{_fmt_offset(abs(entry['offset']))}"
            sheet.put(r, 5, formula=formula, cached=_num(entry["amount"]), fmt=MONEY_FMT)
        total = first + len(self.waterfall_rows) + 1
        sheet.put(total, 2, "Total")
        sheet.put(total, 5, formula=f"SUM(E{first}:E{total - 2})",
                  cached=_num(sum((e["amount"] for e in self.waterfall_rows), Decimal("0"))), fmt=MONEY_FMT)
        self.sheets.append(sheet)

    # -- Mgmt fee tab -----------------------------------------------------------------------

    def fee_cols(self) -> list[str]:
        """Fee-tab amount columns the Allocation pulls: one per quarter billed by the current event."""
        return ["F", "G"] if self.fee_periods == 2 else ["F"]

    def _fee_rows(self) -> dict[str, Any]:
        lps = [i for i in self.investors if not i.is_gp and i.vehicle != self.lookthrough]
        gps = [i for i in self.investors if i.is_gp]
        first = 9
        lp_rows = {inv.name: first + idx for idx, inv in enumerate(lps)}
        lp_sub = first + len(lps) + 1
        gp_rows = {inv.name: lp_sub + 2 + idx for idx, inv in enumerate(gps)}
        gp_sub = lp_sub + 2 + len(gps) + 1
        return {"lp_rows": lp_rows, "lp_sub": lp_sub, "gp_rows": gp_rows, "gp_sub": gp_sub, "total": gp_sub + 1,
                "check": gp_sub + 2}

    def _fee_total_row(self) -> int:
        return self._fee_rows()["total"]

    def _build_fee_tab(self) -> None:
        sheet = SheetBuilder(self.names["mgmt_fee"], "mgmt_fee")
        alloc = self.alloc_sheet
        a = q(alloc.name)
        cols = self.acols
        event = self.current
        period = event.fee_period or self.prior_events[-1].fee_period
        if self.defect == "stale_fee_period":
            period = "Q2 2026"
        quarters = self.fee_quarters if (self.fee_quarters and "mgmt_fee" in event.drivers) else [period]
        fee_is_current = "mgmt_fee" in event.drivers
        rows = self._fee_rows()
        sheet.put(2, 2, formula=f"{a}!{alloc.coord(self.arows['fund_name'], 2)}", cached=FUND_NAME)
        sheet.put(3, 2, "Management Fee Calculation")
        sheet.put(5, 5, "Mgmt Fee %")
        rate_cell = sheet.put(5, 6, _num(MGMT_FEE_RATE), fmt="0.00%")
        sheet.put(6, 5, "% of year")
        frac_cell = sheet.put(6, 6, _num(FEE_PERIOD_FRACTION), fmt="0.00")
        for i, h in enumerate(["Investor", "Vehicle", "Affiliate?", "Commitment Amount"], start=2):
            sheet.put(8, i, h)
        fee_columns = list(zip(self.fee_cols(), quarters))  # (column letter, quarter)
        for index, (letter, quarter) in enumerate(fee_columns):
            sheet.put(8, 6 + index, f"{quarter} Mgmt Fees (2.0%)")
        round_digits = 0 if self.whole_dollar_fees else 2
        lp_total = Decimal("0")
        per_column_total = {letter: Decimal("0") for letter, _ in fee_columns}
        commit_total = Decimal("0")
        gp_commit_total = Decimal("0")
        for inv in self.investors:
            if inv.is_gp or inv.vehicle == self.lookthrough:
                continue
            rr = rows["lp_rows"][inv.name]
            arow = self.investor_row[inv.name]
            sheet.put(rr, 2, formula=f"{a}!{alloc.col(cols['investor'])}{arow}", cached=inv.name)
            sheet.put(rr, 3, inv.vehicle)
            flag = "Y" if inv.affiliate and self.defect != "affiliate_charged_fee" else "N"
            sheet.put(rr, 4, flag)
            sheet.put(rr, 5, formula=f"{a}!{alloc.col(cols['commitment'])}{arow}", cached=_num(inv.commitment),
                      fmt=MONEY_FMT)
            quarter_fee = (_fee(inv, whole_dollars=self.whole_dollar_fees) if fee_is_current else _fee(inv))
            for index, (letter, _) in enumerate(fee_columns):
                if self.defect == "fee_tab_value_wrong" and inv.name == "Meridian Endowment Fund":
                    sheet.put(rr, 6 + index, _num(quarter_fee), fmt=MONEY_FMT)
                else:
                    sheet.put(rr, 6 + index, formula=f"IF($D{rr}=\"N\",ROUND($E{rr}*$F$5*$F$6,{round_digits}),0)",
                              cached=_num(quarter_fee), fmt=MONEY_FMT)
                per_column_total[letter] += quarter_fee
                lp_total += quarter_fee
            commit_total += inv.commitment
        first_lp = min(rows["lp_rows"].values())
        last_lp = max(rows["lp_rows"].values())
        sheet.put(rows["lp_sub"], 2, "Limited Partners")
        sheet.put(rows["lp_sub"], 5, formula=f"SUM(E{first_lp}:E{last_lp})", cached=_num(commit_total), fmt=MONEY_FMT)
        for letter, _ in fee_columns:
            sheet.put(rows["lp_sub"], 6 + self.fee_cols().index(letter), formula=f"SUM({letter}{first_lp}:{letter}{last_lp})",
                      cached=_num(per_column_total[letter]), fmt=MONEY_FMT)
        for inv in self.investors:
            if not inv.is_gp:
                continue
            rr = rows["gp_rows"][inv.name]
            arow = self.investor_row[inv.name]
            sheet.put(rr, 2, formula=f"{a}!{alloc.col(cols['investor'])}{arow}", cached=inv.name)
            sheet.put(rr, 3, inv.vehicle)
            sheet.put(rr, 4, "GP")
            sheet.put(rr, 5, formula=f"{a}!{alloc.col(cols['commitment'])}{arow}", cached=_num(inv.commitment),
                      fmt=MONEY_FMT)
            gp_commit_total += inv.commitment
            for index, _ in enumerate(fee_columns):
                sheet.put(rr, 6 + index, 0, fmt=MONEY_FMT)
        gp_first, gp_last = min(rows["gp_rows"].values()), max(rows["gp_rows"].values())
        sheet.put(rows["gp_sub"], 2, "General Partner (non paying)")
        sheet.put(rows["gp_sub"], 5, formula=f"SUM(E{gp_first}:E{gp_last})", cached=_num(gp_commit_total), fmt=MONEY_FMT)
        sheet.put(rows["total"], 2, "Total")
        sheet.put(rows["total"], 5, formula=f"E{rows['lp_sub']}+E{rows['gp_sub']}",
                  cached=_num(commit_total + gp_commit_total), fmt=MONEY_FMT)
        for index, (letter, _) in enumerate(fee_columns):
            sheet.put(rows["gp_sub"], 6 + index, formula=f"SUM({letter}{gp_first}:{letter}{gp_last})", cached=0, fmt=MONEY_FMT)
            sheet.put(rows["total"], 6 + index, formula=f"{letter}{rows['lp_sub']}+{letter}{rows['gp_sub']}",
                      cached=_num(per_column_total[letter]), fmt=MONEY_FMT)
        sheet.put(7, 6, formula="+".join(f"{letter}{rows['total']}" for letter, _ in fee_columns), cached=_num(lp_total),
                  fmt=MONEY_FMT)
        gt = alloc.row(self.arows["grand_total"])
        sheet.put(rows["check"], 5, formula=f"{a}!{alloc.col(cols['commitment'])}{gt}-E{rows['total']}", cached=0,
                  fmt=MONEY_FMT)
        sheet.view = "pageBreakPreview"
        self.sheets.append(sheet)
        self.layouts[sheet.name] = {
            "role": "mgmt_fee",
            "sheet": sheet.name,
            "header_row": 8,
            "investor_rows": [first_lp, last_lp],
            "gp_rows": sorted(rows["gp_rows"].values()),
            "columns": {"investor": "B", "vehicle": "C", "affiliate_flag": "D", "commitment": "E"},
            "rate_cells": [rate_cell],
            "period_fraction_cells": [frac_cell],
            "fee_columns": [{"column": letter, "period_label": quarter} for letter, quarter in fee_columns],
            "subtotal_rows": {"limited_partners": rows["lp_sub"], "general_partner": rows["gp_sub"],
                              "total": rows["total"]},
            "check_rows": [rows["check"]],
        }
        self.truth["mgmt_fee"] = {"total": str(lp_total), "period": quarters[0]}

    # -- Other sheets -------------------------------------------------------------------------

    def _build_other_sheets(self) -> None:
        notes = SheetBuilder("Notes", "other")
        notes.put(1, 1, "Support and notes for the capital event workpapers.")
        notes.put(3, 1, "Reviewer sign-off tracked in the fund's document system.")
        self.sheets.append(notes)

        tracker = SheetBuilder("Portfolio Investment Tracker", "other")
        for i, h in enumerate(["Deal #", "Event", "Deal", "Amount Called", "Wire Date"], start=2):
            tracker.put(5, i, h)
        deals = self._deals()
        for r, (n, ev, deal, amount, wire) in enumerate(deals, start=6):
            tracker.put(r, 2, n)
            tracker.put(r, 3, ev)
            tracker.put(r, 4, deal)
            tracker.put(r, 5, _num(amount), fmt=MONEY_FMT)
            tracker.put(r, 6, wire, fmt=DATE_SHORT_FMT if isinstance(wire, dt.date) else None)
        self.sheets.append(tracker)

        legacy = SheetBuilder("3rd Close Rebalance", "other")
        legacy.state = "hidden"
        legacy.put(1, 1, "3rd Close Rebalance (2024)")
        legacy.put(3, 1, "Investor")
        legacy.put(3, 2, "Rebalance Amount")
        for r, inv in enumerate(self.lps()[:4], start=4):
            legacy.put(r, 1, inv.name)
            legacy.put(r, 2, _num(r2(inv.commitment * Decimal("0.0021"))), fmt=MONEY_FMT)
        if self.defect in ("ok_hidden_legacy_errors", "referenced_hidden_sheet_error"):
            legacy.put(9, 2, formula="#REF!*2", cached="#REF!")
        self.sheets.append(legacy)

        if self.defect in HIDDEN_MERGE_DEFECTS:
            # A Merge template left over from an earlier vehicle: hidden, headers only, every row #REF!.
            old = SheetBuilder("Old Merge", "other")
            old.state = "hidden"
            for i, h in enumerate(["Investor", "Short Name", "Investor ID", "Fund ID", "File Name", "Cash Due"], start=1):
                old.put(1, i, h)
            for r in range(2, 7):
                for i in range(1, 7):
                    old.put(r, i, formula="#REF!", cached="#REF!")
            self.sheets.append(old)

    def _deals(self) -> list[tuple]:
        deals = [(1, "CC#1", "Northwind Robotics", Decimal("20000000"), dt.date(2024, 1, 25)),
                 (2, "CC#2", "Solace Diagnostics", Decimal("10000000"), dt.date(2024, 6, 27)),
                 (3, "CC#3", "Tidepool Logistics", Decimal("5000000"), dt.date(2026, 3, 11))]
        if self.stage == "current" and self.current.kind == "capital_call":
            wire = "TBD" if self.defect == "ok_tbd_pending" else dt.date(2026, 6, 12)
            deals.append((4, "CC#4", "Ember Grid Storage", Decimal("8500000"), wire))
        return deals

    def tracker_current_row(self) -> int | None:
        """Tracker row of the current event's deal (capital calls only)."""
        if self.current.kind != "capital_call":
            return None
        return 6 + len(self._deals()) - 1

    def fee_tab_total(self) -> Decimal:
        return sum((self.fee_for(i) for i in self.investors if not i.is_gp), Decimal("0"))

    # -- Cell-level defects -------------------------------------------------------------------

    def _apply_cell_defects(self) -> None:
        # Most defects are data-level (applied while computing values, so every dependent
        # cached value stays consistent). Cell-level ones are applied at build time above.
        pass


def _spread_residual(driver: Decimal, amounts: dict[str, Decimal], members: list[Investor]) -> dict[str, Decimal]:
    """Plug pattern seen in the reference sample: the rounding residual spread as cent
    offsets over the top three eligible LPs instead of a single plug."""
    top = sorted([i for i in members if not i.is_gp and not i.affiliate], key=lambda i: -i.commitment)[:3]
    residual = driver - sum(amounts.values())
    pattern = [Decimal("0.02"), Decimal("0.02"), residual - Decimal("0.04")]
    offsets: dict[str, Decimal] = {}
    for inv, off in zip(top, pattern):
        if off:
            amounts[inv.name] += off
            offsets[inv.name] = off
    return offsets


def _spread_units(driver: Decimal, amounts: dict[str, Decimal], members: list[Investor],
                  on_affiliate: bool = False) -> dict[str, Decimal]:
    """Whole-dollar rounding: the residual (a few dollars) is spread as +/-1 units over the largest
    eligible LPs, one unit each (or, seeded defect, the first unit on the affiliate LP)."""
    residual = driver - sum(amounts.values())
    units = int(abs(residual))
    step = Decimal("1") if residual > 0 else Decimal("-1")
    ranked = sorted([i for i in members if not i.is_gp and not i.affiliate], key=lambda i: -i.commitment)
    targets = ranked[:units]
    if on_affiliate:
        affiliate = next(i for i in members if i.affiliate)
        if not units:  # force one unit on the affiliate, balanced on the largest LP
            amounts[affiliate.name] += step
            amounts[ranked[0].name] -= step
            return {affiliate.name: step, ranked[0].name: -step}
        targets = [affiliate] + ranked[:units - 1]
    offsets: dict[str, Decimal] = {}
    for inv in targets:
        amounts[inv.name] += step
        offsets[inv.name] = step
    return offsets


TYPED_CARRY = Decimal("80000")
WAIVER_AMOUNT = Decimal("100000")  # the waiver reduces the LP's remaining commitment (Allocation: -, ITD: +)

C_BY_KEY = {c.key: c for c in (*CALL_COMPONENTS, *DIST_COMPONENTS, C_WITHHOLDING)}


def _referenced_sheets(sheets: list[SheetBuilder], relevant: list[str]) -> list[str]:
    """Sheets outside ``relevant`` linked to them: referenced by a relevant sheet, or a visible
    sheet that references the Allocation or Summary (FA: support tabs linked to/from them)."""
    names = [s.name for s in sheets]
    hubs = [s.name for s in sheets if s.name in relevant and s.role in ("allocation", "summary")]

    def refs(sheet: SheetBuilder) -> set[str]:
        return {n for n in names if n != sheet.name
                for cell in sheet.cells.values() if cell.formula and f"{q(n)}!" in cell.formula}

    found: set[str] = set()
    for sheet in sheets:
        if sheet.name in relevant:
            found |= refs(sheet) - set(relevant)
        elif sheet.state == "visible" and refs(sheet) & set(hubs):
            found.add(sheet.name)
    return [n for n in names if n in found]


def _clean_facts(event_type: str) -> dict[str, dict[str, Any]]:
    facts: dict[str, dict[str, Any]] = {
        "CE-WB-NO-PLACEHOLDERS": {"placeholder_count": 0, "tbd_cells": 0, "tbd_referenced": 0},
        "CE-ITD-EVENT-BLOCK": {"unclassified_current_columns": 0, "overlay_rows": 0},
        "CE-DATE-CONSISTENCY": {"fee_period_mismatch": False, "distinct_event_numbers": [4 if event_type != "distribution" else 2]},
        "CE-ALLOC-STALE-COMPONENTS": {"stale_columns": 0, "active_with_prior_label": 0},
        "CE-ALLOC-REFERENCE-INTEGRITY": {"period_mismatches": 0, "link_pattern_exceptions": 0},
        "CE-WB-MERGE-TABS": {"inactive_ignored": 0},
    }
    if event_type == "distribution":
        facts["CE-DIST-CARRY-SPLIT"] = {"gp_share_matches_rate": True}
    return facts


def _abs(coord: str) -> str:
    m = re.fullmatch(r"([A-Z]+)(\d+)", coord)
    return f"${m.group(1)}${m.group(2)}"


def _abs_row(coord: str) -> str:
    m = re.fullmatch(r"([A-Z]+)(\d+)", coord)
    return f"{m.group(1)}${m.group(2)}"


def _fmt_offset(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _cached_decimal(sheet: SheetBuilder, coord: str) -> Decimal:
    cell = sheet.cells.get(coord)
    if cell is None:
        return Decimal("0")
    value = cell.cached if cell.formula else cell.value
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value))
    return Decimal("0")


def _jsonable(value: Any) -> Any:
    return json.loads(json.dumps(value, default=str))


# ---------------------------------------------------------------------------
# Writing
# ---------------------------------------------------------------------------


def _write_workbook(sheets: list[SheetBuilder], path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.name)
        ws.sheet_state = sheet.state
        if sheet.view:
            ws.sheet_view.view = sheet.view
        for coord, cell in sheet.cells.items():
            target = ws[coord]
            if cell.formula is not None:
                target.value = "=" + cell.formula
            elif isinstance(cell.value, Decimal):
                target.value = _num(cell.value)
            else:
                target.value = cell.value
            if cell.fmt:
                target.number_format = cell.fmt
        for rng in sheet.merges:
            ws.merge_cells(rng)
        for r in sheet.hidden_rows:
            ws.row_dimensions[r].hidden = True
        for c in sheet.hidden_cols:
            ws.column_dimensions[c].hidden = True
        for c in sheet.outlined_cols:
            ws.column_dimensions[c].outlineLevel = 1
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.xlsx"
        wb.save(raw)
        _inject_cached_values(raw, path, sheets)


_FORMULA_CELL_RE = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/></c>', re.S)


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _inject_cached_values(src: Path, dst: Path, sheets: list[SheetBuilder]) -> None:
    cached_by_index = {
        idx: {coord: cell.cached for coord, cell in sheet.cells.items() if cell.formula is not None}
        for idx, sheet in enumerate(sheets, start=1)
    }
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m:
                cached = cached_by_index[int(m.group(1))]

                def repl(match: re.Match) -> str:
                    coord, attrs, formula = match.group(1), match.group(2), match.group(3)
                    value = cached.get(coord)
                    if value is None:
                        return match.group(0)
                    if isinstance(value, bool):
                        return f'<c r="{coord}"{attrs} t="b"><f>{formula}</f><v>{int(value)}</v></c>'
                    if isinstance(value, str) and value.startswith("#"):
                        return f'<c r="{coord}"{attrs} t="e"><f>{formula}</f><v>{_xml_escape(value)}</v></c>'
                    if isinstance(value, str):
                        return f'<c r="{coord}"{attrs} t="str"><f>{formula}</f><v>{_xml_escape(value)}</v></c>'
                    return f'<c r="{coord}"{attrs}><f>{formula}</f><v>{_num(value)!r}</v></c>'

                data = _FORMULA_CELL_RE.sub(repl, data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def build_fixture(spec: FixtureSpec, out_dir: Path) -> FixtureManifest:
    return _Builder(spec).build(Path(out_dir))


def main() -> None:
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    OUT_DIR.mkdir(parents=True)
    index = []
    for spec in default_specs():
        manifest = build_fixture(spec, OUT_DIR)
        (manifest.path.parent / "manifest.json").write_text(json.dumps(manifest.to_json(), indent=2), encoding="utf-8")
        index.append({"fixture_id": spec.fixture_id, "path": str(manifest.path.relative_to(OUT_DIR))})
    (OUT_DIR / "index.json").write_text(json.dumps(index, indent=2), encoding="utf-8")
    print(f"Wrote {len(index)} fixtures to {OUT_DIR}")


if __name__ == "__main__":
    main()
